from __future__ import annotations

import copy
import csv
import hashlib
import io
import json
import logging
import mimetypes
import re
import uuid
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import settings
from app.models.finance import FinanceImport, FinanceSalaryEntry
from app.models.org import OrgPerson, OrgPosition, OrgUnit
from app.models.position_kpi import PositionKpiMetric, PositionKpiProfile
from app.services.admin.salary_extraction import (
    SalaryExtractionError,
    extract_material_document,
    extract_salary_document,
    read_full_text,
)
from app.services.position_kpi.connect import upsert_generated_catalog
from app.services.position_kpi.daily import (
    department_key,
    get_or_compute_position_kpi,
    resolve_profile,
)
from app.services.position_kpi.extract import build_metric
from app.services.workflows.document import DocumentError, load_attachment_bytes

logger = logging.getLogger(__name__)

IMPORT_KINDS = {"salary", "material_incentive"}
MAX_FILE_BYTES = 25 * 1024 * 1024
_RUSSIAN_NAME_RE = re.compile(r"^[А-Яа-яЁё][А-Яа-яЁё\s.'-]*$")


class FinanceError(Exception):
    def __init__(self, message: str, status_code: int = 400) -> None:
        super().__init__(message)
        self.message = message
        self.status_code = status_code


def _is_russian_name(value: str) -> bool:
    return bool(_RUSSIAN_NAME_RE.fullmatch(value.strip()))


def current_salary(
    db: Session, position_id: str, department: str, as_of: date
) -> FinanceSalaryEntry | None:
    """Salary of the department first, then a legacy row imported without a department."""
    if not position_id:
        return None
    own = (department or "").strip()
    for dept in dict.fromkeys((own, "")):
        row = db.scalar(
            select(FinanceSalaryEntry)
            .where(
                FinanceSalaryEntry.position_id == position_id,
                FinanceSalaryEntry.department == dept,
                FinanceSalaryEntry.is_current.is_(True),
                FinanceSalaryEntry.effective_from <= as_of,
            )
            .order_by(FinanceSalaryEntry.effective_from.desc())
            .limit(1)
        )
        if row is not None:
            return row
    return None


def list_departments(db: Session) -> dict[str, Any]:
    names = {
        name.strip()
        for name in db.scalars(select(OrgUnit.name)).all()
        if (name or "").strip()
    }
    names.update(
        name.strip()
        for name in db.scalars(
            select(OrgPerson.department).where(OrgPerson.is_active.is_(True)).distinct()
        ).all()
        if (name or "").strip()
    )
    rows = [{"name": name} for name in sorted(names, key=department_key)]
    return {"rows": rows, "total": len(rows)}


def _departments_by_key(db: Session) -> dict[str, str]:
    return {department_key(row["name"]): row["name"] for row in list_departments(db)["rows"]}


def position_departments(db: Session) -> dict[str, list[str]]:
    rows = db.execute(
        select(OrgPerson.position_id, OrgPerson.department)
        .where(OrgPerson.is_active.is_(True))
        .distinct()
    ).all()
    result: dict[str, set[str]] = {}
    for position_id, department in rows:
        if position_id and (department or "").strip():
            result.setdefault(position_id, set()).add(department.strip())
    return {key: sorted(value, key=department_key) for key, value in result.items()}


def list_employees(
    db: Session,
    *,
    search: str = "",
    department: str = "",
    limit: int = 500,
    login_fios: list[str] | None = None,
) -> dict[str, Any]:
    stmt = select(OrgPerson).where(OrgPerson.is_active.is_(True)).order_by(OrgPerson.fio)
    needle = search.strip()
    if needle:
        like = f"%{needle}%"
        stmt = stmt.where(
            OrgPerson.fio.ilike(like) | OrgPerson.position.ilike(like)
        )
    if department.strip():
        stmt = stmt.where(OrgPerson.department == department.strip())
    if login_fios is None:
        people = list(db.scalars(stmt.limit(max(1, min(limit, 2000)))).all())
        ordered_people: list[OrgPerson | None] = people
        ordered_fios = [person.fio for person in people]
    else:
        keys = [" ".join(fio.lower().replace("ё", "е").split()) for fio in login_fios]
        people = list(
            db.scalars(
                select(OrgPerson).where(
                    OrgPerson.is_active.is_(True),
                    OrgPerson.fio_key.in_(keys),
                )
            ).all()
        )
        people_by_key = {person.fio_key: person for person in people}
        ordered_people = [people_by_key.get(key) for key in keys]
        ordered_fios = login_fios
    rows = []
    for fio, person in zip(ordered_fios, ordered_people, strict=True):
        if not _is_russian_name(fio):
            continue
        if department.strip() and (person is None or person.department != department.strip()):
            continue
        salary = (
            current_salary(db, person.position_id or "", person.department, date.today())
            if person is not None
            else None
        )
        rows.append(
            {
                "id": person.id if person is not None else "",
                "fio": fio,
                "position": person.position if person is not None else "",
                "position_id": (person.position_id or "") if person is not None else "",
                "department": person.department if person is not None else "",
                "salary": str(salary.amount) if salary else "",
                "currency": salary.currency if salary else "RUB",
                "salary_effective_from": salary.effective_from.isoformat() if salary else "",
            }
        )
    departments = sorted({row["department"] for row in rows if row["department"]})
    return {"rows": rows, "departments": departments, "total": len(rows)}


def list_positions(db: Session, *, limit: int = 2000) -> dict[str, Any]:
    positions = db.scalars(
        select(OrgPosition)
        .where(OrgPosition.is_active.is_(True))
        .order_by(OrgPosition.name)
        .limit(max(1, min(limit, 5000)))
    ).all()
    rows = [{"id": position.id, "name": position.name} for position in positions]
    return {"rows": rows, "total": len(rows)}


def employee_salaries(db: Session, person_id: str) -> dict[str, Any]:
    person = _person(db, person_id)
    employee = {
        "id": person.id,
        "fio": person.fio,
        "position": person.position,
        "department": person.department,
    }
    if not person.position_id:
        return {"employee": employee, "rows": []}
    departments = list(dict.fromkeys(((person.department or "").strip(), "")))
    rows = db.scalars(
        select(FinanceSalaryEntry)
        .where(
            FinanceSalaryEntry.position_id == person.position_id,
            FinanceSalaryEntry.department.in_(departments),
        )
        .order_by(
            FinanceSalaryEntry.effective_from.desc(),
            FinanceSalaryEntry.revision.desc(),
        )
    ).all()
    return {
        "employee": employee,
        "rows": [
            {
                "id": row.id,
                "amount": str(row.amount),
                "currency": row.currency,
                "department": row.department,
                "effective_from": row.effective_from.isoformat(),
                "revision": row.revision,
                "is_current": row.is_current,
                "import_id": row.import_id,
                "replaced_entry_id": row.replaced_entry_id or "",
                "created_at": _iso(row.created_at),
            }
            for row in rows
        ],
    }


def employee_kpi(
    db: Session,
    person_id: str,
    *,
    period_from: date | None = None,
    period_to: date | None = None,
) -> dict[str, Any]:
    person = _person(db, person_id)
    if not person.position:
        raise FinanceError("У сотрудника не указана должность", 404)
    as_of = period_to or date.today()
    result = get_or_compute_position_kpi(
        db,
        person.position,
        as_of=as_of,
        date_from=period_from,
        date_to=period_to,
        subject=person.fio,
        allow_stale=True,
        department=person.department,
    )
    if result.get("tiles"):
        return result
    profile = resolve_profile(db, person.position, as_of=as_of, department=person.department)
    if profile is None:
        return result
    metrics = db.scalars(
        select(PositionKpiMetric)
        .where(PositionKpiMetric.profile_id == profile.id)
        .order_by(PositionKpiMetric.sort_order, PositionKpiMetric.code)
    ).all()
    return {
        **result,
        "tiles": [
            {
                "code": metric.code,
                "name": metric.name,
                "weight": metric.weight,
                "unit": metric.unit,
                "plan": metric.plan_value,
                "score": None,
                "contrib": None,
                "fact": None,
                "evidence": "Методика KPI; фактические данные ещё не подключены",
            }
            for metric in metrics
        ],
    }


def create_import(
    db: Session,
    *,
    kind: str,
    filename: str,
    raw: bytes,
    user_id: str,
    user_fio: str,
) -> dict[str, Any]:
    row = start_import(
        db, kind=kind, filename=filename, raw=raw, user_id=user_id, user_fio=user_fio
    )
    return process_import(db, row.id)


def start_import(
    db: Session,
    *,
    kind: str,
    filename: str,
    raw: bytes,
    user_id: str,
    user_fio: str,
) -> FinanceImport:
    """Store the original and an import row in «parsing»; parsing runs separately."""
    import_kind = kind.strip().lower()
    if import_kind not in IMPORT_KINDS:
        raise FinanceError("Неизвестный тип импорта")
    if not raw:
        raise FinanceError("Файл пустой")
    if len(raw) > MAX_FILE_BYTES:
        raise FinanceError("Файл больше 25 МБ")
    import_id = str(uuid.uuid4())
    safe_name = Path(filename or "file").name
    folder = settings.finance_storage_dir / import_id
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / safe_name
    path.write_bytes(raw)
    row = FinanceImport(
        id=import_id,
        kind=import_kind,
        status="parsing",
        original_name=safe_name,
        storage_path=str(path),
        media_type=mimetypes.guess_type(safe_name)[0] or "application/octet-stream",
        sha256=hashlib.sha256(raw).hexdigest(),
        size_bytes=len(raw),
        created_by=user_id,
        created_by_fio=user_fio,
        draft_json={},
        validation_json={},
    )
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


def process_import_in_background(import_id: str) -> None:
    from app.db.session import SessionLocal

    with SessionLocal() as db:
        try:
            process_import(db, import_id)
        except Exception as exc:  # noqa: BLE001
            logger.exception("finance import failed id=%s", import_id)
            db.rollback()
            row = db.get(FinanceImport, import_id)
            if row is not None and row.status == "parsing":
                row.status = "error"
                row.error_text = str(exc)
                db.commit()


def process_import(db: Session, import_id: str) -> dict[str, Any]:
    row = get_import(db, import_id)
    import_kind = row.kind
    safe_name = row.original_name
    draft: dict[str, Any] | None = None
    parse_error = ""
    try:
        raw = Path(row.storage_path).read_bytes()
        draft = (
            _parse_salary(db, safe_name, raw)
            if import_kind == "salary"
            else _parse_material(db, safe_name, raw, text_cache=_ocr_cache_path(row))
        )
    except Exception as exc:  # noqa: BLE001
        parse_error = str(exc)
    row = get_import(db, import_id)
    try:
        if draft is None:
            raise FinanceError(parse_error or "Файл не разобран")
        validation = validate_draft(db, import_kind, draft)
        row.draft_json = draft
        row.validation_json = validation
        row.rows_total = _draft_count(import_kind, draft)
        row.rows_valid = max(0, row.rows_total - len(validation["errors"]))
        row.status = "review"
    except Exception as exc:  # noqa: BLE001
        row.status = "error"
        row.error_text = str(exc)
    db.commit()
    db.refresh(row)
    if import_kind == "material_incentive" and row.status == "review" and not row.validation_json.get("errors"):
        return _auto_confirm_material(db, row)
    return serialize_import(row)


def _auto_confirm_material(db: Session, row: FinanceImport) -> dict[str, Any]:
    """KPI go straight to the DB; the import stays in review only if writing fails."""
    import_id = row.id
    try:
        return confirm_import(db, import_id)
    except Exception as exc:  # noqa: BLE001
        db.rollback()
        row = get_import(db, import_id)
        row.error_text = f"KPI не записаны автоматически: {exc}"
        db.commit()
        db.refresh(row)
        return serialize_import(row)


def list_imports(
    db: Session, *, kind: str = "", status: str = "", limit: int = 200
) -> dict[str, Any]:
    stmt = select(FinanceImport).order_by(FinanceImport.created_at.desc())
    if kind.strip():
        stmt = stmt.where(FinanceImport.kind == kind.strip())
    if status.strip():
        stmt = stmt.where(FinanceImport.status == status.strip())
    rows = db.scalars(stmt.limit(max(1, min(limit, 1000)))).all()
    changes = _changes_by_import(db, rows)
    items = []
    for row in rows:
        item = serialize_import(row, include_draft=False)
        diff = changes.get(row.id)
        item["changes"] = _changes_summary(diff) if diff is not None else None
        items.append(item)
    return {"rows": items, "total": len(rows)}


def import_changes(db: Session, import_id: str) -> dict[str, Any]:
    row = get_import(db, import_id)
    diff = _changes_by_import(db, [row]).get(row.id)
    if diff is None:
        raise FinanceError("Изменения доступны после разбора файла", 409)
    return {"import_id": row.id, "kind": row.kind, **_changes_summary(diff), **diff}


def _changes_summary(diff: dict[str, Any]) -> dict[str, Any]:
    return {
        "summary": {
            "added": len(diff["added"]),
            "updated": len(diff["updated"]),
            "removed": len(diff["removed"]),
            "unchanged": diff["unchanged"],
            "has_baseline": diff["has_baseline"],
        }
    }


def _changes_by_import(db: Session, rows: list[FinanceImport]) -> dict[str, dict[str, Any]]:
    """Diff of each import against the state built from earlier confirmed imports.

    The baseline is limited to the departments the import covers: a file for one
    department must not report every other department as removed.
    """
    result: dict[str, dict[str, Any]] = {}
    targets = [row for row in rows if row.status in {"review", "confirmed"} and row.draft_json]
    for kind in {row.kind for row in targets}:
        history = db.scalars(
            select(FinanceImport)
            .where(
                FinanceImport.kind == kind,
                FinanceImport.status == "confirmed",
                FinanceImport.confirmed_at.is_not(None),
            )
            .order_by(FinanceImport.confirmed_at, FinanceImport.created_at)
        ).all()
        order = {item.id: index for index, item in enumerate(history)}
        # Unconfirmed imports are compared with everything already written to the DB.
        own = sorted(
            (row for row in targets if row.kind == kind),
            key=lambda row: order.get(row.id, len(history)),
        )
        state: dict[tuple[str, str], dict[str, Any]] = {}
        folded = 0
        for row in own:
            stop = order.get(row.id, len(history))
            while folded < stop:
                state.update(_import_entries(history[folded]))
                folded += 1
            result[row.id] = _diff_entries(kind, state, _import_entries(row))
    return result


def _import_entries(row: FinanceImport) -> dict[tuple[str, str], dict[str, Any]]:
    draft = row.draft_json or {}
    entries: dict[tuple[str, str], dict[str, Any]] = {}
    if row.kind == "salary":
        for item in draft.get("rows") or []:
            if not isinstance(item, dict):
                continue
            name = str(item.get("position_name") or "").strip()
            position = str(item.get("position_id") or "") or _position_key(name)
            if not position:
                continue
            department = str(item.get("department") or "").strip()
            entries[(position, department_key(department))] = {
                "position": name,
                "department": department,
                "amount": _amount_label(item.get("amount")),
                "currency": str(item.get("currency") or "RUB").strip() or "RUB",
                "effective_from": str(item.get("effective_from") or ""),
            }
        return entries
    for item in draft.get("profiles") or []:
        if not isinstance(item, dict):
            continue
        name = str(item.get("position_name") or "").strip()
        if not name:
            continue
        department = str(item.get("department") or "").strip()
        entries[(_position_key(name), department_key(department))] = {
            "position": name,
            "department": department,
            "effective_from": str(draft.get("effective_from") or ""),
            "metrics": _metric_entries(item.get("metrics")),
        }
    return entries


def _metric_entries(value: Any) -> list[dict[str, Any]]:
    if isinstance(value, str):
        try:
            value = json.loads(value)
        except ValueError:
            return []
    metrics = []
    for metric in value if isinstance(value, list) else []:
        if not isinstance(metric, dict) or not str(metric.get("name") or "").strip():
            continue
        plan = metric.get("plan_value", metric.get("plan"))
        metrics.append(
            {
                "name": str(metric["name"]).strip(),
                "weight": int(metric.get("weight") or 0),
                "plan": "" if plan in (None, "") else str(plan),
            }
        )
    return metrics


def _amount_label(value: Any) -> str:
    raw = str(value or "").replace("\u00a0", "").replace(" ", "").replace(",", ".")
    try:
        return f"{Decimal(raw):.2f}"
    except (InvalidOperation, ValueError):
        return str(value or "").strip()


def _metric_label(metric: dict[str, Any]) -> str:
    label = f"{metric['name']} — {metric['weight']}%"
    return f"{label}, цель {metric['plan']}" if metric["plan"] else label


def _entry_changes(kind: str, before: dict[str, Any], after: dict[str, Any]) -> list[dict[str, str]]:
    if kind == "salary":
        changes = []
        if before["amount"] != after["amount"]:
            changes.append({"label": "Оклад", "before": before["amount"], "after": after["amount"]})
        if before["currency"] != after["currency"]:
            changes.append({"label": "Валюта", "before": before["currency"], "after": after["currency"]})
        return changes
    old = {_position_key(metric["name"]): metric for metric in before["metrics"]}
    new = {_position_key(metric["name"]): metric for metric in after["metrics"]}
    changes = []
    for key, metric in new.items():
        previous = old.get(key)
        if previous is None:
            changes.append({"label": "KPI добавлен", "before": "", "after": _metric_label(metric)})
            continue
        if previous["weight"] != metric["weight"]:
            changes.append(
                {
                    "label": f"KPI «{metric['name']}»: вес",
                    "before": f"{previous['weight']}%",
                    "after": f"{metric['weight']}%",
                }
            )
        if previous["plan"] != metric["plan"]:
            changes.append(
                {
                    "label": f"KPI «{metric['name']}»: цель",
                    "before": previous["plan"] or "—",
                    "after": metric["plan"] or "—",
                }
            )
    for key, metric in old.items():
        if key not in new:
            changes.append({"label": "KPI удалён", "before": _metric_label(metric), "after": ""})
    return changes


def _entry_order(entry: dict[str, Any]) -> tuple[str, str]:
    return department_key(entry["department"]), _position_key(entry["position"])


def _diff_entries(
    kind: str,
    state: dict[tuple[str, str], dict[str, Any]],
    entries: dict[tuple[str, str], dict[str, Any]],
) -> dict[str, Any]:
    scope = {department for _, department in entries}
    baseline = {key: value for key, value in state.items() if key[1] in scope}
    added, updated, removed = [], [], []
    unchanged = 0
    for key, entry in entries.items():
        previous = baseline.get(key)
        if previous is None:
            added.append(entry)
            continue
        changes = _entry_changes(kind, previous, entry)
        if changes:
            updated.append({**entry, "changes": changes})
        else:
            unchanged += 1
    for key, entry in baseline.items():
        if key not in entries:
            removed.append(entry)
    return {
        "added": sorted(added, key=_entry_order),
        "updated": sorted(updated, key=_entry_order),
        "removed": sorted(removed, key=_entry_order),
        "unchanged": unchanged,
        "has_baseline": bool(baseline),
    }


def get_import(db: Session, import_id: str) -> FinanceImport:
    row = db.get(FinanceImport, import_id)
    if row is None:
        raise FinanceError("Импорт не найден", 404)
    return row


def update_import_draft(
    db: Session, import_id: str, draft: dict[str, Any]
) -> dict[str, Any]:
    row = get_import(db, import_id)
    if row.status == "confirmed":
        raise FinanceError("Подтверждённый импорт нельзя изменить", 409)
    validation = validate_draft(db, row.kind, draft)
    row.draft_json = draft
    row.validation_json = validation
    row.rows_total = _draft_count(row.kind, draft)
    row.rows_valid = max(0, row.rows_total - len(validation["errors"]))
    row.status = "review"
    db.commit()
    db.refresh(row)
    return serialize_import(row)


def confirm_import(db: Session, import_id: str) -> dict[str, Any]:
    row = get_import(db, import_id)
    if row.status == "confirmed":
        return serialize_import(row)
    validation = validate_draft(db, row.kind, row.draft_json or {})
    row.validation_json = validation
    if validation["errors"]:
        db.commit()
        raise FinanceError("Исправьте ошибки перед подтверждением", 409)
    if row.kind == "salary":
        _confirm_salaries(db, row)
    else:
        _confirm_material(db, row)
    row.status = "confirmed"
    row.error_text = ""
    row.confirmed_at = datetime.now(timezone.utc)
    db.commit()
    db.refresh(row)
    return serialize_import(row)


def serialize_import(row: FinanceImport, *, include_draft: bool = True) -> dict[str, Any]:
    result = {
        "id": row.id,
        "kind": row.kind,
        "status": row.status,
        "filename": row.original_name,
        "media_type": row.media_type,
        "sha256": row.sha256,
        "size_bytes": row.size_bytes,
        "created_by": row.created_by,
        "created_by_fio": row.created_by_fio,
        "validation": row.validation_json or {},
        "error": row.error_text,
        "rows_total": row.rows_total,
        "rows_valid": row.rows_valid,
        "created_at": _iso(row.created_at),
        "updated_at": _iso(row.updated_at),
        "confirmed_at": _iso(row.confirmed_at),
    }
    if include_draft:
        result["draft"] = row.draft_json or {}
    return result


def validate_draft(db: Session, kind: str, draft: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    errors: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    known_departments = _departments_by_key(db)
    items = draft.get("rows") if kind == "salary" else draft.get("profiles")
    seen: dict[tuple[str, str], int] = {}
    for index, item in enumerate(items or []):
        if not isinstance(item, dict):
            continue
        department = str(item.get("department") or "").strip()
        if kind == "salary" and not department:
            errors.append(
                {"row": index, "field": "department", "message": "Выберите подразделение"}
            )
        elif department and department_key(department) not in known_departments:
            errors.append(
                {
                    "row": index,
                    "field": "department",
                    "message": f"Подразделение «{department}» не найдено в справочнике",
                }
            )
        position = str(item.get("position_id") or item.get("position_name") or "").strip()
        if position:
            key = (department_key(position), department_key(department))
            if key in seen:
                errors.append(
                    {
                        "row": index,
                        "field": "department",
                        "message": (
                            f"Должность и подразделение повторяют строку {seen[key] + 1}"
                        ),
                    }
                )
            else:
                seen[key] = index
    if kind == "salary":
        for index, item in enumerate(draft.get("rows") or []):
            position_id = str(item.get("position_id") or "")
            if not position_id or db.get(OrgPosition, position_id) is None:
                errors.append(
                    {
                        "row": index,
                        "field": "position_id",
                        "message": "Выберите должность",
                    }
                )
            try:
                if Decimal(str(item.get("amount") or "0").replace(" ", "").replace(",", ".")) <= 0:
                    raise InvalidOperation
            except (InvalidOperation, ValueError):
                errors.append({"row": index, "field": "amount", "message": "Некорректный оклад"})
            try:
                date.fromisoformat(str(item.get("effective_from") or ""))
            except ValueError:
                errors.append({"row": index, "field": "effective_from", "message": "Укажите дату"})
    else:
        try:
            date.fromisoformat(str(draft.get("effective_from") or ""))
        except ValueError:
            errors.append({"field": "effective_from", "message": "Укажите дату вступления"})
        for index, profile in enumerate(draft.get("profiles") or []):
            if not str(profile.get("position_name") or "").strip():
                errors.append({"row": index, "field": "position_name", "message": "Укажите должность"})
            metrics = profile.get("metrics") if isinstance(profile.get("metrics"), list) else []
            total = sum(int(metric.get("weight") or 0) for metric in metrics if isinstance(metric, dict))
            if not metrics:
                errors.append({"row": index, "field": "metrics", "message": "Добавьте KPI"})
            elif total != 100:
                errors.append({"row": index, "field": "metrics", "message": f"Сумма весов KPI: {total}%"})
    return {"errors": errors, "warnings": warnings}


def _parse_salary(db: Session, name: str, raw: bytes) -> dict[str, Any]:
    try:
        extracted = extract_salary_document(name, raw)
    except SalaryExtractionError as exc:
        raise FinanceError(str(exc)) from exc
    positions = db.scalars(
        select(OrgPosition).where(OrgPosition.is_active.is_(True))
    ).all()
    by_name = {_position_key(position.name): position for position in positions}
    departments = _departments_by_key(db)
    staffed = position_departments(db)
    effective_from = str(extracted.get("effective_from") or "")
    rows = []
    for values in extracted.get("rows") or []:
        if not isinstance(values, dict):
            continue
        position_name = str(values.get("position_name") or values.get("position") or "").strip()
        position = by_name.get(_position_key(position_name))
        department = str(values.get("department_name") or values.get("department") or "").strip()
        department = departments.get(department_key(department), department)
        if not department and position is not None and len(staffed.get(position.id, [])) == 1:
            department = staffed[position.id][0]
        rows.append(
            {
                "row_id": str(uuid.uuid4()),
                "position_id": position.id if position else "",
                "position_name": position.name if position else position_name,
                "department": department,
                "amount": str(values.get("amount") or "").replace("\u00a0", " ").strip(),
                "currency": str(values.get("currency") or "RUB").strip() or "RUB",
                "effective_from": str(values.get("effective_from") or effective_from),
            }
        )
    if not rows:
        rows.append(
            {
                "row_id": str(uuid.uuid4()),
                "position_id": "",
                "position_name": "",
                "department": "",
                "amount": "",
                "currency": "RUB",
                "effective_from": effective_from,
            }
        )
    return {"rows": rows}


def _ocr_cache_path(row: FinanceImport) -> Path:
    return Path(row.storage_path).parent / "recognized.txt"


def _parse_material(
    db: Session, name: str, raw: bytes, *, text_cache: Path | None = None
) -> dict[str, Any]:
    """Re-parsing an import reuses the recognized text: OCR of a scan takes minutes."""
    if text_cache is not None and text_cache.is_file():
        text = text_cache.read_text(encoding="utf-8")
    else:
        try:
            text = read_full_text(name, raw)
        except SalaryExtractionError as exc:
            raise FinanceError(str(exc)) from exc
        if text_cache is not None:
            text_cache.write_text(text, encoding="utf-8")
    try:
        extracted = extract_material_document(name, text)
    except SalaryExtractionError as exc:
        raise FinanceError(str(exc)) from exc
    effective = _parse_iso_date(extracted.get("effective_from")) or _find_date(name) or _find_date(text)
    positions = db.scalars(select(OrgPosition).where(OrgPosition.is_active.is_(True))).all()
    by_name = {_position_key(position.name): position for position in positions}
    departments = _departments_by_key(db)
    staffed = position_departments(db)
    document_department = _canonical_department(departments, extracted.get("department_name"))
    profiles = []
    for item in extracted.get("positions") or []:
        if not isinstance(item, dict):
            continue
        title = str(item.get("position_name") or "").strip()
        if not title:
            continue
        position = _match_position(by_name, title)
        department = (
            _canonical_department(departments, item.get("department_name")) or document_department
        )
        # A regulation without a department covers the position wherever it is staffed.
        targets = (
            [department]
            if department
            else sorted(
                {
                    _canonical_department(departments, value)
                    for value in staffed.get(position.id if position else "", [])
                },
                key=department_key,
            )
            or [""]
        )
        metrics = [
            build_metric(
                str(metric.get("name") or ""),
                metric.get("weight"),
                str(metric.get("formula") or ""),
                index,
                metric.get("plan"),
            )
            for index, metric in enumerate(item.get("metrics") or [], start=1)
            if isinstance(metric, dict)
        ]
        for target in targets:
            profiles.append(
                {
                    "position_id": position.id if position else "",
                    "position_name": position.name if position else title,
                    "source_position_name": title,
                    "department": target,
                    "bonus_base_pct": 100,
                    "bonus_kind": "salary_times_crp_times_sum",
                    "metrics": copy.deepcopy(metrics),
                }
            )
    if not profiles:
        profiles.append(
            {
                "position_id": "",
                "position_name": "",
                "department": "",
                "bonus_base_pct": 100,
                "bonus_kind": "salary_times_crp_times_sum",
                "metrics": [],
            }
        )
    return {
        "effective_from": effective.isoformat() if effective else "",
        "source_title": Path(name).stem,
        "profiles": profiles,
    }


def _read_table(name: str, raw: bytes) -> list[dict[str, Any]]:
    suffix = Path(name).suffix.lower()
    matrix: list[list[Any]] = []
    if suffix in {".xlsx", ".xlsm"}:
        workbook = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
        sheet = workbook.active
        matrix = [list(row) for row in sheet.iter_rows(values_only=True)]
    elif suffix in {".csv", ".txt", ".tsv"}:
        text = raw.decode("utf-8-sig", errors="replace")
        if suffix == ".tsv":
            dialect = csv.excel_tab
        else:
            try:
                dialect = csv.Sniffer().sniff(text[:2048], delimiters=";,\t|")
            except csv.Error:
                dialect = csv.excel
        matrix = [list(row) for row in csv.reader(io.StringIO(text), dialect)]
    else:
        attachment = load_attachment_bytes(name, raw, ocr=True)
        lines = [line for line in str(attachment.get("text") or "").splitlines() if line.strip()]
        matrix = [re.split(r"\t|;|\s{2,}", line.strip()) for line in lines]
    return _map_salary_matrix(matrix)


def _map_salary_matrix(matrix: list[list[Any]]) -> list[dict[str, Any]]:
    aliases = {
        "fio": ("фио", "сотрудник", "работник"),
        "position": ("должность",),
        "amount": ("оклад", "сумма", "размер"),
        "currency": ("валюта",),
        "effective_from": ("дата", "действует с"),
    }
    header_index = -1
    mapping: dict[str, int] = {}
    for idx, row in enumerate(matrix[:30]):
        normalized = [str(cell or "").casefold().replace("ё", "е").strip() for cell in row]
        candidate: dict[str, int] = {}
        for field, names in aliases.items():
            for col, value in enumerate(normalized):
                if any(name == value or name in value for name in names):
                    candidate[field] = col
                    break
        if "fio" in candidate and "amount" in candidate:
            header_index, mapping = idx, candidate
            break
    if header_index < 0:
        return []
    rows: list[dict[str, Any]] = []
    for row in matrix[header_index + 1 :]:
        item = {
            field: row[col] if col < len(row) and row[col] is not None else ""
            for field, col in mapping.items()
        }
        if not str(item.get("fio") or "").strip():
            continue
        raw_date = item.get("effective_from")
        if isinstance(raw_date, (date, datetime)):
            item["effective_from"] = raw_date.date().isoformat() if isinstance(raw_date, datetime) else raw_date.isoformat()
        rows.append(item)
    return rows


def _canonical_department(departments: dict[str, str], value: Any) -> str:
    raw = str(value or "").strip()
    return departments.get(department_key(raw), raw) if raw else ""


def _confirm_salaries(db: Session, import_row: FinanceImport) -> None:
    departments = _departments_by_key(db)
    for item in import_row.draft_json.get("rows") or []:
        position = _position(db, str(item.get("position_id") or ""))
        department = _canonical_department(departments, item.get("department"))
        effective_from = date.fromisoformat(str(item["effective_from"]))
        amount = Decimal(str(item["amount"]).replace(" ", "").replace(",", "."))
        previous = db.scalar(
            select(FinanceSalaryEntry)
            .where(
                FinanceSalaryEntry.position_id == position.id,
                FinanceSalaryEntry.department == department,
                FinanceSalaryEntry.effective_from == effective_from,
                FinanceSalaryEntry.is_current.is_(True),
            )
            .order_by(FinanceSalaryEntry.revision.desc())
            .limit(1)
        )
        revision = (previous.revision + 1) if previous else 1
        if previous:
            previous.is_current = False
        db.add(
            FinanceSalaryEntry(
                id=str(uuid.uuid4()),
                position_id=position.id,
                department=department,
                amount=amount,
                currency=str(item.get("currency") or "RUB")[:8],
                effective_from=effective_from,
                revision=revision,
                is_current=True,
                import_id=import_row.id,
                replaced_entry_id=previous.id if previous else None,
                position_snapshot=position.name,
            )
        )


def _confirm_material(db: Session, import_row: FinanceImport) -> None:
    effective_from = date.fromisoformat(str(import_row.draft_json["effective_from"]))
    departments = _departments_by_key(db)
    for item in import_row.draft_json.get("profiles") or []:
        name = str(item.get("position_name") or "").strip()
        department = _canonical_department(departments, item.get("department"))
        current = db.scalars(
            select(PositionKpiProfile).where(
                PositionKpiProfile.position_name == name,
                PositionKpiProfile.effective_from < effective_from,
                PositionKpiProfile.effective_to.is_(None),
            )
        ).all()
        for profile in current:
            if department_key(profile.department) == department_key(department):
                profile.effective_to = effective_from - timedelta(days=1)
        catalog = dict(item)
        catalog.update(
            {
                "position_name": name,
                "department": department,
                "effective_from": effective_from,
                "source_import_id": import_row.id,
                "source_title": str(import_row.draft_json.get("source_title") or import_row.original_name),
            }
        )
        upsert_generated_catalog(db, position=name, catalog=catalog, modules=[])


def _person(db: Session, person_id: str) -> OrgPerson:
    person = db.get(OrgPerson, person_id)
    if person is None or not person.is_active:
        raise FinanceError("Сотрудник не найден", 404)
    return person


def _position(db: Session, position_id: str) -> OrgPosition:
    position = db.get(OrgPosition, position_id)
    if position is None or not position.is_active:
        raise FinanceError("Должность не найдена", 404)
    return position


def _position_key(value: str) -> str:
    return " ".join(value.casefold().replace("ё", "е").split())


def _match_position(by_name: dict[str, OrgPosition], title: str) -> OrgPosition | None:
    """Exact directory match; for «A / B» cells — the first part found in the directory."""
    exact = by_name.get(_position_key(title))
    if exact is not None:
        return exact
    for part in re.split(r"\s*/\s*", title):
        found = by_name.get(_position_key(part))
        if found is not None:
            return found
    return None


def _parse_iso_date(value: Any) -> date | None:
    try:
        return date.fromisoformat(str(value or "").strip()[:10])
    except ValueError:
        return None


def _find_date(text: str) -> date | None:
    match = re.search(r"(?<!\d)(\d{1,2})[./-](\d{1,2})[./-](20\d{2})(?!\d)", text or "")
    if not match:
        return None
    try:
        return date(int(match.group(3)), int(match.group(2)), int(match.group(1)))
    except ValueError:
        return None


def _draft_count(kind: str, draft: dict[str, Any]) -> int:
    return len(draft.get("rows") or []) if kind == "salary" else len(draft.get("profiles") or [])


def _iso(value: datetime | None) -> str:
    if value is None:
        return ""
    return value.isoformat()
