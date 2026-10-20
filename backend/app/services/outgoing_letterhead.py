"""Фирменный бланк исходящего письма и список адресов получателей."""

from __future__ import annotations

import io
import json
from datetime import datetime
from pathlib import Path
from typing import Any

from app.config import BACKEND_ROOT
from app.services.moscow_time import moscow_now

LETTERHEAD_NAME = "Фирменный бланк.docx"
COMPANY_FALLBACK = "НПО «Турбулентность-ДОН»"
SEEDED_RECIPIENTS = ["zhalibin.max@gmail.com"]
_RECIPIENTS_PATH = BACKEND_ROOT / "storage" / "outgoing_recipients.json"
LOGO_PATH = Path(__file__).resolve().parent.parent / "assets" / "turbulence_don_logo.png"
_EMPTY_GUID = "00000000-0000-0000-0000-000000000000"
COMPANY_LINES = (
    "ГК «Турбулентность-ДОН»",
    "Ростовская область, Мясниковский район",
    "1-ый километр Новошахтинского шоссе, база «Спектр», стр. 6/8",
)

DIRECTION_LABELS = {
    "КоммерческийДиректор": "Коммерческий директор",
    "ПрочиеВнутренние": "Прочие внутренние",
    "ДиректорНПО": "Директор НПО",
    "ТехническийДиректор": "Технический директор",
    "ДиректорПроизводства2": "Директор производства",
    "ФинансовыйДиректор": "Финансовый директор",
    "ЗаместительДиректораПоЭкономическойБезопасности": "Зам. директора по экономической безопасности",
    "ИсполнительныйДиректор": "Исполнительный директор",
}


def signature_lines(
    *,
    fio: str = "",
    position: str = "",
    from_email: str = "",
    work_phone: str = "",
    mobile_phone: str = "",
) -> list[str]:
    """Подпись в конце исходящего: сотрудник, адрес ГК, почта отправителя и телефоны из 1С."""
    lines = ["С уважением,"]
    for item in (fio, position):
        text = " ".join(str(item or "").split())
        if text:
            lines.append(text)
    lines.extend(COMPANY_LINES)
    email = " ".join(str(from_email or "").split())
    if "@" in email:
        lines.append(email)
    work = " ".join(str(work_phone or "").split())
    mobile = " ".join(str(mobile_phone or "").split())
    if work:
        lines.append(f"раб. тел: {work}")
    if mobile:
        lines.append(f"тел. моб.: {mobile}")
    return lines


def phone_role(kind: str) -> str:
    text = " ".join(str(kind or "").split()).casefold().replace("ё", "е")
    if not text or "факс" in text or "домаш" in text:
        return ""
    if "мобил" in text or "сотов" in text:
        return "mobile"
    if "тел" in text or "рабоч" in text or "служеб" in text:
        return "work"
    return ""


def sender_phones(client: Any, fio: str) -> dict[str, str]:
    """Рабочий и мобильный телефоны физического лица из контактной информации 1С."""
    found = {"work": "", "mobile": ""}
    query = " ".join(str(fio or "").split())
    if not query:
        return {"work_phone": "", "mobile_phone": ""}
    try:
        rows = client.get(
            "Catalog_Пользователи",
            filt=f"DeletionMark eq false and substringof('{query.replace(chr(39), chr(39) * 2)}', Description)",
            top=20,
        )
    except Exception:  # noqa: BLE001
        return {"work_phone": "", "mobile_phone": ""}
    needle = query.casefold()
    chosen = None
    for row in rows:
        name = str(row.get("Description") or "").strip()
        if name.casefold() == needle:
            chosen = row
            break
        if chosen is None and needle in name.casefold():
            chosen = row
    if chosen is None:
        return {"work_phone": "", "mobile_phone": ""}
    person = str(chosen.get("ФизическоеЛицо_Key") or "")
    if not person or person.startswith(_EMPTY_GUID):
        return {"work_phone": "", "mobile_phone": ""}
    try:
        contacts = client.get(
            "Catalog_ФизическиеЛица_КонтактнаяИнформация",
            filt=f"Ref_Key eq guid'{person}'",
            top=40,
        )
    except Exception:  # noqa: BLE001
        return {"work_phone": "", "mobile_phone": ""}
    for row in contacts:
        if str(row.get("Тип") or "") != "Телефон":
            continue
        number = " ".join(str(row.get("Представление") or row.get("НомерТелефона") or "").split())
        if not number:
            continue
        label = ""
        kind_key = str(row.get("Вид_Key") or "")
        if kind_key and not kind_key.startswith(_EMPTY_GUID):
            try:
                kinds = client.get("Catalog_ВидыКонтактнойИнформации", ref_key=kind_key)
            except Exception:  # noqa: BLE001
                kinds = []
            if kinds:
                label = str(kinds[0].get("Description") or "")
        role = phone_role(label)
        if role and not found[role]:
            found[role] = number
    return {"work_phone": found["work"], "mobile_phone": found["mobile"]}


def direction_label(value: str) -> str:
    text = " ".join(str(value or "").split())
    return DIRECTION_LABELS.get(text, text)


def recipient_emails() -> list[str]:
    stored: list[str] = []
    try:
        raw = json.loads(_RECIPIENTS_PATH.read_text(encoding="utf-8"))
        if isinstance(raw, list):
            stored = [str(item) for item in raw]
    except (OSError, json.JSONDecodeError):
        stored = []
    merged: list[str] = []
    seen: set[str] = set()
    for item in [*SEEDED_RECIPIENTS, *stored]:
        text = " ".join(str(item).split())
        key = text.casefold()
        if "@" not in text or key in seen:
            continue
        seen.add(key)
        merged.append(text)
    return merged


def remember_recipient(email: str) -> None:
    text = " ".join(str(email or "").split())
    if "@" not in text:
        return
    current = recipient_emails()
    if text.casefold() not in {item.casefold() for item in current}:
        current.append(text)
    _RECIPIENTS_PATH.parent.mkdir(parents=True, exist_ok=True)
    _RECIPIENTS_PATH.write_text(json.dumps(current, ensure_ascii=False, indent=1), encoding="utf-8")


def _set_run(run, *, text: str, size: int, bold: bool = False, italic: bool = False, color: str = "") -> None:
    from docx.shared import Pt, RGBColor

    run.text = text
    run.bold = bold
    run.italic = italic
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def build_outgoing_letterhead(
    *,
    organization: str = "",
    theme: str = "",
    body: str = "",
    email: str = "",
    direction: str = "",
    when: datetime | None = None,
    signature: list[str] | None = None,
) -> bytes:
    """Лист A4: шапка организации, кому, тема и текст письма."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)

    org = organization.strip() or COMPANY_FALLBACK
    if len(org) == 36 and org.count("-") == 4:
        org = COMPANY_FALLBACK

    head = document.add_paragraph()
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run(head.add_run(), text=org, size=16, bold=True, color="1A3A6B")

    line = document.add_paragraph()
    line.paragraph_format.space_before = Pt(2)
    line.paragraph_format.space_after = Pt(8)
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1A3A6B")
    borders.append(bottom)
    line._p.get_or_add_pPr().append(borders)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run(title.add_run(), text="ИСХОДЯЩЕЕ ПИСЬМО", size=13, bold=True)

    stamp = document.add_paragraph()
    stamp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    moment = when or moscow_now().replace(tzinfo=None)
    _set_run(stamp.add_run(), text=moment.strftime("%d.%m.%Y"), size=12)

    def meta(label: str, value: str) -> None:
        text = " ".join(value.split())
        if not text:
            return
        row = document.add_paragraph()
        row.paragraph_format.space_after = Pt(2)
        _set_run(row.add_run(), text=f"{label}: ", size=12, bold=True)
        _set_run(row.add_run(), text=text, size=12)

    meta("Кому", email)
    meta("Направление", direction_label(direction))
    meta("Тема", theme)

    document.add_paragraph()
    lines = [item for item in (body or "").splitlines()] or [""]
    for block in lines:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Cm(1.25)
        paragraph.paragraph_format.space_after = Pt(6)
        _set_run(paragraph.add_run(), text=block, size=12)

    sign = [line for line in (signature or []) if str(line).strip()]
    if sign:
        document.add_paragraph()
        for index, line in enumerate(sign):
            row = document.add_paragraph()
            row.paragraph_format.space_after = Pt(0)
            row.paragraph_format.space_before = Pt(10 if index == 0 else 0)
            _set_run(row.add_run(), text=line, size=12, bold=index == 0)
    if LOGO_PATH.is_file():
        picture = document.add_paragraph()
        picture.paragraph_format.space_before = Pt(12)
        picture.add_run().add_picture(str(LOGO_PATH), width=Cm(16))

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def recipients_path() -> Path:
    return _RECIPIENTS_PATH
