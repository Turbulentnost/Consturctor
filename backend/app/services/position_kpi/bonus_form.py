"""Форма «Индивидуальные целевые показатели премирования» для любого сотрудника.

Строки берутся из каталога KPI его должности. Оценка строки — вклад в итог:
вес × выполнение / 100. При полном выполнении вклад совпадает с весом.
"""

from __future__ import annotations

import calendar
from datetime import date
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from app.services.position_kpi.detail_sheets import attach_kpi_detail_sheets, report_hyperlink
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.position_kpi import PositionKpiMetric, PositionKpiProfile
from app.services.position_kpi.daily import (
    PositionKpiNotFound,
    get_or_compute_position_kpi,
    month_bounds,
    resolve_profile,
)
from app.services.profile_overrides import apply_profile_overrides

_MONTHS = (
    "",
    "январь",
    "февраль",
    "март",
    "апрель",
    "май",
    "июнь",
    "июль",
    "август",
    "сентябрь",
    "октябрь",
    "ноябрь",
    "декабрь",
)

_THIN = Border(
    left=Side(style="thin", color="000000"),
    right=Side(style="thin", color="000000"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="thin", color="000000"),
)
_HEADER_FILL = PatternFill("solid", fgColor="C6E0B4")
_TOTAL_FILL = PatternFill("solid", fgColor="F2F2F2")
_TITLE_FONT = Font(name="Calibri", bold=True, size=14)
_LABEL_FONT = Font(name="Calibri", bold=True, size=11)
_CELL_FONT = Font(name="Calibri", size=11)
_HEADER_FONT = Font(name="Calibri", bold=True, size=10)
_DEADLINE_FONT = Font(name="Calibri", size=11, color="FF0000")
_LINK_FONT = Font(name="Calibri", size=11, color="0563C1", underline="single")
_WRAP = Alignment(wrap_text=True, vertical="center", horizontal="left")
_CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")


class ReportSubjectError(Exception):
    def __init__(self, message: str, status_code: int = 404) -> None:
        super().__init__(message)
        self.message = message
        self.status_code = status_code


def resolve_report_subject(*, fio: str, position: str = "") -> dict[str, str]:
    """ФИО → должность. Явная должность не затирает чужого сотрудника."""
    name = " ".join(str(fio or "").split())
    pos = " ".join(str(position or "").split())
    if not name and not pos:
        raise ReportSubjectError("Укажите сотрудника", 400)
    if name:
        from app.services.app_users import find_app_user_by_fio

        user = find_app_user_by_fio(name)
        if user is not None:
            name = " ".join(str(user.fio or name).split()) or name
            if not pos:
                pos = " ".join(str(user.position or "").split())
        _, pos = apply_profile_overrides(name, "", pos)
        if not pos:
            pos = _position_from_erp(name)
            _, pos = apply_profile_overrides(name, "", pos)
    if not pos:
        who = f"«{name}»" if name else "сотрудника"
        raise ReportSubjectError(
            f"Для {who} не найдена должность. В карточке 1С её нет, форма премирования не из чего собрать.",
            404,
        )
    return {"fio": name, "position": pos}


def _position_from_erp(fio: str) -> str:
    try:
        from app.clients.erp_sql import get_user_profile_by_fio

        profile = get_user_profile_by_fio(fio)
    except Exception:
        return ""
    return " ".join(str(getattr(profile, "position", "") or "").split())


def period_caption(start: date, end: date) -> str:
    if (
        start.year == end.year
        and start.month == end.month
        and start.day == 1
        and end.day == calendar.monthrange(start.year, start.month)[1]
    ):
        return f"{_MONTHS[start.month]} {start.year} г."
    return f"{_fmt_date(start)}–{_fmt_date(end)}"


def _fmt_date(value: date) -> str:
    return value.strftime("%d.%m.%Y")


def _pct(value: float | int | None) -> str:
    if value is None:
        return ""
    number = float(value)
    if abs(number - round(number)) < 0.05:
        return f"{int(round(number))}%"
    return f"{number:.1f}".replace(".", ",") + "%"


def _earned(weight: int, tile: dict[str, Any]) -> float | None:
    contrib = tile.get("contrib")
    if contrib is not None:
        return float(contrib)
    score = tile.get("score")
    if score is None or not weight:
        return None
    return round(float(score) * weight / 100.0, 1)


def form_rows(db: Session, profile: PositionKpiProfile, tiles: list[dict[str, Any]]) -> list[dict[str, Any]]:
    by_code = {str(item.get("code") or ""): item for item in tiles if isinstance(item, dict)}
    metrics = (
        db.execute(
            select(PositionKpiMetric)
            .where(PositionKpiMetric.profile_id == profile.id)
            .order_by(PositionKpiMetric.sort_order, PositionKpiMetric.code)
        )
        .scalars()
        .all()
    )
    rows: list[dict[str, Any]] = []
    for metric in metrics:
        tile = by_code.get(metric.code) or {}
        weight = int(metric.weight or 0)
        detail = tile.get("detail") if isinstance(tile.get("detail"), dict) else {}
        rows.append(
            {
                "code": metric.code,
                "name": metric.name,
                "weight": weight,
                "evidence": str(tile.get("evidence") or "").strip(),
                "earned": _earned(weight, tile),
                "detail": detail,
            }
        )
    return rows


def build_bonus_form_xlsx(
    *,
    fio: str,
    position: str,
    period_from: date,
    period_to: date,
    rows: list[dict[str, Any]],
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ИЦПП"
    _write_sheet(
        sheet,
        fio=fio,
        position=position,
        period_from=period_from,
        period_to=period_to,
        rows=rows,
    )
    names = attach_kpi_detail_sheets(
        workbook,
        rows,
        period_from=period_from,
        period_to=period_to,
        cover=sheet.title,
    )
    _link_reports(sheet, names)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def render_bonus_form(
    db: Session,
    *,
    fio: str,
    position: str,
    date_from: date | None,
    date_to: date | None,
) -> tuple[bytes, str]:
    subject = resolve_report_subject(fio=fio, position=position)
    profile = resolve_profile(db, subject["position"])
    if profile is None:
        raise PositionKpiNotFound(subject["position"])
    as_of = date.today()
    start, end = date_from, date_to
    if start is None or end is None:
        start, end = month_bounds(as_of.year, as_of.month)
    payload = get_or_compute_position_kpi(
        db,
        subject["position"],
        as_of=as_of,
        date_from=start,
        date_to=end,
        allow_stale=True,
    )
    tiles = payload.get("tiles") if isinstance(payload.get("tiles"), list) else []
    rows = form_rows(db, profile, tiles)
    if not rows:
        raise PositionKpiNotFound(subject["position"])
    content = build_bonus_form_xlsx(
        fio=subject["fio"],
        position=subject["position"],
        period_from=start,
        period_to=end,
        rows=rows,
    )
    surname = subject["fio"].split()[0] if subject["fio"] else "сотрудник"
    filename = f"ИЦПП_{surname}_{start.isoformat()[:7]}.xlsx"
    return content, filename


def _write_sheet(
    sheet: Worksheet,
    *,
    fio: str,
    position: str,
    period_from: date,
    period_to: date,
    rows: list[dict[str, Any]],
) -> None:
    widths = {"A": 8, "B": 62, "C": 16, "D": 18, "E": 18, "F": 42, "G": 22}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    sheet.merge_cells("A1:G1")
    sheet["A1"] = "Индивидуальные целевые показатели премирования"
    sheet["A1"].font = _TITLE_FONT
    sheet["A1"].alignment = Alignment(vertical="center", horizontal="left")
    sheet.row_dimensions[1].height = 24

    sheet["A2"] = "на"
    sheet["A2"].font = _LABEL_FONT
    sheet["A2"].alignment = _CENTER
    sheet.merge_cells("B2:C2")
    sheet["B2"] = period_caption(period_from, period_to)
    sheet["B2"].font = _LABEL_FONT
    sheet["B2"].alignment = _CENTER

    _person_block(sheet, start_row=4, label="Оцениваемый работник:", position=position, fio=fio, when=period_from)
    _person_block(sheet, start_row=8, label="Цели установил:", position="", fio="", when=period_from)

    header_row = 12
    headers = (
        "№ п/п",
        "Индивидуальные цели\nОписание",
        "Удельный вес (%)",
        "Ожидаемый результат",
        "Срок выполнения",
        "Фактический результат (заполняется по итогам выполнения целей)",
        "Заключительная оценка выполнения (% выполнения целевого значения)",
    )
    for index, text in enumerate(headers, start=1):
        cell = sheet.cell(header_row, index, text)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _CENTER
        cell.border = _THIN
    sheet.row_dimensions[header_row].height = 48

    deadline = _fmt_date(period_to)
    missing = 0
    earned_total = 0.0
    weight_total = 0
    for offset, row in enumerate(rows, start=1):
        excel_row = header_row + offset
        weight = int(row.get("weight") or 0)
        earned = row.get("earned")
        weight_total += weight
        if earned is None:
            missing += 1
        else:
            earned_total += float(earned)
        values = (
            offset,
            str(row.get("name") or ""),
            _pct(weight),
            "отчёт",
            deadline,
            str(row.get("evidence") or ""),
            _pct(earned if isinstance(earned, (int, float)) else None),
        )
        for index, value in enumerate(values, start=1):
            cell = sheet.cell(excel_row, index, value)
            cell.font = _DEADLINE_FONT if index == 5 else _CELL_FONT
            cell.alignment = _CENTER if index != 2 and index != 6 else _WRAP
            cell.border = _THIN
        sheet.row_dimensions[excel_row].height = 36

    total_row = header_row + len(rows) + 1
    sheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=2)
    total_label = sheet.cell(total_row, 1, "ИТОГО ЗА ИНДИВИДУАЛЬНЫЕ ЦЕЛИ:")
    total_label.font = _LABEL_FONT
    total_label.alignment = Alignment(vertical="center", horizontal="right")
    weight_cell = sheet.cell(total_row, 3, _pct(weight_total))
    score_cell = sheet.cell(total_row, 7, _pct(earned_total) if rows else "")
    for index in range(1, 8):
        cell = sheet.cell(total_row, index)
        cell.border = _THIN
        cell.fill = _TOTAL_FILL
        cell.font = _LABEL_FONT
    weight_cell.alignment = _CENTER
    score_cell.alignment = _CENTER
    sheet.row_dimensions[total_row].height = 22

    sign_row = total_row + 2
    _person_block(
        sheet,
        start_row=sign_row,
        label="Выполнение целей оценил:",
        position="",
        fio="",
        when=period_to,
    )
    comment_row = sign_row + 4
    sheet.cell(comment_row, 1, "Комментарии:").font = _LABEL_FONT
    sheet.merge_cells(start_row=comment_row + 1, start_column=1, end_row=comment_row + 2, end_column=7)
    note = ""
    if missing:
        note = "По части целей факт за период ещё не посчитан, оценка этих строк пустая."
    comment = sheet.cell(comment_row + 1, 1, note)
    comment.font = _CELL_FONT
    comment.alignment = _WRAP
    for index in range(1, 8):
        sheet.cell(comment_row + 1, index).border = _THIN
        sheet.cell(comment_row + 2, index).border = _THIN

    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:1"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.oddFooter.center.text = "ИЦПП"
    sheet.freeze_panes = "A13"
    sheet.sheet_view.showGridLines = False


def _link_reports(sheet: Worksheet, names: list[str]) -> None:
    header_row = 12
    for offset, title in enumerate(names, start=1):
        cell = sheet.cell(header_row + offset, 4)
        cell.value = "отчёт"
        cell.hyperlink = report_hyperlink(cell.coordinate, title)
        cell.font = _LINK_FONT
        cell.alignment = _CENTER


def _person_block(
    sheet: Worksheet,
    *,
    start_row: int,
    label: str,
    position: str,
    fio: str,
    when: date,
) -> None:
    sheet.merge_cells(start_row=start_row, start_column=1, end_row=start_row + 2, end_column=1)
    label_cell = sheet.cell(start_row, 1, label)
    label_cell.font = _LABEL_FONT
    label_cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left")
    pairs = (("должность", position), ("Ф И О", fio), ("дата", _fmt_date(when)))
    for offset, (caption, value) in enumerate(pairs):
        row = start_row + offset
        caption_cell = sheet.cell(row, 2, caption)
        caption_cell.font = _CELL_FONT
        caption_cell.alignment = _CENTER
        caption_cell.border = _THIN
        sheet.merge_cells(start_row=row, start_column=3, end_row=row, end_column=5)
        value_cell = sheet.cell(row, 3, value)
        value_cell.font = _CELL_FONT
        value_cell.alignment = _WRAP
        for column in (3, 4, 5):
            sheet.cell(row, column).border = _THIN
        if offset == 2:
            sign = sheet.cell(row, 6, "подпись")
            sign.font = _CELL_FONT
            sign.alignment = _CENTER
        sheet.row_dimensions[row].height = 20
