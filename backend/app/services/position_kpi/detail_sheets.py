"""Подробный лист расчёта на каждый KPI внутри формы премирования.

Сводная колонка показывает имя листа и ссылается на него. Состав колонок берётся
из строк калькулятора, без отдельной вёрстки на каждую должность.
"""

from __future__ import annotations

import re
from datetime import date, datetime
from typing import Any

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.worksheet import Worksheet

_INVALID_SHEET = re.compile(r"[\[\]\:\*\?\/\\]")
_COVER = "ИЦПП"

_THIN = Border(
    left=Side(style="thin", color="000000"),
    right=Side(style="thin", color="000000"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="thin", color="000000"),
)
_HEADER_FILL = PatternFill("solid", fgColor="C6E0B4")
_TITLE_FONT = Font(name="Calibri", bold=True, size=14)
_LABEL_FONT = Font(name="Calibri", bold=True, size=11)
_CELL_FONT = Font(name="Calibri", size=11)
_HEADER_FONT = Font(name="Calibri", bold=True, size=10)
_LINK_FONT = Font(name="Calibri", size=11, color="0563C1", underline="single")
_GOOD_FONT = Font(name="Calibri", size=11, color="08745F")
_BAD_FONT = Font(name="Calibri", size=11, color="A32020")
_WRAP = Alignment(wrap_text=True, vertical="center", horizontal="left")
_CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")

_COLUMN_ORDER = (
    "kind_label",
    "item_kind",
    "series",
    "number",
    "doc_kind",
    "content",
    "subject",
    "topic",
    "name",
    "plan_count",
    "fact_count",
    "missed",
    "rule",
    "plan_date",
    "fact_date",
    "fact_title",
    "appointed",
    "event_date",
    "event_subject",
    "in_calendar",
    "deadline",
    "due",
    "protocol_number",
    "protocol_date",
    "protocol_status",
    "closed_at",
    "date",
    "posted",
    "decided",
    "entered",
    "status",
    "prepared",
    "issued",
    "in_tracker",
    "complete",
    "in_period",
    "on_time",
    "active",
    "control",
    "returned",
)
_SKIP_ROW_KEYS = {"kind", "theme", "ref"}
_COLUMN_LABELS = {
    "kind_label": "Орган",
    "item_kind": "Тип",
    "series": "Серия",
    "number": "Номер",
    "doc_kind": "Вид",
    "content": "Содержание",
    "subject": "Совещание",
    "topic": "Тема",
    "name": "Название",
    "plan_count": "План",
    "fact_count": "Факт",
    "missed": "Нет факта",
    "rule": "Правило",
    "plan_date": "План",
    "fact_date": "Факт",
    "fact_title": "Встреча в календаре",
    "appointed": "В календаре",
    "event_date": "Факт",
    "event_subject": "Встреча в календаре",
    "in_calendar": "В календаре",
    "deadline": "Срок",
    "due": "Срок наступил",
    "protocol_number": "Протокол",
    "protocol_date": "Дата протокола",
    "protocol_status": "Статус протокола",
    "closed_at": "Закрыт",
    "date": "Дата",
    "posted": "Проведён",
    "decided": "Дата решения",
    "entered": "Внесено",
    "status": "Статус",
    "prepared": "Подготовлен",
    "issued": "Выпущен",
    "in_tracker": "В Excel",
    "complete": "Заполнено",
    "in_period": "В периоде",
    "on_time": "Вовремя",
    "active": "Активное",
    "control": "Контроль",
    "returned": "Возврат",
}
_ITEM_KIND = {"protocol": "протокол", "assignment": "поручение"}
_SUMMARY_ORDER = (
    "z_total",
    "z_on_time",
    "p_total",
    "p_on_time",
    "r_total",
    "r24",
    "r_in_tracker",
    "r_active",
    "r_control",
    "kpi3_1_pct",
    "kpi3_2_pct",
    "v_total",
    "v_errors",
    "target_pct",
    "plan_total",
    "fact_total",
    "in_calendar",
    "violations",
    "fact_pct",
    "score_pct",
)
_SUMMARY_LABELS = {
    "z_total": "Zвсего",
    "z_on_time": "Zвовремя",
    "p_total": "Pвсего",
    "p_on_time": "Pвовремя",
    "r_total": "Rвсего",
    "r24": "R24",
    "r_in_tracker": "В Excel",
    "r_active": "Rактив",
    "r_control": "Rконтроль",
    "kpi3_1_pct": "KPI3.1",
    "kpi3_2_pct": "KPI3.2",
    "v_total": "Vвсего",
    "v_errors": "Vоши",
    "target_pct": "Цель",
    "plan_total": "План, шт",
    "fact_total": "Факт, шт",
    "in_calendar": "В календаре, шт",
    "violations": "Нарушения",
    "fact_pct": "Факт",
    "score_pct": "Оценка",
}
_PERCENT_KEYS = {"kpi3_1_pct", "kpi3_2_pct", "target_pct", "fact_pct", "score_pct"}


def sheet_name(index: int, name: str, used: set[str]) -> str:
    clean = _INVALID_SHEET.sub(" ", str(name or "KPI"))
    clean = " ".join(clean.split()) or "KPI"
    prefix = f"{index}. "
    title = (prefix + clean)[:31].strip()
    if title not in used:
        return title
    suffix = f" {index}"
    return (prefix + clean)[: 31 - len(suffix)].strip() + suffix


def attach_kpi_detail_sheets(
    workbook: Workbook,
    rows: list[dict[str, Any]],
    *,
    period_from: date,
    period_to: date,
    cover: str = _COVER,
) -> list[str]:
    """Добавляет по листу на каждую цель и возвращает имена листов по порядку строк."""
    used = {sheet.title for sheet in workbook.worksheets}
    names: list[str] = []
    for index, row in enumerate(rows, start=1):
        title = sheet_name(index, str(row.get("name") or row.get("code") or "KPI"), used)
        used.add(title)
        sheet = workbook.create_sheet(title)
        write_kpi_detail_sheet(
            sheet,
            row,
            period_from=period_from,
            period_to=period_to,
            cover=cover,
        )
        names.append(title)
    return names


def report_hyperlink(cell_ref: str, sheet_title: str) -> Hyperlink:
    escaped = sheet_title.replace("'", "''")
    return Hyperlink(ref=cell_ref, location=f"'{escaped}'!A1")


def write_kpi_detail_sheet(
    sheet: Worksheet,
    row: dict[str, Any],
    *,
    period_from: date,
    period_to: date,
    cover: str = _COVER,
) -> None:
    detail = row.get("detail") if isinstance(row.get("detail"), dict) else {}
    records = [item for item in (detail.get("rows") or []) if isinstance(item, dict)]
    columns = _columns(records)

    sheet["A1"] = str(row.get("name") or "KPI")
    sheet["A1"].font = _TITLE_FONT
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(columns), 4))
    sheet.row_dimensions[1].height = 24

    back = sheet.cell(2, 1, "К форме")
    back.font = _LINK_FONT
    back.hyperlink = report_hyperlink("A2", cover)
    period = sheet.cell(2, 2, f"{_fmt(period_from)} — {_fmt(period_to)}")
    period.font = _CELL_FONT
    sheet.merge_cells(start_row=2, start_column=2, end_row=2, end_column=4)

    summary = _summary_pairs(detail, row)
    header_at = 4
    for index, (label, value) in enumerate(summary, start=1):
        label_cell = sheet.cell(header_at, index, label)
        label_cell.font = _HEADER_FONT
        label_cell.fill = _HEADER_FILL
        label_cell.alignment = _CENTER
        label_cell.border = _THIN
        value_cell = sheet.cell(header_at + 1, index, value)
        value_cell.font = _LABEL_FONT
        value_cell.alignment = _CENTER
        value_cell.border = _THIN
    sheet.row_dimensions[header_at].height = 22
    sheet.row_dimensions[header_at + 1].height = 22

    table_at = header_at + 3
    if not records or not columns:
        empty = sheet.cell(table_at, 1, "За период нет построчного расчёта.")
        empty.font = _CELL_FONT
        sheet.column_dimensions["A"].width = 42
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1
        sheet.freeze_panes = "A3"
        return

    for index, key in enumerate(columns, start=1):
        cell = sheet.cell(table_at, index, _COLUMN_LABELS.get(key, key))
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _CENTER
        cell.border = _THIN
        sheet.column_dimensions[cell.column_letter].width = _width(key)
    for offset, record in enumerate(records, start=1):
        excel_row = table_at + offset
        for index, key in enumerate(columns, start=1):
            value, kind = _cell_value(key, record.get(key))
            cell = sheet.cell(excel_row, index, value)
            cell.font = _GOOD_FONT if kind is True else _BAD_FONT if kind is False else _CELL_FONT
            cell.alignment = (
                _WRAP
                if key in {"subject", "topic", "name", "content", "rule", "fact_title", "event_subject", "status", "protocol_status"}
                else _CENTER
            )
            cell.border = _THIN
        sheet.row_dimensions[excel_row].height = 20
    sheet.auto_filter.ref = f"A{table_at}:{_column_letter(len(columns))}{table_at + len(records)}"
    sheet.freeze_panes = f"A{table_at + 1}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.print_title_rows = f"{table_at}:{table_at}"
    sheet.oddFooter.center.text = str(row.get("name") or "")[:40]


def _columns(records: list[dict[str, Any]]) -> list[str]:
    present: set[str] = set()
    for record in records:
        present.update(str(key) for key in record)
    present -= _SKIP_ROW_KEYS
    if "kind_label" in present:
        present.discard("kind")
    ordered = [key for key in _COLUMN_ORDER if key in present]
    rest = sorted(present - set(ordered))
    return ordered + rest


def _summary_pairs(detail: dict[str, Any], row: dict[str, Any]) -> list[tuple[str, str]]:
    pairs = [("Вес", _percent(row.get("weight")) if row.get("weight") is not None else "")]
    for key in _SUMMARY_ORDER:
        if key not in detail or detail.get(key) is None:
            continue
        value = detail.get(key)
        text = _percent(value) if key in _PERCENT_KEYS else _plain(value)
        pairs.append((_SUMMARY_LABELS.get(key, key), text))
    if row.get("earned") is not None:
        pairs.append(("Вклад", _percent(row.get("earned"))))
    return pairs


def _cell_value(key: str, value: Any) -> tuple[str, bool | None]:
    if key == "item_kind":
        return _ITEM_KIND.get(str(value or ""), _plain(value)), None
    if isinstance(value, bool):
        return ("да" if value else "нет"), value
    return _plain(value), None


def _plain(value: Any) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, bool):
        return "да" if value else "нет"
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        if abs(value - round(value)) < 0.05:
            return str(int(round(value)))
        return f"{value:.1f}".replace(".", ",")
    if isinstance(value, (dict, list)):
        return ""
    return str(value)


def _percent(value: Any) -> str:
    if value is None or value == "":
        return ""
    try:
        number = float(value)
    except (TypeError, ValueError):
        return str(value)
    if abs(number - round(number)) < 0.05:
        return f"{int(round(number))}%"
    return f"{number:.1f}".replace(".", ",") + "%"


def _fmt(value: date) -> str:
    return value.strftime("%d.%m.%Y")


def _width(key: str) -> int:
    if key in {"subject", "topic", "name", "content", "fact_title", "event_subject"}:
        return 42
    if key == "rule":
        return 28
    if key in {"protocol_number", "number", "status", "protocol_status"}:
        return 22
    return 16


def _column_letter(index: int) -> str:
    letters = ""
    while index:
        index, rest = divmod(index - 1, 26)
        letters = chr(65 + rest) + letters
    return letters
