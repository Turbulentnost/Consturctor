"""Лист показателя из уже собранного отчёта ИЦПП: как сложился факт."""

from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

_MONTH = re.compile(r"^\d{4}-\d{2}$")
_PERCENT_LABELS = {"План", "Факт"}


def report_dir() -> Path:
    return Path(__file__).resolve().parents[3] / "kpi" / "generated"


def fold(text: str) -> str:
    return " ".join(str(text or "").replace("ё", "е").replace("Ё", "Е").casefold().split())


def find_report(fio: str, month: str, root: Path | None = None) -> Path | None:
    """Файл ИЦПП этого сотрудника за месяц. Чужие книги не берём."""
    token = str(month or "").strip()
    person = fold(fio)
    if not person or not _MONTH.fullmatch(token):
        return None
    folder = root or report_dir()
    if not folder.is_dir():
        return None
    for path in sorted(folder.glob(f"ИЦПП_*_{token}.xlsx")):
        if path.name.startswith("~$"):
            continue
        if _cover_fio(path) == person:
            return path
    return None


class ReportFactMissing(Exception):
    def __init__(self, message: str) -> None:
        self.message = message
        super().__init__(message)


def load_metric_fact(fio: str, month: str, metric_name: str, root: Path | None = None) -> dict[str, Any]:
    path = find_report(fio, month, root)
    if path is None:
        raise ReportFactMissing("Отчёт за этот месяц ещё не собран")
    from openpyxl import load_workbook

    book = load_workbook(path, data_only=True)
    try:
        found = breakdown_for_metric(book, metric_name)
    finally:
        book.close()
    if found is None:
        raise ReportFactMissing("В отчёте нет листа этого показателя")
    return found


def breakdown_for_metric(book: Workbook, metric_name: str) -> dict[str, Any] | None:
    sheet = _match_sheet(book, metric_name)
    if sheet is None:
        return None
    return _read_sheet(sheet)


def _cover_fio(path: Path) -> str:
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    try:
        cover = book.worksheets[0]
        for index, row in enumerate(cover.iter_rows(min_row=4, max_row=6, max_col=3, values_only=True), start=4):
            if index == 5 and len(row) >= 3:
                return fold(row[2])
        return ""
    finally:
        book.close()


def _match_sheet(book: Workbook, metric_name: str) -> Worksheet | None:
    wanted = fold(metric_name)
    if not wanted:
        return None
    exact: list[Worksheet] = []
    partial: list[Worksheet] = []
    for sheet in book.worksheets:
        if fold(sheet.title) == "ицпп":
            continue
        title = fold(sheet.cell(1, 1).value)
        if not title:
            continue
        if title == wanted:
            exact.append(sheet)
        elif title in wanted or wanted in title:
            partial.append(sheet)
    if exact:
        return exact[0]
    if len(partial) == 1:
        return partial[0]
    return None


def _read_sheet(sheet: Worksheet) -> dict[str, Any]:
    summary: list[dict[str, str]] = []
    for column in range(1, 8):
        label = _text(sheet.cell(4, column).value)
        if not label:
            break
        value = _text(sheet.cell(5, column).value)
        if label in _PERCENT_LABELS and value and not value.endswith("%"):
            value = f"{value}%"
        summary.append({"label": label, "value": value})
    columns: list[str] = []
    for column in range(1, (sheet.max_column or 0) + 1):
        header = _text(sheet.cell(9, column).value)
        if not header:
            break
        columns.append(header)
    rows: list[list[str]] = []
    for row_no in range(10, (sheet.max_row or 9) + 1):
        values = [_text(sheet.cell(row_no, column).value) for column in range(1, len(columns) + 1)]
        if not any(values):
            break
        rows.append(values)
    return {
        "name": _text(sheet.cell(1, 1).value),
        "period": _text(sheet.cell(2, 2).value),
        "fact_text": _text(sheet.cell(7, 1).value),
        "summary": summary,
        "columns": columns,
        "rows": rows,
        "included": included_flags(columns, rows),
    }


def included_flags(columns: list[str], rows: list[list[str]]) -> list[bool]:
    """Строка вошла в факт. Остальные на экране розово-красные."""
    index = {name: position for position, name in enumerate(columns)}

    def yes(row: list[str], name: str) -> bool:
        position = index.get(name)
        if position is None or position >= len(row):
            return False
        return row[position].casefold() == "да"

    if "Возврат" in index and "Вовремя" not in index:
        return [not yes(row, "Возврат") for row in rows]
    if {"В периоде", "Вовремя", "Активное", "Контроль"} <= set(index):
        period = on_time = active = control = 0
        for row in rows:
            if yes(row, "В периоде"):
                period += 1
                on_time += int(yes(row, "Вовремя"))
            if yes(row, "Активное"):
                active += 1
                control += int(yes(row, "Контроль"))
        kpi31 = on_time / period if period else 1.0
        kpi32 = control / active if active else 1.0
        if kpi32 <= kpi31:
            return [yes(row, "Активное") and yes(row, "Контроль") for row in rows]
        return [yes(row, "В периоде") and yes(row, "Вовремя") for row in rows]
    if "Вовремя" in index and "Срок наступил" in index:
        return [yes(row, "Срок наступил") and yes(row, "Вовремя") for row in rows]
    if "Вовремя" in index:
        return [yes(row, "Вовремя") for row in rows]
    return [True for _row in rows]


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, float):
        return f"{value:.4f}".rstrip("0").rstrip(".")
    if isinstance(value, int):
        return str(value)
    return str(value).strip()
