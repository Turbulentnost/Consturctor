"""Fill the official СТО-34-003 regulation (регламент) Word template."""

from __future__ import annotations

import json
import re
import shutil
from pathlib import Path
from typing import Any

BACKEND_ROOT = Path(__file__).resolve().parents[3]

STO_TEMPLATE_PATH = BACKEND_ROOT / "app" / "static" / "sto_regulation_template.docx"
STO_STRUCTURE_PATH = BACKEND_ROOT / "app" / "static" / "sto_regulation_structure.json"

STO_SECTIONS: tuple[str, ...] = (
    "Назначение и область применения",
    "Нормативные ссылки",
    "Термины и определения",
    "Сокращения",
    "Ответственность",
    "Организация работы",
    "Ресурсы",
    "Документирование и архивирование",
    "Приложения",
)

_ARCHIVE_BOILERPLATE = (
    "8.1 Первый экземпляр Регламента хранится в Архиве. Выдачу учетных экземпляров "
    "на бумажном носителе выдает специалист по ПУ с отметкой о выдаче в Листе выдачи "
    "первого экземпляра документа. Электронная версия выложена в локальной сети: "
    "\\\\192.168.1.198\\Files\\10.СКТБ\\НОРМАТИВНЫЕ ДОКУМЕНТЫ ОРГАНИЗАЦИИ\\"
    "НОРМАТИВНЫЕ ДОКУМЕНТЫ\\Нормативные документы по подразделениям. "
    "Сохранение документов осуществляется ежедневно на сервере системным администратором. "
    "Внесение изменений в документы осуществляет специалист по ПУ по предъявленному "
    "Извещению об изменении в соответствии с СТО-34-003 «Управление документированной информацией»."
)

_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _qn(tag: str) -> str:
    return f"{{{_W_NS}}}{tag}"


def _has_sect_pr(el) -> bool:
    if el.find(_qn("sectPr")) is not None:
        return True
    return el.find(f".//{_qn('sectPr')}") is not None


def _normalize_heading(text: str) -> str:
    cleaned = re.sub(r"^\s*\d+(?:\.\d+)*\s*", "", str(text or "").strip())
    return re.sub(r"\s+", " ", cleaned).casefold()


def _sto_index(title: str) -> int | None:
    key = _normalize_heading(title)
    for idx, name in enumerate(STO_SECTIONS, start=1):
        if key == _normalize_heading(name):
            return idx
    return None


def load_sto_structure() -> dict[str, Any]:
    if not STO_STRUCTURE_PATH.is_file():
        return {"templateKind": "regulation", "sections": [{"title": name} for name in STO_SECTIONS]}
    return json.loads(STO_STRUCTURE_PATH.read_text(encoding="utf-8"))


def fill_sto_regulation(
    document: dict[str, Any],
    out_path: Path,
    *,
    template_path: Path | None = None,
) -> Path:
    """Write a regulation DOCX from Constructor document JSON into the album template."""
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("Для формирования DOCX требуется python-docx") from exc

    template = Path(template_path or STO_TEMPLATE_PATH)
    if not template.is_file():
        raise FileNotFoundError(f"Шаблон регламента не найден: {template}")

    payload = normalize_sto_document(document)
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template, out_path)
    doc = Document(str(out_path))
    _fill_title_page(doc, payload)
    _fill_headers(doc, payload)
    _fill_toc(doc, payload)
    _fill_body(doc, payload)
    doc.save(str(out_path))
    return out_path


def normalize_sto_document(document: dict[str, Any]) -> dict[str, Any]:
    """Map free-form creation JSON onto the nine СТО регламент sections."""
    title = str(document.get("title") or "Регламент").strip() or "Регламент"
    meta = document.get("meta") if isinstance(document.get("meta"), dict) else {}
    incoming = [item for item in (document.get("sections") or []) if isinstance(item, dict)]
    buckets: list[dict[str, Any]] = [
        {"number": str(idx), "title": name, "paragraphs": [], "items": [], "tables": [], "sections": []}
        for idx, name in enumerate(STO_SECTIONS, start=1)
    ]
    leftovers: list[dict[str, Any]] = []
    for section in incoming:
        idx = _sto_index(_section_heading(section))
        if idx is None:
            leftovers.append(section)
            continue
        _merge_section(buckets[idx - 1], section)
    if leftovers:
        work = buckets[5]
        if not work["paragraphs"] and not work["items"] and not work["tables"] and not work["sections"]:
            work["sections"] = leftovers
        else:
            work["sections"].extend(leftovers)
    if not buckets[7]["paragraphs"]:
        buckets[7]["paragraphs"] = [_ARCHIVE_BOILERPLATE]
    if not buckets[8]["paragraphs"] and not buckets[8]["items"] and not buckets[8]["tables"]:
        buckets[8]["paragraphs"] = ["Приложений нет"]
    return {
        "title": title,
        "code": str(document.get("code") or meta.get("code") or "").strip() or _default_code(title),
        "version": str(document.get("version") or meta.get("version") or "01").strip() or "01",
        "year": str(document.get("year") or meta.get("year") or "").strip() or _current_year(),
        "status": str(document.get("status") or meta.get("status") or "проект на согласование").strip(),
        "city": str(document.get("city") or meta.get("city") or "Ростов-на-Дону").strip(),
        "developedBy": str(meta.get("developedBy") or document.get("developedBy") or "").strip(),
        "agreedDept": str(meta.get("agreedDept") or document.get("agreedDept") or "").strip(),
        "agreedDirector": str(
            meta.get("agreedDirector") or document.get("agreedDirector") or "Директор по направлению"
        ).strip(),
        "checkedBy": str(
            meta.get("checkedBy") or document.get("checkedBy") or "Специалист по процессному управлению"
        ).strip(),
        "sections": buckets,
    }


def make_blank_template(source_path: Path, dest_path: Path) -> Path:
    """Copy album chrome from an existing regulation and restore empty СТО sections."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.shared import Cm, Pt
    from docx.text.paragraph import Paragraph

    source_path = Path(source_path)
    dest_path = Path(dest_path)
    if not source_path.is_file():
        raise FileNotFoundError(source_path)
    dest_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source_path, dest_path)
    doc = Document(str(dest_path))
    _reset_title_page(doc)
    _reset_headers(doc)
    body = doc.element.body
    change_tbl = _find_change_log_table(body)
    start = _first_content_paragraph(doc)
    to_remove = []
    started = False
    for child in list(body):
        if child is start._p:
            started = True
        if not started:
            continue
        if change_tbl is not None and child is change_tbl:
            break
        if _has_sect_pr(child):
            continue
        to_remove.append(child)
    for el in to_remove:
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)
    insert_before = change_tbl if change_tbl is not None else body.find(_qn("sectPr"))
    if insert_before is None:
        insert_before = list(body)[-1]
    last_placeholder = None
    for idx, name in enumerate(STO_SECTIONS, start=1):
        heading = OxmlElement("w:p")
        insert_before.addprevious(heading)
        hp = Paragraph(heading, doc)
        _write_paragraph(
            hp,
            f"{idx} {name}",
            size=14,
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
            first_line=Cm(0),
        )
        placeholder = OxmlElement("w:p")
        heading.addnext(placeholder)
        text = _ARCHIVE_BOILERPLATE if idx == 8 else ("Приложений нет" if idx == 9 else "\t")
        _write_paragraph(
            Paragraph(placeholder, doc),
            text,
            size=14,
            align=WD_ALIGN_PARAGRAPH.JUSTIFY,
            first_line=Cm(0) if idx in {8, 9} else Cm(1.25),
        )
        last_placeholder = placeholder
        insert_before = last_placeholder.getnext() if last_placeholder.getnext() is not None else insert_before
    doc.save(str(dest_path))
    return dest_path


def _current_year() -> str:
    from datetime import date

    return str(date.today().year)


def _default_code(_title: str) -> str:
    return "РГ-"


def _section_heading(section: dict[str, Any]) -> str:
    title = str(section.get("title") or "").strip()
    number = str(section.get("number") or "").strip()
    return f"{number} {title}".strip() if title else ""


def _merge_section(target: dict[str, Any], source: dict[str, Any]) -> None:
    target["paragraphs"].extend(_clean_text_list(source.get("paragraphs")))
    target["items"].extend(_item_texts(source.get("items")))
    target["tables"].extend(_tables_of(source))
    for key in ("sections", "subsections", "children"):
        nested = source.get(key)
        if isinstance(nested, list):
            target["sections"].extend(item for item in nested if isinstance(item, dict))


def _clean_text_list(value: Any) -> list[str]:
    if not isinstance(value, list):
        return []
    return [str(item).strip() for item in value if str(item).strip()]


def _item_texts(value: Any) -> list[str]:
    out: list[str] = []
    if not isinstance(value, list):
        return out
    for item in value:
        if isinstance(item, dict):
            text = str(item.get("text") or item.get("title") or "").strip()
            if text:
                out.append(text)
            children = item.get("items") or item.get("children") or []
            if isinstance(children, list):
                out.extend(_item_texts(children))
        else:
            text = str(item or "").strip()
            if text:
                out.append(text)
    return out


def _tables_of(section: dict[str, Any]) -> list[dict[str, Any]]:
    tables = section.get("tables")
    if not isinstance(tables, list):
        return []
    out: list[dict[str, Any]] = []
    for table in tables:
        if not isinstance(table, dict):
            continue
        headers = [str(h).strip() for h in (table.get("headers") or [])]
        rows = table.get("rows") or []
        if not headers or not isinstance(rows, list):
            continue
        out.append({"headers": headers, "rows": rows})
    return out


def _set_run_font(run: Any, *, size: int = 14, bold: bool | None = None) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        r_fonts.set(qn(attr), "Times New Roman")


def _clear_paragraph(paragraph: Any) -> None:
    el = paragraph._p
    for child in list(el):
        if child.tag != _qn("pPr"):
            el.remove(child)


def _write_paragraph(
    paragraph: Any,
    text: str,
    *,
    size: int = 14,
    bold: bool = False,
    align: Any = None,
    first_line: Any = None,
    space_after: Any = None,
    space_before: Any = None,
) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.shared import Cm, Pt

    _clear_paragraph(paragraph)
    paragraph.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = paragraph.paragraph_format
    pf.space_after = Pt(6) if space_after is None else space_after
    pf.space_before = Pt(0) if space_before is None else space_before
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.first_line_indent = Cm(1.25) if first_line is None else first_line
    run = paragraph.add_run(text)
    _set_run_font(run, size=size, bold=bold)


def _insert_paragraph_after(paragraph: Any) -> Any:
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph

    new_el = OxmlElement("w:p")
    paragraph._p.addnext(new_el)
    return Paragraph(new_el, paragraph._parent)


def _insert_table_after(paragraph: Any, doc: Any, headers: list[str], rows: list[Any]) -> Any:
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    from docx.text.paragraph import Paragraph

    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    paragraph._p.addnext(table._tbl)
    for i, header in enumerate(headers):
        _set_cell_text(table.cell(0, i), header, size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _shade_cell(table.cell(0, i))
    for r_i, row in enumerate(rows, start=1):
        values = list(row) if not isinstance(row, str) else [row]
        for c_i, header in enumerate(headers):
            value = values[c_i] if c_i < len(values) else ""
            if isinstance(value, list):
                value = " / ".join(str(part) for part in value)
            _set_cell_text(table.cell(r_i, c_i), str(value))
    tbl_pr = table._tbl.find(qn("w:tblPr"))
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        table._tbl.insert(0, tbl_pr)
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "5000")
    tbl_w.set(qn("w:type"), "pct")
    spacer = OxmlElement("w:p")
    table._tbl.addnext(spacer)
    sp = Paragraph(spacer, paragraph._parent)
    _write_paragraph(sp, "", first_line=None, space_after=Pt(4))
    return sp


def _set_cell_text(cell: Any, text: str, *, size: int = 11, bold: bool = False, align: Any = None) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align or WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    _set_run_font(run, size=size, bold=bold)
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")


def _shade_cell(cell: Any, fill: str = "D9D9D9") -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def _set_header_runs(cell: Any, text: str, *, size: int = 10) -> None:
    paragraph = cell.paragraphs[0]
    _clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    _set_run_font(run, size=size, bold=True)


def _unique_row_cells(row: Any) -> list[Any]:
    seen: set[int] = set()
    out: list[Any] = []
    for cell in row.cells:
        ident = id(cell._tc)
        if ident in seen:
            continue
        seen.add(ident)
        out.append(cell)
    return out


def _fill_title_page(doc: Any, payload: dict[str, Any]) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    found_reglament = False
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == "РЕГЛАМЕНТ":
            found_reglament = True
            continue
        if found_reglament and not paragraph.text.strip():
            _write_paragraph(
                paragraph,
                payload["title"],
                size=16,
                bold=True,
                align=WD_ALIGN_PARAGRAPH.CENTER,
                first_line=Cm(0),
                space_before=Pt(12),
            )
            break
        if found_reglament and _sto_index(paragraph.text) is None and "Ростов" not in paragraph.text:
            if paragraph.alignment is not None and int(paragraph.alignment) == 1:
                _write_paragraph(
                    paragraph,
                    payload["title"],
                    size=16,
                    bold=True,
                    align=WD_ALIGN_PARAGRAPH.CENTER,
                    first_line=Cm(0),
                )
                break

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("20") or text == "2026":
            _write_paragraph(
                paragraph,
                payload["year"],
                size=14,
                bold=True,
                align=WD_ALIGN_PARAGRAPH.CENTER,
                first_line=Cm(0),
            )
            break
        if text == "Ростов-на-Дону" and payload.get("city"):
            _write_paragraph(
                paragraph,
                payload["city"],
                size=14,
                bold=True,
                align=WD_ALIGN_PARAGRAPH.CENTER,
                first_line=Cm(0),
            )

    if len(doc.tables) > 1:
        status = doc.tables[1]
        if "Статус" in status.cell(0, 0).text:
            _set_cell_text(status.cell(0, 1), payload.get("status") or "")

    if len(doc.tables) > 2:
        approval = doc.tables[2]
        if payload.get("developedBy"):
            _set_cell_text(approval.cell(1, 1), payload["developedBy"])
        if payload.get("agreedDept"):
            _set_cell_text(approval.cell(2, 1), payload["agreedDept"])
        _set_cell_text(approval.cell(3, 1), payload.get("agreedDirector") or "Директор по направлению")
        _set_cell_text(approval.cell(4, 1), payload.get("checkedBy") or "Специалист по процессному управлению")


def _fill_headers(doc: Any, payload: dict[str, Any]) -> None:
    code = payload.get("code") or "РГ-"
    short = payload["title"]
    version = payload.get("version") or "01"
    for section in doc.sections:
        header = section.header
        if header.tables:
            cells = _unique_row_cells(header.tables[0].rows[0])
            if len(cells) >= 4:
                _set_header_runs(cells[0], f"{code} {short}")
                _set_header_runs(cells[2], f"Версия            {version}")
            elif len(cells) >= 3:
                _set_header_runs(cells[0], f"{code} РЕГЛАМЕНТ")
                _set_header_runs(cells[1], f"Версия            {version}")
        if section.different_first_page_header_footer and section.first_page_header.tables:
            table = section.first_page_header.tables[0]
            if len(table.rows) >= 2:
                cells = _unique_row_cells(table.rows[1])
                if len(cells) >= 1:
                    _set_header_runs(cells[0], f"{code} РЕГЛАМЕНТ")
                if len(cells) >= 2:
                    _set_header_runs(cells[1], f"Версия                   {version}")


def _reset_title_page(doc: Any) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm

    after_reglament = False
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == "РЕГЛАМЕНТ":
            after_reglament = True
            continue
        if after_reglament and paragraph.text.strip() and "Ростов" not in paragraph.text and not paragraph.text.strip().startswith("20"):
            if _sto_index(paragraph.text) is not None:
                break
            _write_paragraph(paragraph, "", align=WD_ALIGN_PARAGRAPH.CENTER, first_line=Cm(0))
            break
    if len(doc.tables) > 1 and "Статус" in doc.tables[1].cell(0, 0).text:
        _set_cell_text(doc.tables[1].cell(0, 1), "")


def _reset_headers(doc: Any) -> None:
    for section in doc.sections:
        if not section.header.tables:
            continue
        cells = _unique_row_cells(section.header.tables[0].rows[0])
        if cells:
            _set_header_runs(cells[0], "РГ-          РЕГЛАМЕНТ")


def _fill_toc(doc: Any, payload: dict[str, Any]) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    toc_lines = [f"{idx} {name}" for idx, name in enumerate(STO_SECTIONS, start=1)]
    toc_lines.extend(["Лист регистрации изменений", "Лист выдачи и ознакомления"])
    toc_table = None
    for table in doc.tables:
        blob = table.cell(0, 0).text
        if "Назначение" in blob and "Нормативные" in blob:
            toc_table = table
            break
    if toc_table is None:
        return
    packed = [
        "\n\n".join(toc_lines[0:3]),
        toc_lines[3],
        "\n\n".join(toc_lines[4:7]),
        toc_lines[7],
        toc_lines[8],
        toc_lines[9],
        toc_lines[10],
    ]
    for row_i, text in enumerate(packed):
        if row_i >= len(toc_table.rows):
            break
        cell = toc_table.cell(row_i, 0)
        while len(cell.paragraphs) > 1:
            extra = cell.paragraphs[-1]._element
            extra.getparent().remove(extra)
        paragraph = cell.paragraphs[0]
        _clear_paragraph(paragraph)
        lines = text.split("\n")
        for i, line in enumerate(lines):
            target = paragraph if i == 0 else cell.add_paragraph()
            target.paragraph_format.space_after = Pt(0)
            target.paragraph_format.space_before = Pt(0)
            target.paragraph_format.first_line_indent = Cm(0)
            run = target.add_run(line)
            _set_run_font(run, size=14, bold=True)


def _is_toc_like(paragraph: Any) -> bool:
    style = (paragraph.style.name if paragraph.style else "") or ""
    if style.casefold().startswith("toc"):
        return True
    text = paragraph.text or ""
    if "\t" in text and re.search(r"\d+\s*$", text.strip()):
        return True
    return False


def _paragraph_in_table(paragraph: Any) -> bool:
    parent = paragraph._p.getparent()
    while parent is not None:
        if parent.tag == _qn("tbl"):
            return True
        parent = parent.getparent()
    return False


def _first_content_paragraph(doc: Any) -> Any:
    for paragraph in doc.paragraphs:
        if _paragraph_in_table(paragraph) or _is_toc_like(paragraph):
            continue
        if _sto_index(paragraph.text) == 1 or paragraph.text.strip() in {
            "Назначение и область применения",
            "1 Назначение и область применения",
        }:
            return paragraph
    raise KeyError("Не найден раздел «Назначение и область применения»")


def _find_change_log_table(body: Any) -> Any | None:
    for child in body:
        if child.tag != _qn("tbl"):
            continue
        texts = "".join((node.text or "") for node in child.iter(_qn("t")))
        if "Лист регистрации извещений" in texts:
            return child
    return None


def _fill_body(doc: Any, payload: dict[str, Any]) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.shared import Cm, Pt
    from docx.text.paragraph import Paragraph

    heading_map: dict[int, Any] = {}
    for paragraph in doc.paragraphs:
        if _paragraph_in_table(paragraph) or _is_toc_like(paragraph):
            continue
        idx = _sto_index(paragraph.text)
        if idx is not None and idx not in heading_map:
            heading_map[idx] = paragraph
            _write_paragraph(
                paragraph,
                f"{idx} {STO_SECTIONS[idx - 1]}",
                size=14,
                bold=True,
                align=WD_ALIGN_PARAGRAPH.CENTER,
                first_line=Cm(0),
                space_before=Pt(12),
                space_after=Pt(8),
            )
    for idx, section in enumerate(payload["sections"], start=1):
        heading = heading_map.get(idx)
        if heading is None:
            continue
        _clear_until_next_heading(heading, heading_map, idx)
        cur = heading
        nxt = heading._p.getnext()
        first_written = False
        if nxt is not None and nxt.tag == _qn("p"):
            placeholder = Paragraph(nxt, heading._parent)
            if _sto_index(placeholder.text) is None:
                cur = placeholder
                first_written = True
        cur = _write_section_content(doc, cur, section, first_into=cur if first_written else None)


def _clear_until_next_heading(heading: Any, heading_map: dict[int, Any], idx: int) -> None:
    from docx.oxml import OxmlElement

    stop = heading_map.get(idx + 1)
    stop_el = stop._p if stop is not None else _find_change_log_table(heading.part.document.element.body)
    child = heading._p.getnext()
    while child is not None and child is not stop_el:
        nxt = child.getnext()
        if _has_sect_pr(child):
            child = nxt
            continue
        if child.tag == _qn("p") and _sto_index("".join((t.text or "") for t in child.iter(_qn("t")))) == idx + 1:
            break
        parent = child.getparent()
        if parent is not None:
            parent.remove(child)
        child = nxt
    placeholder = OxmlElement("w:p")
    heading._p.addnext(placeholder)


def _write_section_content(doc: Any, cur: Any, section: dict[str, Any], *, first_into: Any | None) -> Any:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    blocks: list[tuple[str, Any]] = []
    for paragraph in section.get("paragraphs") or []:
        blocks.append(("p", str(paragraph)))
    for item in section.get("items") or []:
        text = item if isinstance(item, str) else str(item)
        if text and not text.startswith("—"):
            text = f"— {text}"
        blocks.append(("p", text))
    for table in section.get("tables") or []:
        blocks.append(("t", table))
    for nested in section.get("sections") or []:
        if not isinstance(nested, dict):
            continue
        heading = _section_heading(nested) or str(nested.get("title") or "").strip()
        if heading:
            blocks.append(("h", heading))
        for paragraph in _clean_text_list(nested.get("paragraphs")):
            blocks.append(("p", paragraph))
        for item in _item_texts(nested.get("items")):
            blocks.append(("p", item if item.startswith("—") else f"— {item}"))
        for table in _tables_of(nested):
            blocks.append(("t", table))

    target = first_into
    for kind, value in blocks:
        if kind == "p":
            if target is not None:
                _write_paragraph(target, value)
                cur = target
                target = None
            else:
                cur = _insert_paragraph_after(cur)
                _write_paragraph(cur, value)
        elif kind == "h":
            if target is not None:
                _write_paragraph(
                    target,
                    value,
                    bold=True,
                    align=WD_ALIGN_PARAGRAPH.LEFT,
                    first_line=Cm(0),
                    space_before=Pt(10),
                )
                cur = target
                target = None
            else:
                cur = _insert_paragraph_after(cur)
                _write_paragraph(
                    cur,
                    value,
                    bold=True,
                    align=WD_ALIGN_PARAGRAPH.LEFT,
                    first_line=Cm(0),
                    space_before=Pt(10),
                )
        elif kind == "t":
            if target is not None:
                cur = target
                target = None
            spacer = _insert_paragraph_after(cur)
            _write_paragraph(spacer, "", first_line=Cm(0), space_after=Pt(4))
            cur = _insert_table_after(spacer, doc, value.get("headers") or [], value.get("rows") or [])
    if target is not None and not blocks:
        _write_paragraph(target, "")
        cur = target
    return cur
