from __future__ import annotations

import json
import re
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any


_VAGUE_PATTERNS = (
    r"\bпри необходимости\b",
    r"\bпо возможности\b",
    r"\bсвоевременно\b",
    r"\bв кратчайшие сроки\b",
    r"\bоперативно\b",
)


@dataclass(slots=True)
class ValidationIssue:
    code: str
    severity: str
    message: str
    hint: str = ""


def _is_table_separator(line: str) -> bool:
    row = line.strip()
    if not row.startswith("|") or not row.endswith("|"):
        return False
    cells = [part.strip() for part in row.strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell or "") for cell in cells)


def _split_table_row(line: str) -> list[str]:
    return [part.strip() for part in line.strip().strip("|").split("|")]


def _normalize_heading(text: str) -> str:
    return re.sub(r"\s+", " ", text.strip().lower())


def _extract_headings(markdown_text: str) -> list[str]:
    headings: list[str] = []
    for raw in markdown_text.splitlines():
        match = re.match(r"^\s{0,3}#{1,6}\s+(.+?)\s*$", raw)
        if match:
            headings.append(_normalize_heading(match.group(1)))
    return headings


def _extract_required_titles(structure: dict[str, Any]) -> list[str]:
    titles: list[str] = []
    for sec in structure.get("sections") or []:
        if not isinstance(sec, dict):
            continue
        title = str(sec.get("title") or "").strip()
        if title:
            titles.append(_normalize_heading(title))
    return titles


def validate_regulation_markdown(
    markdown_text: str,
    structure: dict[str, Any] | None = None,
) -> dict[str, Any]:
    issues: list[ValidationIssue] = []
    text = markdown_text or ""
    headings = _extract_headings(text)
    heading_set = set(headings)

    if not headings:
        issues.append(
            ValidationIssue(
                code="no_headings",
                severity="error",
                message="В документе не найдены markdown-заголовки.",
                hint="Используйте разделы вида '# ...' и '## ...'.",
            )
        )

    # Required section coverage from structure JSON.
    required = _extract_required_titles(structure or {})
    missing_required = [title for title in required if title not in heading_set]
    for title in missing_required:
        issues.append(
            ValidationIssue(
                code="missing_section",
                severity="error",
                message=f"Отсутствует обязательный раздел: {title}",
                hint="Добавьте раздел в markdown в соответствии с шаблоном структуры.",
            )
        )

    # Vague wording.
    for pattern in _VAGUE_PATTERNS:
        if re.search(pattern, text, flags=re.IGNORECASE):
            issues.append(
                ValidationIssue(
                    code="vague_wording",
                    severity="warn",
                    message=f"Найдена расплывчатая формулировка: '{pattern}'.",
                    hint="Замените на измеримый критерий/срок/условие.",
                )
            )

    # Placeholders and TODO markers indicate incompleteness.
    tbd_hits = len(re.findall(r"<\s*TBD[^>]*>", text, flags=re.IGNORECASE))
    todo_hits = len(re.findall(r"\bTODO\b|\bFIXME\b", text, flags=re.IGNORECASE))
    if tbd_hits > 0:
        issues.append(
            ValidationIssue(
                code="tbd_placeholders",
                severity="warn",
                message=f"Найдены незаполненные placeholder-поля: {tbd_hits}.",
                hint="Закройте placeholders или явно назначьте владельца заполнения.",
            )
        )
    if todo_hits > 0:
        issues.append(
            ValidationIssue(
                code="todo_markers",
                severity="warn",
                message=f"Найдены маркеры TODO/FIXME: {todo_hits}.",
                hint="Замените TODO на финальные формулировки до публикации.",
            )
        )

    # Detect empty sections (heading with no content until next heading).
    lines = text.splitlines()
    for idx, line in enumerate(lines):
        if not re.match(r"^\s{0,3}#{1,6}\s+.+$", line):
            continue
        content_chars = 0
        j = idx + 1
        while j < len(lines) and not re.match(r"^\s{0,3}#{1,6}\s+.+$", lines[j]):
            content_chars += len(lines[j].strip())
            j += 1
        if content_chars < 20:
            heading = re.sub(r"^\s{0,3}#{1,6}\s+", "", line).strip()
            issues.append(
                ValidationIssue(
                    code="thin_section",
                    severity="warn",
                    message=f"Раздел '{heading}' выглядит неполным (мало содержимого).",
                    hint="Добавьте конкретику: действия, роли, критерии, сроки, артефакты.",
                )
            )

    # Table consistency checks.
    table_lines = [ln for ln in lines if ln.strip().startswith("|") and ln.strip().endswith("|")]
    if table_lines:
        i = 0
        while i < len(lines):
            ln = lines[i].strip()
            if not (ln.startswith("|") and ln.endswith("|")):
                i += 1
                continue
            block: list[str] = []
            while i < len(lines):
                cur = lines[i].strip()
                if cur.startswith("|") and cur.endswith("|"):
                    block.append(cur)
                    i += 1
                    continue
                break
            if len(block) >= 2 and _is_table_separator(block[1]):
                cols = len(_split_table_row(block[0]))
                for row in block[2:]:
                    if len(_split_table_row(row)) != cols:
                        issues.append(
                            ValidationIssue(
                                code="table_columns_mismatch",
                                severity="error",
                                message="В таблице найдено несовпадение количества колонок.",
                                hint="Проверьте строки таблицы: число ячеек должно быть одинаковым.",
                            )
                        )
                        break
            elif len(block) >= 1:
                issues.append(
                    ValidationIssue(
                        code="invalid_table",
                        severity="warn",
                        message="Найдены строки, похожие на таблицу, без корректного markdown-разделителя.",
                        hint="Добавьте вторую строку таблицы вида '|---|---|'.",
                    )
                )

    errors = [item for item in issues if item.severity == "error"]
    warnings = [item for item in issues if item.severity == "warn"]
    total_required = len(required)
    covered_required = max(0, total_required - len(missing_required))
    completeness_score = 100 if total_required == 0 else round((100 * covered_required) / total_required)

    return {
        "ok": len(errors) == 0,
        "stats": {
            "errors": len(errors),
            "warnings": len(warnings),
            "required_sections_total": total_required,
            "required_sections_covered": covered_required,
            "completeness_score": completeness_score,
            "tbd_count": tbd_hits,
            "todo_count": todo_hits,
        },
        "issues": [asdict(item) for item in issues],
    }


def markdown_to_docx(markdown_text: str, out_path: Path) -> None:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("Для формирования DOCX требуется python-docx") from exc

    doc = Document()
    lines = markdown_text.splitlines()
    i = 0
    in_code = False
    code_buffer: list[str] = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Code fence blocks.
        if stripped.startswith("```"):
            if in_code:
                if code_buffer:
                    block = doc.add_paragraph("\n".join(code_buffer))
                    block.style = "Intense Quote"
                    code_buffer = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buffer.append(line)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        # Headings.
        hm = re.match(r"^\s{0,3}(#{1,6})\s+(.+?)\s*$", line)
        if hm:
            level = min(len(hm.group(1)), 6)
            doc.add_heading(hm.group(2).strip(), level=level)
            i += 1
            continue

        # Table block.
        if stripped.startswith("|") and stripped.endswith("|") and i + 1 < len(lines):
            sep = lines[i + 1].strip()
            if _is_table_separator(sep):
                header = _split_table_row(lines[i])
                rows: list[list[str]] = []
                i += 2
                while i < len(lines):
                    cur = lines[i].strip()
                    if not (cur.startswith("|") and cur.endswith("|")):
                        break
                    rows.append(_split_table_row(lines[i]))
                    i += 1
                table = doc.add_table(rows=1, cols=len(header))
                table.style = "Table Grid"
                for idx, cell in enumerate(header):
                    table.rows[0].cells[idx].text = cell
                for row in rows:
                    tr = table.add_row().cells
                    for idx in range(len(header)):
                        tr[idx].text = row[idx] if idx < len(row) else ""
                continue

        # Bulleted list.
        bullet = re.match(r"^\s*[-*]\s+(.+)$", line)
        if bullet:
            doc.add_paragraph(bullet.group(1).strip(), style="List Bullet")
            i += 1
            continue

        # Numbered list.
        numbered = re.match(r"^\s*\d+\.\s+(.+)$", line)
        if numbered:
            doc.add_paragraph(numbered.group(1).strip(), style="List Number")
            i += 1
            continue

        # Paragraph (merge continuous lines until boundary).
        para = [stripped]
        j = i + 1
        while j < len(lines):
            nxt = lines[j].strip()
            if (
                not nxt
                or nxt.startswith("#")
                or nxt.startswith("|")
                or re.match(r"^\s*[-*]\s+.+$", lines[j])
                or re.match(r"^\s*\d+\.\s+.+$", lines[j])
                or nxt.startswith("```")
            ):
                break
            para.append(nxt)
            j += 1
        doc.add_paragraph(" ".join(para))
        i = j

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def load_structure_json(path: Path | None) -> dict[str, Any]:
    if not path:
        return {}
    raw = json.loads(path.read_text(encoding="utf-8"))
    return raw if isinstance(raw, dict) else {}
