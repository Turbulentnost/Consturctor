"""Сопоставление голосов расшифровки с сотрудниками и сохранение отпечатка.

Уже известный голос узнаётся по отпечатку. Неизвестный — только если агент
уверен по тексту и человек подтвердил. После подтверждения отпечаток пишется
в базу, и в следующий раз этот человек находится без угадывания.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any
from uuid import uuid4

from sqlalchemy import inspect
from sqlalchemy.orm import Session

from app.models.meeting_audio import MeetingTranscript, VoiceProfile
from app.models.user import AppUser

# Косинус нормализованных отпечатков ECAPA. Ниже порога голос считается новым.
# Если два сотрудника почти одинаково близки, угадывать нельзя — решает контекст.
MATCH_MIN = 0.62
MATCH_MARGIN = 0.05
_VOICE_RE = re.compile(r"Голос\s+(\d+)")


def fio_key(value: str) -> str:
    return " ".join((value or "").casefold().replace("ё", "е").split())


def cosine(left: list[float], right: list[float]) -> float:
    size = min(len(left), len(right))
    if size == 0:
        return 0.0
    dot = 0.0
    left_norm = 0.0
    right_norm = 0.0
    for index in range(size):
        a = float(left[index])
        b = float(right[index])
        dot += a * b
        left_norm += a * a
        right_norm += b * b
    if left_norm <= 0.0 or right_norm <= 0.0:
        return 0.0
    return dot / ((left_norm ** 0.5) * (right_norm ** 0.5))


def voice_labels(segments: list[dict[str, Any]]) -> list[str]:
    found: list[str] = []
    seen: set[str] = set()
    for row in segments:
        for match in _VOICE_RE.finditer(str(row.get("speaker") or "")):
            label = f"Голос {match.group(1)}"
            if label not in seen:
                seen.add(label)
                found.append(label)
    return found


def best_voice_match(
    vector: list[float],
    profiles: list[tuple[str, str, str, list[float]]],
) -> dict[str, Any] | None:
    """profiles: (id, user_id, fio, embedding). None, если нет явного победителя."""
    scored: list[tuple[float, str, str, str]] = []
    for profile_id, user_id, fio, embedding in profiles:
        if not embedding:
            continue
        scored.append((cosine(vector, embedding), profile_id, user_id, fio))
    scored.sort(key=lambda item: item[0], reverse=True)
    if not scored or scored[0][0] < MATCH_MIN:
        return None
    if len(scored) > 1 and scored[0][0] - scored[1][0] < MATCH_MARGIN:
        return None
    score, profile_id, user_id, fio = scored[0]
    return {
        "profile_id": profile_id,
        "user_id": user_id,
        "fio": fio,
        "similarity": round(score, 3),
    }


def fio_initials(fio: str) -> str:
    """«Амураль Игорь Борисович» → «Амураль И.Б.»; одно слово остаётся как есть."""
    parts = (fio or "").split()
    if len(parts) < 2:
        return " ".join(parts)
    return parts[0] + " " + "".join(f"{part[0].upper()}." for part in parts[1:3])


def relabel_speaker(speaker: str, mapping: dict[str, str]) -> str:
    def replace(match: re.Match[str]) -> str:
        label = f"Голос {match.group(1)}"
        return mapping.get(label, label)

    return _VOICE_RE.sub(replace, speaker or "")


def _table_ready(db: Session) -> bool:
    try:
        bind = db.get_bind()
        return inspect(bind).has_table(MeetingTranscript.__tablename__)
    except Exception:
        return False


def _clean_embeddings(raw: dict[str, Any] | None) -> dict[str, list[float]]:
    cleaned: dict[str, list[float]] = {}
    for key, value in (raw or {}).items():
        if not isinstance(value, list) or not value:
            continue
        try:
            cleaned[str(key)] = [float(item) for item in value]
        except (TypeError, ValueError):
            continue
    return cleaned


def _known_and_unknown(
    db: Session, labels: list[str], embeddings: dict[str, list[float]]
) -> tuple[list[dict[str, Any]], list[str]]:
    profiles = [
        (row.id, row.user_id, row.fio, list(row.embedding or []))
        for row in db.query(VoiceProfile).all()
    ]
    known: list[dict[str, Any]] = []
    unknown: list[str] = []
    for label in labels:
        vector = embeddings.get(label) or []
        match = best_voice_match(vector, profiles) if vector else None
        if match is None:
            unknown.append(label)
            continue
        known.append({"speaker": label, **match})
    return known, unknown


def save_transcript(
    db: Session,
    *,
    user_id: str,
    filename: str,
    sha256: str,
    duration_sec: float,
    transcript_path: str,
    full_text: str,
    segments: list[dict[str, Any]],
    embeddings: dict[str, Any] | None,
) -> dict[str, Any] | None:
    if not _table_ready(db):
        return None
    vectors = _clean_embeddings(embeddings)
    labels = voice_labels(segments)
    row = MeetingTranscript(
        id=str(uuid4()),
        user_id=user_id,
        filename=filename,
        sha256=sha256,
        duration_sec=float(duration_sec or 0.0),
        transcript_path=transcript_path,
        full_text=full_text,
        segments=segments,
        embeddings=vectors,
        speaker_map={},
    )
    db.add(row)
    db.commit()
    db.refresh(row)
    known, unknown = _known_and_unknown(db, labels, vectors)
    return {
        "transcript_id": row.id,
        "known_speakers": known,
        "unknown_speakers": unknown,
    }


_UNBOUND_SPEAKER_RE = re.compile(r"Голос\s+\d+|одновременно|[\s+()]")


def _pristine(row: MeetingTranscript) -> bool:
    """В репликах только «Голос N»: bind_speakers эту запись ещё не переписывал."""
    return all(
        not _UNBOUND_SPEAKER_RE.sub("", str(segment.get("speaker") or ""))
        for segment in row.segments or []
        if isinstance(segment, dict)
    )


def claim_transcript(transcript_id: str, user_id: str, sha256: str) -> dict[str, Any] | None:
    """Расшифровка из кэша — запись этого пользователя с метками «Голос N».

    Кэш общий для всех, а audio.speakers и bind_speakers открывают только свою запись.
    Своей нетронутой записи нет — заводим копию с чистой чужой (или с той, что в кэше).
    """
    from app.db.session import SessionLocal

    db = SessionLocal()
    try:
        if not _table_ready(db):
            return None
        rows = (
            db.query(MeetingTranscript)
            .filter(MeetingTranscript.sha256 == sha256)
            .order_by(MeetingTranscript.created_at.desc())
            .all()
        )
        pristine = [row for row in rows if _pristine(row)]
        mine = next((row for row in pristine if row.user_id == user_id), None)
        if mine is None:
            source = pristine[0] if pristine else db.get(MeetingTranscript, (transcript_id or "").strip())
            if source is None:
                return None
            mine = MeetingTranscript(
                id=str(uuid4()),
                user_id=user_id,
                filename=source.filename,
                sha256=source.sha256,
                duration_sec=source.duration_sec,
                transcript_path=source.transcript_path,
                full_text=source.full_text,
                segments=list(source.segments or []),
                embeddings=dict(source.embeddings or {}),
                speaker_map={},
            )
            db.add(mine)
            db.commit()
            db.refresh(mine)
        segments = list(mine.segments or [])
        vectors = _clean_embeddings(mine.embeddings if isinstance(mine.embeddings, dict) else {})
        known, unknown = _known_and_unknown(db, voice_labels(segments), vectors)
        return {
            "transcript_id": mine.id,
            "segments": segments,
            "pristine": _pristine(mine),
            "known_speakers": known,
            "unknown_speakers": unknown,
        }
    finally:
        db.close()


def _bound_path(path: str, transcript_id: str) -> Path:
    """Файл с ФИО — свой у каждой записи: исходный остаётся в кэше с метками «Голос N»."""
    source = Path(path)
    tag = transcript_id[:8]
    if source.stem.endswith(f"-{tag}"):
        return source
    return source.with_name(f"{source.stem}-{tag}{source.suffix}")


def _own_transcript(db: Session, transcript_id: str, user_id: str) -> MeetingTranscript:
    row = db.get(MeetingTranscript, (transcript_id or "").strip())
    if row is None or row.user_id != user_id:
        raise RuntimeError("Расшифровка не найдена")
    return row


def _lines_for(segments: list[dict[str, Any]], label: str, limit: int) -> list[str]:
    lines: list[str] = []
    for row in segments:
        speaker = str(row.get("speaker") or "")
        if label not in speaker:
            continue
        text = str(row.get("text") or "").strip()
        if not text:
            continue
        start = float(row.get("start") or 0.0)
        minutes, seconds = divmod(int(start), 60)
        hours, minutes = divmod(minutes, 60)
        clock = f"{hours:02d}:{minutes:02d}:{seconds:02d}" if hours else f"{minutes:02d}:{seconds:02d}"
        lines.append(f"[{clock}] {text}")
        if len(lines) >= limit:
            break
    return lines


def _stored_proposals(row: MeetingTranscript) -> list[dict[str, Any]]:
    raw = row.speaker_map if isinstance(row.speaker_map, dict) else {}
    items = raw.get("proposals")
    if not isinstance(items, list):
        return []
    return [item for item in items if isinstance(item, dict)]


def propose_speakers(
    db: Session,
    *,
    transcript_id: str,
    user_id: str,
    voices: list[dict[str, Any]],
) -> dict[str, Any]:
    """Сохранить предложенные ФИО. В 1С и в базу голосов ещё не пишет — это делает автор на карточке."""
    row = _own_transcript(db, transcript_id, user_id)
    labels = voice_labels(list(row.segments or []))
    incoming: dict[str, dict[str, Any]] = {}
    for item in voices:
        speaker = str(item.get("speaker") or "").strip()
        if speaker:
            incoming[speaker] = item
    missing = [label for label in labels if label not in incoming]
    if missing:
        raise RuntimeError("Предложи человека для каждого голоса: " + ", ".join(missing))
    proposals = []
    for label in labels:
        item = incoming[label]
        fio = " ".join(str(item.get("fio") or "").split())
        proposals.append(
            {
                "speaker": label,
                "fio": fio,
                "user_id": str(item.get("user_id") or "").strip(),
                "source": str(item.get("source") or "").strip() or ("voice" if fio else ""),
            }
        )
    row.speaker_map = {"proposals": proposals}
    db.commit()
    return speaker_context(db, transcript_id, user_id)


def transcript_docx_bytes(title: str, segments: list[dict[str, Any]]) -> bytes:
    import io

    from docx import Document

    document = Document()
    document.add_heading(title or "Расшифровка совещания", level=1)
    for row in segments:
        speaker = str(row.get("speaker") or "").strip()
        text = str(row.get("text") or "").strip()
        if not text:
            continue
        document.add_paragraph(f"{speaker}: {text}" if speaker else text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def export_transcript_docx(transcript_id: str, user_id: str, workflow_id: str) -> dict[str, Any]:
    from app.db.session import SessionLocal
    from app.models.workflow import Workflow
    from app.services.workflow_files import register_agent_files

    db = SessionLocal()
    try:
        row = _own_transcript(db, transcript_id, user_id)
        workflow = db.get(Workflow, (workflow_id or "").strip())
        if workflow is None or workflow.user_id != user_id:
            raise RuntimeError("Нет агента, к которому прикрепить файл расшифровки")
        title = Path(row.filename).stem if row.filename else "Расшифровка"
        raw = transcript_docx_bytes(f"Расшифровка {title}", list(row.segments or []))
        filename = f"Расшифровка-{title}.docx"
        saved = register_agent_files(
            db,
            row=workflow,
            files=[(filename, raw)],
            run_id="",
            origin="transcript",
        )
        db.commit()
        item = saved[0]
        return {
            "ok": True,
            "file_id": item.id,
            "filename": item.filename,
            "workflow_id": workflow.id,
        }
    finally:
        db.close()


def speakers_for_user(transcript_id: str, user_id: str) -> dict[str, Any]:
    from app.db.session import SessionLocal

    db = SessionLocal()
    try:
        return speaker_context(db, transcript_id, user_id)
    finally:
        db.close()


def confirm_speakers(
    transcript_id: str,
    user_id: str,
    voices: list[dict[str, Any]],
    *,
    save_voices: bool = True,
    initials: bool = False,
) -> dict[str, Any]:
    from app.db.session import SessionLocal

    db = SessionLocal()
    try:
        return bind_speakers(
            db,
            transcript_id=transcript_id,
            user_id=user_id,
            voices=voices,
            save_voices=save_voices,
            initials=initials,
        )
    finally:
        db.close()


def speaker_context(db: Session, transcript_id: str, user_id: str) -> dict[str, Any]:
    row = _own_transcript(db, transcript_id, user_id)
    segments = list(row.segments or [])
    embeddings = _clean_embeddings(row.embeddings if isinstance(row.embeddings, dict) else {})
    labels = voice_labels(segments)
    known, unknown = _known_and_unknown(db, labels, embeddings)
    known_by = {item["speaker"]: item for item in known}
    proposed = {str(item.get("speaker") or ""): item for item in _stored_proposals(row)}
    voices = []
    for label in labels:
        match = known_by.get(label)
        suggestion = proposed.get(label) or {}
        fio = str(suggestion.get("fio") or (match or {}).get("fio") or "").strip()
        lines = _lines_for(segments, label, 12)
        voices.append(
            {
                "speaker": label,
                "known": bool(match),
                "fio": fio,
                "user_id": str(suggestion.get("user_id") or (match or {}).get("user_id") or "").strip(),
                "source": str(suggestion.get("source") or ("voice" if match else "")).strip(),
                "similarity": (match or {}).get("similarity"),
                "said": " ".join(lines)[:700],
                "lines": lines,
            }
        )
    return {
        "ok": True,
        "transcript_id": row.id,
        "transcript_path": row.transcript_path,
        "filename": row.filename,
        "duration_sec": row.duration_sec,
        "voices": voices,
        "unknown_speakers": unknown,
        "note": (
            "known=true — этот голос уже есть в базе, по тексту его не переназначай. "
            "Для остальных предлагай сотрудника только если из реплик прямо ясно, кто это, "
            "и спроси подтверждение с ключевой фразой."
        ),
    }


def _directory_fio(db: Session, fio: str, user_id: str) -> tuple[str, str, bool]:
    wanted_id = (user_id or "").strip()
    if wanted_id:
        user = db.get(AppUser, wanted_id)
        if user is not None and (user.fio or "").strip():
            return user.id, user.fio.strip(), True
    key = fio_key(fio)
    if not key:
        raise RuntimeError("У голоса нет ФИО")
    users = db.query(AppUser).all()
    exact = [user for user in users if fio_key(user.fio) == key]
    if len(exact) == 1:
        return exact[0].id, exact[0].fio.strip(), True
    if len(exact) > 1:
        raise RuntimeError(f"ФИО «{fio.strip()}» есть у нескольких сотрудников. Передай user_id.")
    return "", " ".join(fio.split()), False


def bind_speakers(
    db: Session,
    *,
    transcript_id: str,
    user_id: str,
    voices: list[dict[str, Any]],
    save_voices: bool = True,
    initials: bool = False,
) -> dict[str, Any]:
    """Подставить ФИО вместо «Голос N». save_voices=False — догадка агента без проверки человеком:
    отпечатки голосов в банк не пишутся. initials=True — в репликах «Фамилия И.О.»."""
    row = _own_transcript(db, transcript_id, user_id)
    segments = list(row.segments or [])
    labels = voice_labels(segments)
    if not labels:
        raise RuntimeError("В расшифровке нет голосов")
    incoming: dict[str, dict[str, Any]] = {}
    for item in voices:
        if not isinstance(item, dict):
            continue
        speaker = str(item.get("speaker") or "").strip()
        if speaker:
            incoming[speaker] = item
    mapping: dict[str, str] = {}
    people: list[dict[str, Any]] = []
    embeddings = _clean_embeddings(row.embeddings if isinstance(row.embeddings, dict) else {})
    for label in labels:
        item = incoming.get(label) or {}
        fio_raw = str(item.get("fio") or "").strip()
        if not fio_raw:
            continue
        employee_id, fio, in_directory = _directory_fio(
            db, fio_raw, str(item.get("user_id") or "")
        )
        mapping[label] = fio
        vector = embeddings.get(label) or []
        voice_saved = False
        if vector and save_voices:
            key = fio_key(fio)
            profile = db.query(VoiceProfile).filter(VoiceProfile.fio_key == key).first()
            if profile is None:
                db.add(
                    VoiceProfile(
                        id=str(uuid4()),
                        user_id=employee_id,
                        fio=fio,
                        fio_key=key,
                        embedding=vector,
                    )
                )
                voice_saved = True
            elif employee_id and not profile.user_id:
                profile.user_id = employee_id
        preview = " ".join(_lines_for(segments, label, 3))
        people.append(
            {
                "speaker": label,
                "fio": fio,
                "user_id": employee_id,
                "in_directory": in_directory,
                "voice_saved": voice_saved,
                "said_preview": preview[:500],
            }
        )
    if initials:
        short = {label: fio_initials(fio) for label, fio in mapping.items()}
        taken = list(short.values())
        mapping = {label: name if taken.count(name) == 1 else mapping[label] for label, name in short.items()}
    for person in people:
        person["label"] = mapping.get(person["speaker"], person["fio"])
    rewritten = []
    for segment in segments:
        item = dict(segment)
        item["speaker"] = relabel_speaker(str(item.get("speaker") or ""), mapping)
        rewritten.append(item)
    full = relabel_speaker(row.full_text or "", mapping)
    path = row.transcript_path
    if path:
        try:
            file_path = Path(path)
            if file_path.is_file():
                text = file_path.read_text(encoding="utf-8")
                target = _bound_path(path, row.id)
                target.write_text(relabel_speaker(text, mapping), encoding="utf-8")
                path = str(target)
        except OSError:
            path = ""
    row.segments = rewritten
    row.full_text = full
    row.speaker_map = mapping
    row.transcript_path = path
    db.commit()
    return {
        "ok": True,
        "transcript_id": row.id,
        "transcript_path": path,
        "people": people,
        "note": (
            "Дальше в 1С идёт эта расшифровка: в репликах уже ФИО. "
            "audio.transcribe повторно не вызывай."
        ),
    }
