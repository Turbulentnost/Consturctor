from __future__ import annotations

import base64
import io
import logging
import mimetypes
import re
import tempfile
from pathlib import Path

logger = logging.getLogger(__name__)

TEXT_SUFFIXES = {
    ".txt",
    ".md",
    ".markdown",
    ".rst",
    ".csv",
    ".tsv",
    ".json",
    ".jsonl",
    ".xml",
    ".html",
    ".htm",
    ".yaml",
    ".yml",
    ".log",
    ".ini",
    ".cfg",
    ".conf",
    ".rtf",
}
DOC_SUFFIXES = {".pdf", ".docx", ".doc"}
SPREADSHEET_SUFFIXES = {".xlsx", ".xlsm", ".xls"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff"}
# Медиа хранится как есть: текст из него не извлекается (расшифровку делает
# audio.transcribe). Без этого байты декодируются как latin-1 и NUL (0x00)
# роняет вставку в текстовые поля PostgreSQL.
MEDIA_SUFFIXES = {
    ".wav",
    ".mp3",
    ".m4a",
    ".aac",
    ".ogg",
    ".opus",
    ".flac",
    ".wma",
    ".amr",
    ".webm",
    ".mp4",
    ".mkv",
    ".mov",
}
SUPPORTED_SUFFIXES = (
    TEXT_SUFFIXES | DOC_SUFFIXES | SPREADSHEET_SUFFIXES | IMAGE_SUFFIXES | MEDIA_SUFFIXES
)
MAX_IMAGES = 5
MAX_IMAGE_BYTES = 15 * 1024 * 1024
MAX_FILE_BYTES = 25 * 1024 * 1024

_MIME_BY_SUFFIX = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".webp": "image/webp",
    ".bmp": "image/bmp",
    ".tif": "image/tiff",
    ".tiff": "image/tiff",
}


class DocumentError(Exception):
    pass


def load_attachment_bytes(name: str, raw: bytes, *, ocr: bool = True) -> dict:
    suffix = Path(name).suffix.lower()
    file_name = Path(name).name or "file"
    if not raw:
        raise DocumentError("Файл пустой.")
    if suffix in MEDIA_SUFFIXES:
        return {
            "name": file_name,
            "text": f"Прикреплён медиафайл {file_name} ({len(raw)} байт).",
            "kind": "audio",
            "mime_type": mimetypes.guess_type(name)[0] or "application/octet-stream",
            "data_b64": "",
        }
    if suffix not in SUPPORTED_SUFFIXES:
        text = _read_text_bytes(raw)
        if text.strip():
            return {
                "name": file_name,
                "text": text.strip(),
                "kind": "text",
                "mime_type": _guess_text_mime(suffix or ".txt"),
                "data_b64": "",
            }
        return {
            "name": file_name,
            "text": f"Прикреплён файл {file_name} ({len(raw)} байт).",
            "kind": "binary",
            "mime_type": mimetypes.guess_type(name)[0] or "application/octet-stream",
            "data_b64": base64.b64encode(raw).decode("ascii"),
        }
    if suffix in IMAGE_SUFFIXES:
        return _load_image(name, raw, suffix, ocr=ocr)
    if len(raw) > MAX_FILE_BYTES:
        raise DocumentError(
            f"{file_name}: слишком большой файл ({len(raw) // (1024 * 1024)} МБ). "
            f"Лимит — {MAX_FILE_BYTES // (1024 * 1024)} МБ."
        )
    if suffix == ".pdf":
        try:
            text = _read_pdf_bytes(raw, ocr=ocr)
        except DocumentError:
            text = ""
        if not text.strip() and ocr:
            text = _ocr_visual(file_name, raw)
        kind = "text"
        mime = "application/pdf"
    elif suffix == ".docx":
        text = _read_docx_bytes(raw)
        kind = "text"
        mime = _guess_text_mime(suffix)
    elif suffix in SPREADSHEET_SUFFIXES:
        try:
            text = _read_xlsx_bytes(raw)
        except DocumentError:
            if suffix != ".xls":
                raise
            text = ""
        kind = "text"
        mime = _guess_text_mime(suffix)
    elif suffix in TEXT_SUFFIXES:
        text = _read_text_bytes(raw)
        kind = "text"
        mime = _guess_text_mime(suffix)
    else:
        text = ""
        kind = "binary"
        mime = mimetypes.guess_type(file_name)[0] or "application/octet-stream"
    text = (text or "").strip()
    if not text:
        text = f"[файл {file_name}: текст не извлечён, исходный файл сохранён]"
    return {
        "name": file_name,
        "text": text,
        "kind": kind,
        "mime_type": mime,
        "data_b64": "",
    }


def compose_document(attachments: list[dict], notes: str = "") -> tuple[str, str]:
    parts: list[str] = []
    names: list[str] = []
    image_n = 0
    for att in attachments:
        name = str(att.get("name") or "file")
        names.append(name)
        if att.get("kind") == "image":
            image_n += 1
            parts.append(
                f"===== IMAGE: {name} =====\n"
                f"(изображение #{image_n}, mime={att.get('mime_type') or 'image'}; "
                f"передано агенту как vision-вложение)\n"
                f"===== END IMAGE ====="
            )
            continue
        body = str(att.get("text") or "").strip()
        if not body:
            continue
        parts.append(f"===== FILE: {name} =====\n{body}\n===== END FILE =====")
    notes_text = (notes or "").strip()
    if notes_text:
        parts.append(f"===== NOTES =====\n{notes_text}\n===== END NOTES =====")
    if not parts and not names:
        return "", ""
    if names:
        real = [name for name in names if name.casefold() not in {"notes.txt", "notes", "file"}]
        if real:
            document_name = real[0] if len(real) == 1 else f"{len(real)} файлов"
        else:
            document_name = "материалы"
    else:
        document_name = "материалы"
    return document_name, "\n\n".join(parts)


def collect_prompt_images(attachments: list[dict]) -> list[dict[str, str]]:
    images: list[dict[str, str]] = []
    for att in attachments:
        if att.get("kind") != "image":
            continue
        data = str(att.get("data_b64") or "").strip()
        mime = str(att.get("mime_type") or "").strip()
        if not data:
            # Fallback: reload bytes from stored path (e.g. after payload slim/migrate).
            path = str(att.get("path") or "").strip()
            if path and Path(path).is_file():
                raw = Path(path).read_bytes()
                if raw and len(raw) <= MAX_IMAGE_BYTES:
                    data = base64.b64encode(raw).decode("ascii")
                    if not mime:
                        suffix = Path(path).suffix.lower()
                        mime = _MIME_BY_SUFFIX.get(suffix) or "image/png"
        if not data or not mime:
            continue
        images.append({"data": data, "mimeType": mime})
        if len(images) >= MAX_IMAGES:
            break
    return images


def _ocr_visual(name: str, raw: bytes) -> str:
    try:
        from app.services.ocr_extract import extract_visual_text

        return (extract_visual_text(name, raw) or "").strip()
    except Exception as exc:  # noqa: BLE001
        logger.warning(
            "attachment ocr failed name=%s detail=%s",
            ascii(name),
            ascii(str(exc)),
        )
        return ""


def _load_image(name: str, raw: bytes, suffix: str, *, ocr: bool = True) -> dict:
    mime = _MIME_BY_SUFFIX.get(suffix) or mimetypes.guess_type(name)[0] or "image/png"
    file_name = Path(name).name
    if not raw:
        raise DocumentError("Изображение пустое.")
    if len(raw) > MAX_IMAGE_BYTES:
        raise DocumentError(
            f"{file_name}: слишком большой файл ({len(raw) // (1024 * 1024)} МБ). "
            f"Лимит API — {MAX_IMAGE_BYTES // (1024 * 1024)} МБ."
        )
    text = _ocr_visual(file_name, raw) if ocr else ""
    if not text:
        text = f"[изображение: {file_name}]"
    return {
        "name": file_name,
        "text": text,
        "kind": "image",
        "mime_type": mime,
        "data_b64": base64.b64encode(raw).decode("ascii"),
    }


def _guess_text_mime(suffix: str) -> str:
    return {
        ".json": "application/json",
        ".xml": "application/xml",
        ".html": "text/html",
        ".htm": "text/html",
        ".csv": "text/csv",
        ".md": "text/markdown",
        ".markdown": "text/markdown",
        ".yaml": "text/yaml",
        ".yml": "text/yaml",
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ".xlsm": "application/vnd.ms-excel.sheet.macroEnabled.12",
    }.get(suffix, "text/plain")


def _read_text_bytes(raw: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "cp1251", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentError("Не удалось прочитать текстовый файл (кодировка).")


def _read_pdf_bytes(raw: bytes, *, ocr: bool = True) -> str:
    try:
        import fitz  # pymupdf
    except ImportError as exc:
        raise DocumentError("Для PDF нужен pymupdf") from exc
    try:
        doc = fitz.open(stream=raw, filetype="pdf")
        parts = [page.get_text() or "" for page in doc]
        doc.close()
    except Exception as exc:  # noqa: BLE001
        raise DocumentError(f"Не удалось разобрать PDF: {exc}") from exc
    text = "\n\n".join(parts).strip()
    if len(text) >= 80 or not ocr:
        return text
    ocr_text = _read_pdf_bytes_ocr(raw)
    if ocr_text.strip():
        return ocr_text
    return text


def _read_pdf_bytes_ocr(raw: bytes) -> str:
    """OCR fallback for scan PDFs when text layer is empty."""
    try:
        from app.services.regulation.detect import is_scan_pdf
        from app.services.regulation.pdf_ocr import extract_pdf_scan
    except ImportError:
        return ""
    if not raw:
        return ""
    with tempfile.TemporaryDirectory(prefix="wf-pdf-ocr-") as tmp:
        path = Path(tmp) / "attachment.pdf"
        path.write_bytes(raw)
        try:
            is_scan, _pages = is_scan_pdf(path)
        except Exception:
            is_scan = len(_read_pdf_text_layer(raw)) < 80
        if not is_scan and _read_pdf_text_layer(raw).strip():
            return ""
        try:
            extracted = extract_pdf_scan(path, work_dir=Path(tmp))
        except Exception as exc:  # noqa: BLE001
            logger.warning("pdf ocr fallback failed: %s", exc)
            return ""
        parts = [block.text.strip() for block in extracted.blocks if getattr(block, "text", "")]
        return "\n\n".join(part for part in parts if part)


def _read_pdf_text_layer(raw: bytes) -> str:
    try:
        import fitz

        doc = fitz.open(stream=raw, filetype="pdf")
        parts = [page.get_text() or "" for page in doc]
        doc.close()
        return "\n\n".join(parts).strip()
    except Exception:
        return ""


def _read_docx_bytes(raw: bytes) -> str:
    try:
        import docx
    except ImportError as exc:
        raise DocumentError("Для DOCX нужен python-docx") from exc
    try:
        document = docx.Document(io.BytesIO(raw))
        parts = [para.text for para in document.paragraphs if para.text]
        for table in document.tables:
            for row in table.rows:
                cells = [" ".join(cell.text.split()) for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
    except Exception as exc:  # noqa: BLE001
        raise DocumentError(f"Не удалось разобрать DOCX: {exc}") from exc
    return "\n".join(parts)


def _read_xlsx_bytes(raw: bytes) -> str:
    try:
        import openpyxl
    except ImportError as exc:
        raise DocumentError("Для XLSX нужен openpyxl") from exc
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        parts: list[str] = []
        for sheet in workbook.worksheets:
            parts.append(f"===== SHEET: {sheet.title} =====")
            for row in sheet.iter_rows(values_only=True):
                values = [str(cell).strip() if cell is not None else "" for cell in row]
                while values and not values[-1]:
                    values.pop()
                if values:
                    parts.append("\t".join(values))
        workbook.close()
    except Exception as exc:  # noqa: BLE001
        raise DocumentError(f"Не удалось разобрать XLSX: {exc}") from exc
    return "\n".join(parts)


_HEADER_HINT = re.compile(
    r"^(id|код|источник|дата|поручен|заказчик|владелец|срок|приоритет|статус|"
    r"комментар|риск|ссылка|исполнитель|тема|название|результат|артефакт|"
    r"документ|ответств|описание|номер)",
    re.I,
)
_KPI_HINT = re.compile(r"карточек|строк|срок|всего|открыт|просроч|сегодня|закрыт|выполн|итог|kpi", re.I)
_NUMERIC_RE = re.compile(r"^[+-]?(?:\d{1,3}(?:[\s\u00a0]\d{3})*|\d+)(?:[.,]\d+)?%?$")
_MAX_PREVIEW_ROWS = 400


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).replace("\xa0", " ").strip()


def _filled(cells: list[str]) -> list[str]:
    return [cell for cell in cells if cell]


def _looks_like_header(cells: list[str]) -> bool:
    filled = _filled(cells)
    if len(filled) < 2:
        return False
    hits = sum(1 for cell in filled if _HEADER_HINT.search(cell))
    if hits >= 2:
        return True
    return len(filled) >= 4 and all(
        len(cell) <= 48 and not _NUMERIC_RE.match(cell) and not cell.upper().startswith("ACT")
        for cell in filled
    )


def _looks_like_kpi_labels(cells: list[str]) -> bool:
    filled = _filled(cells)
    if len(filled) < 2 or len(filled) > 8:
        return False
    if any(len(cell) > 42 or cell.upper().startswith("ACT") for cell in filled):
        return False
    return any(_KPI_HINT.search(cell) for cell in filled)


def _looks_like_kpi_values(cells: list[str]) -> bool:
    filled = _filled(cells)
    if not filled:
        return False
    numeric = sum(1 for cell in filled if _NUMERIC_RE.match(cell))
    return numeric >= max(1, -(-len(filled) * 3 // 5))


def _looks_like_banner(cells: list[str]) -> bool:
    filled = _filled(cells)
    if not filled or len(filled) > 4:
        return False
    if _looks_like_header(cells) or _looks_like_kpi_labels(cells) or _looks_like_kpi_values(cells):
        return False
    text = " ".join(filled)
    return len(text) >= 6 and not _NUMERIC_RE.match(text)


def _structure_sheet(name: str, rows: list[list[str]]) -> dict[str, object]:
    index = 0
    title = ""
    subtitle = ""
    kpis: list[dict[str, str]] = []
    notes: list[str] = []
    if index < len(rows) and _looks_like_banner(rows[index]):
        title = " ".join(_filled(rows[index]))
        index += 1
    if index < len(rows) and _looks_like_banner(rows[index]):
        subtitle = " ".join(_filled(rows[index]))
        index += 1
    if (
        index + 1 < len(rows)
        and _looks_like_kpi_labels(rows[index])
        and _looks_like_kpi_values(rows[index + 1])
    ):
        labels = _filled(rows[index])
        values = _filled(rows[index + 1])
        for label, value in zip(labels, values):
            kpis.append({"label": label, "value": value})
        index += 2
    header_index = next((i for i in range(index, len(rows)) if _looks_like_header(rows[i])), -1)
    headers: list[str] = []
    data: list[list[str]] = []
    if header_index >= 0:
        for row in rows[index:header_index]:
            note = " — ".join(_filled(row))
            if note:
                notes.append(note)
        headers = rows[header_index]
        data = [row for row in rows[header_index + 1 :] if _filled(row)][:_MAX_PREVIEW_ROWS]
        width = max([len(headers), *(len(row) for row in data)], default=0)
        headers = (headers + [""] * width)[:width]
        headers = [cell or f"Колонка {i + 1}" for i, cell in enumerate(headers)]
        data = [(row + [""] * width)[:width] for row in data]
    else:
        for row in rows[index:]:
            note = " — ".join(_filled(row))
            if note:
                notes.append(note)
    return {
        "name": name,
        "title": title,
        "subtitle": subtitle,
        "kpis": kpis,
        "notes": notes,
        "headers": headers,
        "rows": data,
    }


def extract_xlsx_preview(raw: bytes) -> dict[str, object] | None:
    if not raw:
        return None
    try:
        import openpyxl
    except ImportError:
        return None
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    except Exception:
        return None
    try:
        sheets: list[dict[str, object]] = []
        for sheet in workbook.worksheets:
            rows: list[list[str]] = []
            for row in sheet.iter_rows(values_only=True):
                values = [_cell_text(cell) for cell in row]
                while values and not values[-1]:
                    values.pop()
                if values:
                    rows.append(values)
            if not rows:
                continue
            structured = _structure_sheet(str(sheet.title or "Лист"), rows)
            if structured["headers"] or structured["kpis"] or structured["title"]:
                sheets.append(structured)
        if not sheets:
            return None
        return {"kind": "workbook", "sheets": sheets}
    finally:
        workbook.close()
