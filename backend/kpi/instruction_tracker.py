"""Реестр и контроль поручений СД+РК.

  py -m kpi.instruction_tracker --month 2026-09 --tracker ActionTracker.xlsx --odata
  py -m kpi.instruction_tracker --month 2026-09 --tracker rows.json --cards cards.json --protocols protocols.json
"""

from __future__ import annotations

import argparse
import calendar
import json
import os
import sys
from datetime import date
from pathlib import Path
from typing import Any

from kpi.sources.sd_rk_instructions import format_report, score_instruction_kpi


def month_bounds(year: int, month: int) -> tuple[date, date]:
    last = calendar.monthrange(year, month)[1]
    return date(year, month, 1), date(year, month, last)


def load_json_rows(path: Path) -> list[dict]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(payload, list):
        return [row for row in payload if isinstance(row, dict)]
    if isinstance(payload, dict):
        for key in ("rows", "assignments", "cards", "protocols", "value"):
            rows = payload.get(key)
            if isinstance(rows, list):
                return [row for row in rows if isinstance(row, dict)]
    return []


def load_tracker_xlsx(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    ws = wb["Action Tracker"] if "Action Tracker" in wb.sheetnames else wb.active
    headers: list[str] = []
    header_idx = 0
    for index, raw in enumerate(ws.iter_rows(values_only=True), start=1):
        cells = ["" if cell is None else str(cell).strip() for cell in raw]
        if any("дата решения" in cell.casefold() for cell in cells):
            headers = cells
            header_idx = index
            break
    if not headers:
        return []
    rows: list[dict] = []
    for raw in ws.iter_rows(min_row=header_idx + 1, values_only=True):
        if not any(cell not in (None, "") for cell in raw):
            continue
        item = {}
        for header, value in zip(headers, raw):
            if header:
                item[header] = value
        rows.append(item)
    return rows


def load_tracker(path: Path) -> list[dict]:
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        return load_tracker_xlsx(path)
    return load_json_rows(path)


def fetch_odata_cards() -> list[dict]:
    from app.services.erp_assignments import handle_assignments
    from app.services.workflows.daily_assignment_playbook import CUSTOMER

    result = handle_assignments(
        {
            "action": "list",
            "customer": CUSTOMER,
            "include_all": True,
            "only_open": False,
            "include_files": True,
            "limit": 100,
        }
    )
    if result.get("error"):
        raise RuntimeError(result["error"])
    return [row for row in (result.get("assignments") or []) if isinstance(row, dict)]


def find_tracker() -> Path | None:
    env = str(os.environ.get("ACTION_TRACKER_PATH") or "").strip()
    candidates = [Path(env)] if env else []
    home = Path.home()
    candidates.extend(
        [
            home / "Downloads" / "ActionTracker.xlsx",
            home / "Desktop" / "ActionTracker.xlsx",
            Path("ActionTracker.xlsx"),
        ]
    )
    local = Path(os.environ.get("LOCALAPPDATA") or "")
    workspaces = local / "Constructor" / "agent_workspaces"
    if workspaces.is_dir():
        candidates.extend(path for path in workspaces.glob("*/ActionTracker.xlsx") if path.is_file())
    found = [path for path in candidates if path.is_file()]
    if not found:
        return None
    return max(found, key=lambda path: path.stat().st_mtime)


def compute_report(
    *,
    as_of: date | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    tracker: list[dict] | None = None,
    tracker_path: Path | None = None,
    cards: list[dict] | None = None,
    protocols: list[dict] | None = None,
) -> dict[str, Any]:
    as_of = as_of or date.today()
    if date_from is None or date_to is None:
        date_from, date_to = month_bounds(as_of.year, as_of.month)
    path = tracker_path or (None if tracker is not None else find_tracker())
    if tracker is None:
        tracker = load_tracker(path) if path else []
    if cards is None:
        cards = fetch_odata_cards()
    if protocols is None:
        protocols = fetch_odata_psd_protocols(date_from, date_to)
    report = score_instruction_kpi(
        tracker,
        cards,
        protocols,
        as_of=as_of,
        date_from=date_from,
        date_to=date_to,
    )
    report["tracker_path"] = str(path or "")
    report["date_from"] = date_from.isoformat()
    report["date_to"] = date_to.isoformat()
    return report


def evidence_line(report: dict[str, Any]) -> str:
    k1 = report.get("kpi3_1_pct")
    k2 = report.get("kpi3_2_pct")
    fact = report.get("fact_pct")
    score = report.get("score_pct")
    path = report.get("tracker_path") or "Action Tracker"
    return (
        f"{report.get('date_from')} — {report.get('date_to')}: "
        f"1С {report.get('r_total')} позиций, в Excel {report.get('r_in_tracker')}. "
        f"KPI3.1 {report.get('r24')}/{report.get('r_total')} = "
        f"{'нет данных' if k1 is None else f'{k1}%'}. "
        f"KPI3.2 {report.get('r_control')}/{report.get('r_active')} = "
        f"{'нет данных' if k2 is None else f'{k2}%'}. "
        f"min {'нет данных' if fact is None else f'{fact}%'} → "
        f"оценка {'нет данных' if score is None else f'{score}%'}. "
        f"Позиции 1С считаем внесёнными в файл. Файл: {path}"
    )


def compute_tile_update(**kwargs: Any) -> dict[str, Any]:
    report = compute_report(**kwargs)
    return {
        "id": "instructions",
        "fact": {"value": report.get("fact_pct"), "unit": "%"},
        "score_percent": report.get("score_pct"),
        "evidence": evidence_line(report),
    }


def fetch_odata_psd_protocols(date_from: date, date_to: date) -> list[dict]:
    from app.services.meeting_protocols import list_meeting_protocols

    result = list_meeting_protocols(
        {
            "meeting_kind": "sd",
            "psd_mark": True,
            "review_only": False,
            "include_closed": True,
            "date_from": date_from.isoformat(),
            "date_to": date_to.isoformat(),
            "max_results": 100,
        }
    )
    if result.get("error"):
        raise RuntimeError(result["error"])
    return [row for row in (result.get("protocols") or []) if isinstance(row, dict)]


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Реестр и контроль поручений СД+РК")
    parser.add_argument("--month", help="Период YYYY-MM")
    parser.add_argument("--from", dest="date_from")
    parser.add_argument("--to", dest="date_to")
    parser.add_argument("--as-of", dest="as_of", help="Дата расчёта YYYY-MM-DD")
    parser.add_argument("--tracker", type=Path, help="ActionTracker.xlsx или JSON строк")
    parser.add_argument("--cards", type=Path, help="JSON карточек Document_ТД_Поручения")
    parser.add_argument("--protocols", type=Path, help="JSON протоколов ПСД")
    parser.add_argument("--odata", action="store_true", help="Поручения и протоколы ПСД из 1С")
    parser.add_argument("--out", type=Path, help="UTF-8 отчёт в файл")
    parser.add_argument("--json", dest="as_json", action="store_true", help="Печатать JSON отчёта")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    if args.month:
        year_s, month_s = args.month.split("-", 1)
        date_from, date_to = month_bounds(int(year_s), int(month_s))
    elif args.date_from and args.date_to:
        date_from, date_to = date.fromisoformat(args.date_from), date.fromisoformat(args.date_to)
    else:
        today = date.today()
        date_from, date_to = month_bounds(today.year, today.month)
    as_of = date.fromisoformat(args.as_of) if args.as_of else date.today()

    if args.tracker:
        tracker = load_tracker(args.tracker)
        tracker_source = str(args.tracker)
    else:
        tracker = []
        tracker_source = "трекер не читали"

    if args.cards:
        cards = load_json_rows(args.cards)
        card_source = str(args.cards)
    elif args.odata:
        cards = fetch_odata_cards()
        card_source = f"OData поручения АСТ00, {len(cards)} шт."
    else:
        cards = []
        card_source = "поручения 1С не читали"

    if args.protocols:
        protocols = load_json_rows(args.protocols)
        protocol_source = str(args.protocols)
    elif args.odata:
        protocols = fetch_odata_psd_protocols(date_from, date_to)
        protocol_source = f"OData протоколы ПСД, {len(protocols)} шт."
    else:
        protocols = []
        protocol_source = "протоколы ПСД не читали"

    report = score_instruction_kpi(
        tracker,
        cards,
        protocols,
        as_of=as_of,
        date_from=date_from,
        date_to=date_to,
    )
    header = (
        f"Реестр и контроль поручений СД+РК  {date_from} — {date_to}  на {as_of}\n"
        f"Excel: {tracker_source}. План 1С: {card_source}; {protocol_source}\n"
    )
    body = json.dumps(report, ensure_ascii=False, indent=2) if args.as_json else format_report(report)
    text = f"{header}\n{body}\n"
    if args.out:
        args.out.write_text(text, encoding="utf-8")
    sys.stdout.reconfigure(encoding="utf-8")
    print(text, end="")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
