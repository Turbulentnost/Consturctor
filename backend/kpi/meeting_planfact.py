"""План/факт совещаний круга управления (OData 1С).

  py -m kpi.meeting_planfact --month 2026-09
  py -m kpi.meeting_planfact --from 2026-08-01 --to 2026-08-31
  py -m kpi.meeting_planfact --month 2026-08 --excel path/to/п-ф.xlsx
"""

from __future__ import annotations

import argparse
import calendar
import sys
from datetime import date
from pathlib import Path

from kpi.sources.onec_meetings import collect_plan_fact, format_table, normalize_theme_name


def month_bounds(year: int, month: int) -> tuple[date, date]:
    last = calendar.monthrange(year, month)[1]
    return date(year, month, 1), date(year, month, last)


def parse_iso_date(raw: str) -> date:
    return date.fromisoformat(raw)


def read_excel_rows(path: Path) -> list[tuple[str, int, int]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows: list[tuple[str, int, int]] = []
    for index in range(1, (sheet.max_row or 1) + 1):
        name = sheet.cell(index, 1).value
        plan = sheet.cell(index, 4).value
        fact = sheet.cell(index, 5).value
        if name and isinstance(plan, (int, float)) and isinstance(fact, (int, float)):
            rows.append((str(name).strip(), int(plan), int(fact)))
    return rows


def format_compare(report: dict, excel_rows: list[tuple[str, int, int]]) -> str:
    calc = {normalize_theme_name(row["theme"]): row for row in report["rows"]}
    lines = [
        f"{'Тема совещания':<72} {'excel':>8} {'расчёт':>8} {'ок':>4}",
        "-" * 96,
    ]
    matched = 0
    for name, plan, fact in excel_rows:
        got = calc.get(normalize_theme_name(name))
        pair = (got["plan"], got["fact"]) if got else None
        ok = pair == (plan, fact)
        if ok:
            matched += 1
        calc_text = f"{pair[0]}/{pair[1]}" if pair else "-/-"
        lines.append(f"{name[:72]:<72} {plan:>3}/{fact:<3} {calc_text:>8} {'да' if ok else 'нет':>4}")
    excel_keys = {normalize_theme_name(name) for name, _, _ in excel_rows}
    calc_keys = set(calc)
    only_excel = sorted(excel_keys - calc_keys)
    only_calc = sorted(calc_keys - excel_keys)
    lines.append("-" * 96)
    lines.append(f"совпало {matched}/{len(excel_rows)}")
    if only_excel:
        lines.append("только в excel: " + "; ".join(only_excel))
    if only_calc:
        lines.append("только в расчёте: " + "; ".join(only_calc))
    excel_plan = sum(plan for _, plan, _ in excel_rows)
    excel_fact = sum(fact for _, _, fact in excel_rows)
    lines.append(f"итого excel {excel_plan}/{excel_fact}  расчёт {report['totals']['plan']}/{report['totals']['fact']}")
    return "\n".join(lines)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="План/факт совещаний круга управления")
    parser.add_argument("--month", help="Период YYYY-MM")
    parser.add_argument("--from", dest="date_from", help="Начало периода YYYY-MM-DD")
    parser.add_argument("--to", dest="date_to", help="Конец периода YYYY-MM-DD")
    parser.add_argument("--leader", default="Донцова Анна Егоровна")
    parser.add_argument("--excel", type=Path, help="Сверить с выгрузкой 1С (.xlsx)")
    return parser


def resolve_period(args: argparse.Namespace) -> tuple[date, date]:
    if args.month:
        year_s, month_s = args.month.split("-", 1)
        return month_bounds(int(year_s), int(month_s))
    if args.date_from and args.date_to:
        return parse_iso_date(args.date_from), parse_iso_date(args.date_to)
    today = date.today()
    return month_bounds(today.year, today.month)


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    date_from, date_to = resolve_period(args)
    report = collect_plan_fact(leader_fio=args.leader, date_from=date_from, date_to=date_to)
    text = format_table(report)
    if args.excel:
        text = text + "\n\n" + format_compare(report, read_excel_rows(args.excel))
    sys.stdout.reconfigure(encoding="utf-8")
    print(text)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
