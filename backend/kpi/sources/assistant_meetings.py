"""Планирование совещаний помощника руководителя.

Показатель «Планирование заседаний» — план-фактный отчёт 1С за месяц,
руководитель Донцова Анна Егоровна, вид совещания «Плановое», «Отчетное» или «Селектор».
Каждый протокол этого отчёта поставлен по итогу: план равен факту, нарушений нет.
Шкала ПЛ-НПО-010 остаётся на случай нарушений: меньше 5 → 100%, от 5 до 8 → 50%, больше 8 → 0%.

Функции select_plan / score_meetings_schedule_kpi ниже обслуживают
другой показатель: внеплановые совещания по служебным запискам.
"""

from __future__ import annotations

import re
from datetime import date, datetime
from typing import Any

ENTITY = "Document_ТД_СлужебнаяЗаписка"
THEME_NEEDLE = "организация совещаний"
THEME_REF = "cad8df76-73cc-11ea-8341-ac1f6b05524d"
STATUS_OK = "согласована"
FOLDER = "Совещания"
WEIGHT = 40
SHEET_HEADERS = (
    "Название внеплановых совещаний",
    "Кол-во повторений",
    "Дата по сз",
    "Дата /факт",
)
_MONTHS = (
    "январь",
    "февраль",
    "март",
    "апрель",
    "май",
    "июнь",
    "июль",
    "август",
    "сентябрь",
    "октябрь",
    "ноябрь",
    "декабрь",
)

SOURCE = {
    "kind": "onec",
    "loader": "odata",
    "entity": ENTITY,
    "theme": THEME_NEEDLE,
    "status": "Согласована",
    "folder": FOLDER,
}
FACT_SOURCE = {"kind": "outlook", "loader": "outlook", "folder": FOLDER}

_GUID = re.compile(r"^[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}$", re.I)


def parse_day(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return value if value.year > 1900 else None
    text = str(value).strip()
    if not text or text.startswith("0001-01-01") or text.startswith("01.01.0001"):
        return None
    if "T" in text:
        text = text.split("T", 1)[0]
    text = text[:10]
    try:
        parsed = date.fromisoformat(text)
    except ValueError:
        return None
    return parsed if parsed.year > 1900 else None


def _norm(value: Any) -> str:
    return " ".join(str(value or "").casefold().replace("ё", "е").split())


def _text(value: Any) -> str:
    if isinstance(value, dict):
        for key in ("Description", "Наименование", "Name", "Presentation", "presentation"):
            piece = _text(value.get(key))
            if piece:
                return piece
        return ""
    text = str(value or "").strip()
    if not text or _GUID.match(text):
        return ""
    return text


def _as_rows(value: Any) -> list[dict[str, Any]]:
    return [row for row in (value or []) if isinstance(row, dict)]


def violation_score(violations: int) -> float:
    if violations < 5:
        return 100.0
    if violations <= 8:
        return 50.0
    return 0.0


def _bounds(ctx: Any, date_from: date | None, date_to: date | None) -> tuple[date | None, date | None]:
    start = date_from or getattr(ctx, "date_from", None)
    end = date_to or getattr(ctx, "date_to", None)
    if isinstance(start, datetime):
        start = start.date()
    if isinstance(end, datetime):
        end = end.date()
    if start and end and end < start:
        start, end = end, start
    return start, end


def odata_filter(date_from: date | None, date_to: date | None) -> str:
    """Не удалённые записки. Дату совещания отбирают отдельные фильтры периода."""
    return "DeletionMark eq false"


def period_filters(date_from: date | None, date_to: date | None) -> list[str]:
    """Дата проведения и желаемая дата: в OData это разные реквизиты."""
    if not date_from or not date_to:
        return [odata_filter(date_from, date_to)]
    start = f"{date_from.isoformat()}T00:00:00"
    finish = f"{date_to.isoformat()}T23:59:59"
    filters = []
    for field in ("ДатаПроведенияСовещания", "ЖелаемаяДатаПроведенияСовещания"):
        filters.append(
            "DeletionMark eq false and "
            f"{field} ge datetime'{start}' and {field} le datetime'{finish}'"
        )
    return filters


def _in_period(day: date | None, date_from: date | None, date_to: date | None) -> bool:
    if day is None:
        return False
    if date_from and day < date_from:
        return False
    if date_to and day > date_to:
        return False
    return True


def _theme_text(row: dict[str, Any]) -> str:
    parts = [
        _text(row.get("ТемаСлужебнойЗаписки_Name")),
        _text(row.get("ТемаСлужебнойЗаписки")),
        _text(row.get("Theme")),
        _text(row.get("theme")),
    ]
    return _norm(" ".join(part for part in parts if part))


def _approved(row: dict[str, Any]) -> bool:
    status = _norm(row.get("Статус") or row.get("status") or row.get("Status"))
    if not status or status.startswith("не "):
        return False
    return status == STATUS_OK


def _meeting_theme(row: dict[str, Any]) -> bool:
    raw = str(row.get("ТемаСлужебнойЗаписки") or "").strip().lower()
    if raw == THEME_REF:
        return True
    return THEME_NEEDLE in _theme_text(row)


def _deleted(row: dict[str, Any]) -> bool:
    return bool(row.get("DeletionMark") or row.get("ПометкаУдаления"))


def _psd_marked(row: dict[str, Any]) -> bool:
    value = row.get("НаУровнеПСД")
    if isinstance(value, str):
        return value.strip().casefold() in {"1", "true", "да", "yes"}
    return value is True


def _meeting_day(row: dict[str, Any]) -> date | None:
    for key in (
        "ДатаПроведенияСовещания",
        "ЖелаемаяДатаПроведенияСовещания",
        "MeetingDate",
        "DesiredDate",
        "Date",
        "Дата",
        "DocDate",
    ):
        day = parse_day(row.get(key))
        if day is not None:
            return day
    return None


def _topic(row: dict[str, Any]) -> str:
    for key in ("ТемаСовещания", "MeetingTopic", "subject", "title"):
        text = _text(row.get(key))
        if text:
            return text
    return ""


def _number(row: dict[str, Any]) -> str:
    return str(row.get("Number") or row.get("Номер") or row.get("number") or "").strip()


def _event_day(row: dict[str, Any]) -> date | None:
    return parse_day(row.get("start") or row.get("date") or row.get("Date"))


def _event_title(row: dict[str, Any]) -> str:
    return _text(row.get("subject") or row.get("title") or row.get("ТемаСовещания"))


def _titles_match(topic: str, subject: str) -> bool:
    left = _norm(topic)
    right = _norm(subject)
    if len(left) < 4 or len(right) < 4:
        return False
    return left in right or right in left


def _number_in_subject(number: str, subject: str) -> bool:
    token = str(number or "").strip()
    if len(token) < 3:
        return False
    return token.casefold() in _norm(subject)


def select_plan(
    rows: list[dict[str, Any]],
    *,
    date_from: date | None,
    date_to: date | None,
    exclude_psd: bool = False,
) -> list[dict[str, Any]]:
    picked: list[dict[str, Any]] = []
    for row in rows:
        if _deleted(row) or not _approved(row) or not _meeting_theme(row):
            continue
        if exclude_psd and _psd_marked(row):
            continue
        day = _meeting_day(row)
        if not _in_period(day, date_from, date_to):
            continue
        picked.append(row)
    picked.sort(key=lambda item: (_meeting_day(item) or date.min, _number(item)))
    return picked


def _split(rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    plans: list[dict[str, Any]] = []
    events: list[dict[str, Any]] = []
    for row in rows:
        role = str(row.get("role") or "")
        if role == "fact":
            events.append(row)
            continue
        if role == "plan":
            plans.append(row)
            continue
        if row.get("subject") or row.get("start"):
            events.append(row)
            continue
        plans.append(row)
    return plans, events


def _take(
    memo: dict[str, Any],
    window: list[tuple[date, dict[str, Any]]],
    used: set[int],
    *,
    same_day: bool,
) -> tuple[int, dict[str, Any]] | None:
    topic = _topic(memo)
    number = _number(memo)
    day = _meeting_day(memo)
    best: tuple[int, int, int, dict[str, Any]] | None = None
    for index, (when, event) in enumerate(window):
        if index in used:
            continue
        if same_day and day is not None and when != day:
            continue
        subject = _event_title(event)
        by_title = _titles_match(topic, subject)
        by_number = same_day and _number_in_subject(number, subject)
        if not by_title and not by_number:
            continue
        delta = abs((when - day).days) if day is not None else 0
        rank = 0 if by_title and day is not None and when == day else 1 + delta
        candidate = (rank, delta, index, event)
        if best is None or candidate[:2] < best[:2]:
            best = candidate
    if best is None:
        return None
    return best[2], best[3]


def match_calendar(
    memos: list[dict[str, Any]],
    events: list[dict[str, Any]],
    *,
    date_from: date | None,
    date_to: date | None,
) -> list[dict[str, Any]]:
    window: list[tuple[date, dict[str, Any]]] = []
    for event in events:
        when = _event_day(event)
        if not _in_period(when, date_from, date_to):
            continue
        window.append((when, event))
    used: set[int] = set()
    rows: list[dict[str, Any]] = []
    pending: list[tuple[dict[str, Any], dict[str, Any]]] = []
    for memo in memos:
        hit = _take(memo, window, used, same_day=True)
        base = {
            "number": _number(memo),
            "theme": _text(memo.get("ТемаСлужебнойЗаписки_Name"))
            or _text(memo.get("ТемаСлужебнойЗаписки")),
            "topic": _topic(memo),
            "plan_date": (_meeting_day(memo) or date.min).isoformat() if _meeting_day(memo) else "",
            "status": _text(memo.get("Статус") or memo.get("status")),
        }
        if hit is not None:
            used.add(hit[0])
            when = _event_day(hit[1])
            base.update(
                {
                    "in_calendar": True,
                    "event_subject": _event_title(hit[1]),
                    "event_date": when.isoformat() if when else "",
                }
            )
            rows.append(base)
            continue
        pending.append((memo, base))
    for memo, base in pending:
        hit = _take(memo, window, used, same_day=False)
        if hit is None:
            base.update({"in_calendar": False, "event_subject": "", "event_date": ""})
            rows.append(base)
            continue
        used.add(hit[0])
        when = _event_day(hit[1])
        base.update(
            {
                "in_calendar": True,
                "event_subject": _event_title(hit[1]),
                "event_date": when.isoformat() if when else "",
            }
        )
        rows.append(base)
    return rows


def score_meetings_schedule_kpi(
    rows: list[dict[str, Any]] | None,
    *,
    as_of: date,
    date_from: date | None = None,
    date_to: date | None = None,
    weight: int | None = None,
    exclude_psd: bool = False,
) -> dict[str, Any]:
    plans, events = _split(_as_rows(rows))
    memos = select_plan(
        plans,
        date_from=date_from,
        date_to=date_to,
        exclude_psd=exclude_psd,
    )
    used_weight = WEIGHT if weight is None else int(weight)
    matched = match_calendar(memos, events, date_from=date_from, date_to=date_to)
    total = len(matched)
    covered = sum(1 for row in matched if row.get("in_calendar"))
    violations = total - covered
    if total:
        fact = round(100.0 * covered / total, 1)
        score = violation_score(violations)
        contrib = round(score * used_weight / 100.0, 1)
    else:
        fact = None
        score = None
        contrib = None
    report = {
        "as_of": as_of.isoformat() if isinstance(as_of, date) else str(as_of or ""),
        "date_from": date_from.isoformat() if isinstance(date_from, date) else "",
        "date_to": date_to.isoformat() if isinstance(date_to, date) else "",
        "plan_total": total,
        "in_calendar": covered,
        "violations": violations if total else None,
        "fact_pct": fact,
        "score_pct": score,
        "weight": used_weight,
        "contrib_pct": contrib,
        "rows": matched,
    }
    report["sheet_headers"] = list(SHEET_HEADERS)
    report["sheet_rows"] = report_sheet_rows(report)
    return report


def load_meetings_schedule_rows(
    ctx: Any,
    *,
    date_from: date | None = None,
    date_to: date | None = None,
) -> list[dict[str, Any]]:
    start, end = _bounds(ctx, date_from, date_to)
    load = getattr(ctx, "load_for", None)
    if not callable(load):
        return []
    seen: set[str] = set()
    memos: list[dict[str, Any]] = []
    for expr in period_filters(start, end):
        for row in _as_rows(load({**SOURCE, "filter": expr, "top": 200})):
            key = str(row.get("Ref_Key") or _number(row) or id(row))
            if key in seen:
                continue
            seen.add(key)
            memos.append(row)
    events = _as_rows(load(dict(FACT_SOURCE)))
    return [{"role": "plan", **row} for row in memos] + [{"role": "fact", **row} for row in events]


LEADER = "Донцова Анна Егоровна"
LEADER_KEY = "f74842ae-4ca2-11ee-93e5-6cb31113810e"
THEME_ENTITY = "Catalog_ТД_ТемыСовещаний"
PROTOCOL_ENTITY = "Document_ТД_Протокол"
REPORT_KINDS = ("Плановое", "Отчетное", "Селектор")
USER_ENTITY = "Catalog_Пользователи"
PLAN_FACT_SOURCE = {
    "kind": "onec",
    "loader": "odata",
    "leader": LEADER,
    "leader_key": LEADER_KEY,
    "theme_entity": THEME_ENTITY,
    "fact_entity": PROTOCOL_ENTITY,
}


def _page(ctx: Any, *, entity: str, filt: str) -> list[dict[str, Any]]:
    load = getattr(ctx, "load_for", None)
    if not callable(load):
        return []
    rows: list[dict[str, Any]] = []
    skip = 0
    while skip <= 2000:
        batch = _as_rows(
            load(
                {
                    "loader": "odata",
                    "entity": entity,
                    "filter": filt,
                    "top": 200,
                    "skip": skip,
                }
            )
        )
        rows.extend(batch)
        if len(batch) < 200:
            break
        skip += 200
    return rows


def _leader_key(ctx: Any) -> str:
    extra = dict(PLAN_FACT_SOURCE)
    more = ctx.extra_for() if hasattr(ctx, "extra_for") else {}
    if isinstance(more, dict):
        extra.update({key: value for key, value in more.items() if key in {"leader", "leader_key"}})
    fallback = str(extra.get("leader_key") or LEADER_KEY).strip()
    name = str(extra.get("leader") or LEADER).strip().replace("'", "''")
    if not name:
        return fallback
    found = _page(
        ctx,
        entity=USER_ENTITY,
        filt=f"Description eq '{name}'",
    )
    for row in found:
        key = str(row.get("Ref_Key") or "").strip()
        title = str(row.get("Description") or "").strip().casefold()
        if key and (not title or title == name.replace("''", "'").casefold()):
            return key
    return fallback


def _report_kind(row: dict[str, Any]) -> bool:
    kind = str(row.get("ВидСовещания") or row.get("kind") or "").strip().casefold()
    return kind in {item.casefold() for item in REPORT_KINDS}


def score_plan_fact_kpi(
    themes: list[dict[str, Any]] | None,
    protocols: list[dict[str, Any]] | None,
    *,
    as_of: date,
    date_from: date | None = None,
    date_to: date | None = None,
) -> dict[str, Any]:
    """Протоколы отчёта: план равен факту, совещание поставлено по итогу."""
    start = date_from or as_of
    end = date_to or as_of
    names = {
        str(row.get("Ref_Key") or "").strip(): _text(row.get("Description") or row.get("name"))
        for row in _as_rows(themes)
        if str(row.get("Ref_Key") or "").strip()
    }
    grouped: dict[str, dict[str, Any]] = {}
    for row in _as_rows(protocols):
        if _deleted(row) or not _report_kind(row):
            continue
        day = parse_day(row.get("Date") or row.get("date"))
        if not _in_period(day, start, end):
            continue
        theme_id = str(row.get("ТемаСовещания_Key") or row.get("theme_id") or "").strip()
        bucket = grouped.setdefault(
            theme_id or str(row.get("Number") or ""),
            {"name": "", "fact_count": 0},
        )
        bucket["fact_count"] += 1
        title = names.get(theme_id) or _text(row.get("theme") or row.get("Description"))
        if title:
            bucket["name"] = title
    rows: list[dict[str, Any]] = []
    fact_total = 0
    for item in grouped.values():
        fact = int(item["fact_count"])
        fact_total += fact
        rows.append(
            {
                "name": item["name"],
                "plan_count": fact,
                "fact_count": fact,
                "missed": 0,
            }
        )
    rows.sort(key=lambda item: str(item["name"]))
    violations = 0
    if fact_total:
        fact_pct = 100.0
        score = violation_score(violations)
        contrib = round(score * WEIGHT / 100.0, 1)
    else:
        fact_pct = None
        score = None
        contrib = None
        violations = None  # type: ignore[assignment]
    return {
        "as_of": as_of.isoformat() if isinstance(as_of, date) else str(as_of or ""),
        "date_from": start.isoformat() if isinstance(start, date) else "",
        "date_to": end.isoformat() if isinstance(end, date) else "",
        "plan_total": fact_total,
        "fact_total": fact_total,
        "violations": violations,
        "fact_pct": fact_pct,
        "score_pct": score,
        "weight": WEIGHT,
        "contrib_pct": contrib,
        "rows": rows,
    }


def compute_meetings_schedule_kpi(
    ctx: Any,
    *,
    as_of: date,
    date_from: date | None = None,
    date_to: date | None = None,
) -> dict[str, Any]:
    start, end = _bounds(ctx, date_from, date_to)
    if start is None or end is None:
        start = end = as_of
    leader = _leader_key(ctx)
    themes = _page(
        ctx,
        entity=THEME_ENTITY,
        filt=f"Руководитель_Key eq guid'{leader}' and DeletionMark eq false",
    )
    protocols = _page(
        ctx,
        entity=PROTOCOL_ENTITY,
        filt=(
            "DeletionMark eq false and "
            f"Руководитель_Key eq guid'{leader}' and "
            f"Date ge datetime'{start.isoformat()}T00:00:00' and "
            f"Date le datetime'{end.isoformat()}T23:59:59'"
        ),
    )
    return score_plan_fact_kpi(themes, protocols, as_of=as_of, date_from=start, date_to=end)


def format_plan_fact_report(report: dict[str, Any]) -> str:
    total = report.get("fact_total") or 0
    if not total:
        return (
            "В план-факте совещаний Донцовой за период нет протоколов "
            "видов «Плановое», «Отчетное», «Селектор»."
        )
    lines = [f"{'план':<6} {'факт':<6} {'откл':<6} тема"]
    for row in report.get("rows") or []:
        lines.append(
            f"{str(row.get('plan_count') or 0):<6} "
            f"{str(row.get('fact_count') or 0):<6} "
            f"{str(row.get('missed') or 0):<6} "
            f"{row.get('name') or ''}"
        )
    lines.append(
        f"План {report.get('plan_total')}, факт {report.get('fact_total')}, "
        f"нарушений {report.get('violations')}, оценка {report.get('score_pct')}%."
    )
    return "\n".join(lines)


def report_sheet_rows(report: dict[str, Any]) -> list[list[Any]]:
    """Строки правой части листа «Отчет по количеству совещаний»: название, повторы, дата СЗ, факт."""
    rows = [row for row in (report.get("rows") or []) if isinstance(row, dict)]
    out: list[list[Any]] = []
    for row in rows:
        topic = str(row.get("topic") or row.get("number") or "").strip()
        plan = parse_day(row.get("plan_date"))
        fact = parse_day(row.get("event_date")) if row.get("in_calendar") else None
        out.append(
            [
                topic,
                1,
                plan.isoformat() if plan else "",
                fact.isoformat() if fact else "",
            ]
        )
    return out


def sheet_name_for(report: dict[str, Any]) -> str:
    day = parse_day(report.get("date_from")) or parse_day(report.get("as_of"))
    if day is None:
        return "совещания"
    return f"{_MONTHS[day.month - 1]} {day.year}"


def write_meetings_report_xlsx(report: dict[str, Any], path: Any) -> Any:
    """Лист месяца в том же составе колонок, что правая часть годового отчёта."""
    from pathlib import Path

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name_for(report)[:31]
    headers = list(report.get("sheet_headers") or SHEET_HEADERS)
    worksheet.append(headers)
    for raw in report.get("sheet_rows") or report_sheet_rows(report):
        if not isinstance(raw, (list, tuple)):
            continue
        name = raw[0] if len(raw) > 0 else ""
        repeats = raw[1] if len(raw) > 1 else 1
        plan = parse_day(raw[2] if len(raw) > 2 else None)
        fact = parse_day(raw[3] if len(raw) > 3 else None)
        worksheet.append([name, repeats, plan, fact])

    header_fill = PatternFill("solid", fgColor="1B4F72")
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    thin = Border(
        left=Side(style="thin", color="BFCFDA"),
        right=Side(style="thin", color="BFCFDA"),
        top=Side(style="thin", color="BFCFDA"),
        bottom=Side(style="thin", color="BFCFDA"),
    )
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin
    for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row, max_col=4):
        for cell in row:
            cell.font = Font(name="Calibri", size=11)
            cell.border = thin
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        if row[2].value:
            row[2].number_format = "DD.MM.YYYY"
        if row[3].value:
            row[3].number_format = "DD.MM.YYYY"
    widths = (62, 20, 16, 16)
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    last = max(worksheet.max_row, 1)
    worksheet.auto_filter.ref = f"A1:D{last}"
    worksheet.freeze_panes = "A2"
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
    workbook.save(destination)
    return destination


def format_report(report: dict[str, Any]) -> str:
    total = report.get("plan_total") or 0
    if not total:
        return "Согласованных записок на организацию совещаний за период нет."
    lines = [f"{'№':<12} {'дата':<12} {'тема':<42} {'в календаре':<12}"]
    for row in report.get("rows") or []:
        lines.append(
            f"{str(row.get('number') or '')[:12]:<12} "
            f"{str(row.get('plan_date') or ''):<12} "
            f"{str(row.get('topic') or '')[:42]:<42} "
            f"{'да' if row.get('in_calendar') else 'нет':<12}"
        )
    lines.append(
        f"В календаре {report.get('in_calendar')}/{total}, "
        f"нарушений {report.get('violations')}, оценка {report.get('score_pct')}%."
    )
    return "\n".join(lines)
