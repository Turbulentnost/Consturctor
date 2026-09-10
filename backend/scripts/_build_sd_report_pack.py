"""Build SD GK report templates from PL-34-242 v.03."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

OUT = Path(r"C:\Users\mdj\Desktop\Отчет СД ГК ПЛ-34-242")
COPY = Path(r"C:\Users\mdj\Desktop\конструктор\sd-report-pl-34-242")

PERIOD = "август 2026"
MEETING = "СД-ГК-2026-08"
MEETING_DATES = "24–25 сентября 2026"
DEADLINE_PACK = "22.09.2026"
SOURCE = "ПЛ-34-242, версия 03"

COMPANIES_D1 = [
    "ООО НПО «Турбулентность-ДОН»",
    "ООО «Алмаз»",
    "ООО «Метрогазсервис»",
    "ООО «Амурская легенда»",
    "ООО «Милака»",
    "ООО «Малика»",
    "ООО «Мурар Террус»",
]
COMPANIES_D2 = ["ООО «ИТЦ»", "ООО «Авион»", "ИП Магакян Е.И."]
COMPANIES = COMPANIES_D1 + COMPANIES_D2

HEADER = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
TITLE = Font(name="Calibri", bold=True, size=14, color="1F4E79")
BOLD = Font(name="Calibri", bold=True, size=11)
NORMAL = Font(name="Calibri", size=11)
FILL = PatternFill("solid", fgColor="1F4E79")
FILL2 = PatternFill("solid", fgColor="D6E3F0")
FILL_YES = PatternFill("solid", fgColor="C6EFCE")
FILL_NO = PatternFill("solid", fgColor="FFC7CE")
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
WRAP = Alignment(wrap_text=True, vertical="center")


def style_header(ws: Worksheet, row: int, cols: int) -> None:
    for col in range(1, cols + 1):
        cell = ws.cell(row, col)
        cell.font = HEADER
        cell.fill = FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = THIN


def style_body(ws: Worksheet, start: int, end: int, cols: int) -> None:
    zebra = PatternFill("solid", fgColor="F3F7FB")
    for row in range(start, end + 1):
        for col in range(1, cols + 1):
            cell = ws.cell(row, col)
            cell.font = NORMAL
            cell.alignment = WRAP
            cell.border = THIN
            if row % 2 == 0:
                cell.fill = zebra


def autosize(ws: Worksheet, widths: list[int]) -> None:
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:3"


def banner(ws: Worksheet, title: str, note: str, cols: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=cols)
    ws["A1"] = title
    ws["A1"].font = TITLE
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=cols)
    ws["A2"] = note
    ws["A2"].font = Font(name="Calibri", italic=True, size=10, color="5B7A99")
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[3].height = 32


def write_table(ws: Worksheet, headers: list[str], rows: list[list[object]], widths: list[int]) -> None:
    cols = len(headers)
    for col, name in enumerate(headers, start=1):
        ws.cell(3, col, name)
    style_header(ws, 3, cols)
    for ridx, row in enumerate(rows, start=4):
        for cidx, value in enumerate(row, start=1):
            ws.cell(ridx, cidx, value)
    if rows:
        style_body(ws, 4, 3 + len(rows), cols)
    autosize(ws, widths)


def new_wb() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def add_checklist(wb: Workbook) -> None:
    ws = wb.create_sheet("Комплектность", 0)
    banner(
        ws,
        f"Чек-лист комплекта {MEETING} · отчётный период {PERIOD}",
        f"Обязательный пакет п. 6.4 {SOURCE}. Рассылка не позднее {DEADLINE_PACK}. "
        "Пункт пакета есть только если в колонке «Файл» указано имя или на листе формы заполнены цифры. "
        "Файл запуска — не класть в постоянное знание.",
        7,
    )
    rows = [
        [1, "Краткое резюме (1 стр.)", "п. 6.4 / пакет 1", "да", "", "нет", ""],
        [2, "Финансы план/факт: ОПУ, ДДС, ключевые отклонения", "п. 6.4 / пакет 2", "да", "", "нет", ""],
        [3, "Капитал и инвестиции: вложения, бизнес-кейсы, статус проектов", "п. 6.4 / пакет 3", "да", "", "нет", ""],
        [4, "Продажи/клиенты: контракты, риски, дебиторка", "п. 6.4 / пакет 4", "да", "", "нет", ""],
        [5, "Производство/качество: выпуск, себестоимость, рекламации", "п. 6.4 / пакет 5", "да", "", "нет", ""],
        [6, "Проекты/НИОКР: портфель, эффекты, риски", "п. 6.4 / пакет 6", "да", "", "нет", ""],
        [7, "Риски и комплаенс: ТОП-10, КИР, инциденты", "п. 6.4 / пакет 7", "да", "", "нет", ""],
        [8, "Персонал и преемственность: ключевые позиции, резерв, текучесть", "п. 6.4 / пакет 8", "да", "", "нет", ""],
        [9, "Статус исполнения решений/поручений ПСД", "п. 6.4 / пакет 9", "да", "", "нет", ""],
        [10, "Решения на согласование (класс А / сверх DoA)", "п. 6.4 / пакет 10", "да", "", "нет", ""],
        [11, "Приложение А. План-факт бюджета доходов и расходов", "прил. А", "да", "", "нет", ""],
        [12, "Приложение Б. Бюджет доходов и расходов", "прил. Б", "да", "", "нет", ""],
        [13, "БДР, БДДС, анализ ДЗ/КЗ по компаниям", "п. 6.1 / подготовка", "да", "", "нет", ""],
        [14, "Задолженность и претензии свыше 1 000 000 ₽ от одного контрагента", "п. 5.3, 6.3", "да", "", "нет", ""],
        [15, "Подпись гл. бухгалтера, директора и финансового директора", "п. 6.4", "да", "", "нет", ""],
        [16, "Доклад инспекционной группы (проверка ИИ)", "п. 6.4", "да", "", "нет", ""],
        [17, "Повестка на оба дня СД ГК", "п. 6.3", "да", "", "нет", ""],
        [18, "Реестр рассылки участникам", "п. 5.2, 6.4", "да", "", "нет", ""],
        [19, "Протокол предыдущего заседания", "п. 5.2, 6.5", "да", "", "нет", ""],
        [20, "Черновик протокола текущего заседания", "п. 6.5", "после заседания", "", "ожидается", ""],
        [21, "Decision Log (новые решения)", "п. 6.5", "после заседания", "", "ожидается", ""],
        [22, "Action Tracker (новые поручения)", "п. 6.5", "после заседания", "", "ожидается", ""],
    ]
    write_table(
        ws,
        ["№", "Документ / материал", "Норма ПЛ-34-242", "Обязательный", "Файл", "Статус", "Замечание агента"],
        rows,
        [6, 62, 22, 18, 28, 16, 36],
    )
    ws["A27"] = "Правило наличия:"
    ws["B27"] = (
        "Пункт пакета считается приложенным только если в колонке «Файл» есть имя "
        "или на соответствующем листе (ОПУ, ДДС, А, Б, инвестиции, продажи/ДЗ, "
        "производство, НИОКР, риски, персонал, класс А, претензии) заполнены цифры. "
        "Пустой шаблон листа — пункта нет. Не хранить этот Excel в постоянном знании."
    )
    ws["A27"].font = BOLD
    ws.merge_cells("B27:G27")
    ws["B27"].alignment = WRAP
    ws.row_dimensions[27].height = 48
    ws["A28"] = "Итог комплекта:"
    ws["B28"] = "полный / неполный"
    ws["A28"].font = BOLD


def add_pack(wb: Workbook) -> None:
    ws = wb.create_sheet("ОПУ план-факт")
    banner(ws, f"Финансы план/факт · ОПУ · {PERIOD}", f"Пакет 2 {SOURCE}. По компаниям ГК.", 8)
    rows = []
    for name in COMPANIES:
        rows.append([name, "Выручка", "", "", "", "", "", ""])
        rows.append([name, "Себестоимость", "", "", "", "", "", ""])
        rows.append([name, "Валовая прибыль", "", "", "", "", "", ""])
        rows.append([name, "Операционные расходы", "", "", "", "", "", ""])
        rows.append([name, "Прибыль", "", "", "", "", "", ""])
        rows.append([name, "Рентабельность, %", "", "", "", "", "", ""])
    write_table(
        ws,
        ["Компания", "Статья", "План месяца", "Факт месяца", "Откл., ₽", "Откл., %", "План с нач. года", "Факт с нач. года"],
        rows,
        [36, 26, 16, 16, 14, 12, 18, 18],
    )

    ws = wb.create_sheet("ДДС")
    banner(ws, f"Бюджет движения денежных средств · {PERIOD}", f"Пакет 2 / подготовка п. 6.1 {SOURCE}.", 8)
    rows = []
    for name in COMPANIES:
        for article in ("Приход", "Расход", "Чистый денежный поток"):
            rows.append([name, article, "", "", "", "", "", ""])
    write_table(
        ws,
        ["Организация", "Статья / папка статьи", "Документ движения", "Контрагент", "План", "Факт", "Откл., ₽", "Откл., %"],
        rows,
        [36, 28, 24, 28, 14, 14, 14, 12],
    )

    ws = wb.create_sheet("Прил А план-факт БДР")
    banner(ws, "Приложение А. План-факт бюджета доходов и расходов", SOURCE, 9)
    write_table(
        ws,
        ["Статья", "План месяца", "Сопост. план", "Факт месяца", "%", "План с нач. года", "Сопост. план с нач. года", "Факт с нач. года", "%"],
        [
            ["Выручка", "", "", "", "", "", "", "", ""],
            ["Доход", "", "", "", "", "", "", "", ""],
            ["Прибыль", "", "", "", "", "", "", "", ""],
            ["Рентабельность", "", "", "", "", "", "", "", ""],
            ["Себестоимость", "", "", "", "", "", "", "", ""],
            ["Прямые затраты", "", "", "", "", "", "", "", ""],
            ["Производственные расходы", "", "", "", "", "", "", "", ""],
            ["Общехозяйственные расходы", "", "", "", "", "", "", "", ""],
            ["Прочие расходы", "", "", "", "", "", "", "", ""],
            ["Расходы за счёт прибыли", "", "", "", "", "", "", "", ""],
        ],
        [32, 16, 16, 16, 10, 18, 22, 18, 10],
    )

    ws = wb.create_sheet("Прил Б БДР")
    banner(ws, "Приложение Б. Бюджет доходов и расходов", SOURCE, 6)
    rows = []
    for name in COMPANIES:
        rows.append([name, "", "", "", "", ""])
    write_table(
        ws,
        ["Компания", "План месяца", "Факт месяца", "Выполнение, %", "План с нач. года", "Факт с нач. года"],
        rows,
        [40, 16, 16, 16, 20, 20],
    )

    ws = wb.create_sheet("Инвестиции")
    banner(ws, f"Капитал и инвестиции · {PERIOD}", "Пакет 3. Сверх лимита DoA — эскалация на СД.", 9)
    write_table(
        ws,
        ["Проект / бизнес-кейс", "Компания", "Статус", "Лимит DoA, ₽", "Заявка, ₽", "Откл., ₽", "Нужно СД", "Риск", "Решение"],
        [["", "", "", "", "", "", "", "", ""]],
        [32, 28, 16, 16, 14, 12, 12, 20, 22],
    )

    ws = wb.create_sheet("Продажи и ДЗ")
    banner(ws, f"Продажи / клиенты / дебиторка · {PERIOD}", "Пакет 4.", 8)
    write_table(
        ws,
        ["Компания", "Ключевой контракт", "Клиент", "Сумма, ₽", "Дебиторка, ₽", "Просрочка, ₽", "Риск", "Мера"],
        [[name, "", "", "", "", "", "", ""] for name in COMPANIES],
        [36, 28, 24, 14, 16, 14, 18, 24],
    )

    ws = wb.create_sheet("Производство")
    banner(ws, f"Производство / качество · {PERIOD}", "Пакет 5.", 8)
    write_table(
        ws,
        ["Компания / ЦФО", "Выпуск", "План", "Факт", "Себестоимость", "Рекламации", "Брак", "Комментарий"],
        [
            ["Производство 1 (ЦФО 1)", "", "", "", "", "", "", ""],
            ["Производство 2 АЛМАЗ (ЦФО 1)", "", "", "", "", "", "", ""],
            ["Производство БМИ (ЦФО 1)", "", "", "", "", "", "", ""],
            ["Экспериментальный цех (ЦФО 1)", "", "", "", "", "", "", ""],
            ["Сервисная служба (ЦФО 1)", "", "", "", "", "", "", ""],
            ["Метрологическая служба (ЦФО 1)", "", "", "", "", "", "", ""],
        ],
        [34, 14, 12, 12, 16, 14, 12, 28],
    )

    ws = wb.create_sheet("НИОКР")
    banner(ws, f"Проекты / НИОКР · {PERIOD}", "Пакет 6.", 7)
    write_table(
        ws,
        ["Проект", "Владелец", "Статус", "Эффект", "Риск", "Срок", "Нужно решение СД"],
        [["", "", "", "", "", "", ""]],
        [28, 24, 16, 20, 22, 14, 20],
    )

    ws = wb.create_sheet("Риски ТОП-10")
    banner(ws, f"Риски и комплаенс · {PERIOD}", "Пакет 7.", 7)
    write_table(
        ws,
        ["№", "Риск / инцидент", "Компания", "КИР", "Уровень", "Владелец", "Мера"],
        [[i, "", "", "", "", "", ""] for i in range(1, 11)],
        [6, 36, 28, 16, 12, 22, 28],
    )

    ws = wb.create_sheet("Персонал")
    banner(ws, f"Персонал и преемственность · {PERIOD}", "Пакет 8.", 7)
    write_table(
        ws,
        ["Компания", "Ключевая позиция", "Резерв", "Текучесть, %", "Индекс вовлечённости", "Риск", "Мера"],
        [[name, "", "", "", "", "", ""] for name in COMPANIES],
        [36, 26, 18, 14, 22, 18, 24],
    )

    ws = wb.create_sheet("Класс А")
    banner(ws, "Решения на согласование (класс А)", "Пакет 10. Вопросы сверх лимитов DoA.", 8)
    write_table(
        ws,
        ["№", "Вопрос", "Компания", "Сумма, ₽", "Лимит DoA", "Почему СД", "Проект решения", "Статус"],
        [[1, "", "", "", "", "", "", "на согласование"]],
        [6, 36, 28, 14, 14, 22, 32, 18],
    )

    ws = wb.create_sheet("Претензии >1 млн")
    banner(
        ws,
        "Задолженность и претензии свыше 1 000 000 ₽ от одного контрагента",
        "Обязательный доклад директоров и главных бухгалтеров, п. 5.3 и 6.3.",
        8,
    )
    write_table(
        ws,
        ["Компания", "Контрагент", "Вид (ДЗ / КЗ / претензия)", "Сумма, ₽", "Дата возникновения", "Срок", "Риск", "Предлагаемая мера"],
        [[name, "", "", "", "", "", "", ""] for name in COMPANIES],
        [36, 28, 22, 14, 18, 14, 16, 28],
    )


def add_registers(wb: Workbook) -> None:
    ws = wb.create_sheet("Decision Log")
    banner(ws, f"Decision Log · {MEETING}", "п. 5.2, 6.5. Фиксация в течение 1 рабочего дня после заседания.", 9)
    write_table(
        ws,
        ["ID", "Дата СД", "Уровень (СД ГК / ОКМ / ИК / CAB / KPI)", "Формулировка решения", "Компания", "Класс (A/B)", "DoA", "Статус", "Ссылка на протокол"],
        [["DEC-", MEETING_DATES, "СД ГК", "", "", "", "", "принято", MEETING]],
        [12, 16, 28, 40, 28, 12, 12, 14, 16],
    )

    ws = wb.create_sheet("Action Tracker")
    banner(ws, f"Action Tracker · {MEETING}", "Владелец, срок, критерий приёмки, статус. 1 рабочий день после заседания.", 10)
    write_table(
        ws,
        ["ID", "Решение DEC", "Поручение", "Владелец", "Компания", "Срок", "Критерий приёмки", "Статус", "Просрочка", "Комментарий ПСД"],
        [["AT-", "DEC-", "", "", "", "", "", "открыто", "", ""]],
        [10, 12, 36, 22, 28, 14, 28, 14, 12, 24],
    )

    ws = wb.create_sheet("Рассылка")
    banner(ws, f"Реестр рассылки комплекта {MEETING}", f"Не позднее {DEADLINE_PACK} (2 рабочих дня до заседания).", 7)
    people = [
        ("Председатель совета директоров", "ПСД", "оба дня", ""),
        ("Помощник ПСД", "аппарат ПСД", "оба дня", ""),
        ("Финансовый директор", "ГК", "оба дня", ""),
        ("Руководитель инспекционной группы", "ГК", "оба дня", ""),
    ]
    for name in COMPANIES_D1:
        people.append((f"Директор {name}", name, "день 1", ""))
        people.append((f"Главный бухгалтер {name}", name, "день 1", ""))
    for name in COMPANIES_D2:
        people.append((f"Директор {name}", name, "день 2", ""))
        people.append((f"Главный бухгалтер {name}", name, "день 2", ""))
    write_table(
        ws,
        ["Получатель", "Компания / роль", "День СД", "Канал (ЭДО / почта)", "Дата отправки", "Подтверждение", "Файл"],
        [[role, company, day, "почта", "", "", ""] for role, company, day, _ in people],
        [42, 36, 12, 20, 16, 16, 24],
    )

    ws = wb.create_sheet("График СД")
    banner(ws, "Распределение двух дней СД ГК", "п. 6.3: оба дня не позднее 25-го числа месяца, следующего за отчётным.", 4)
    write_table(
        ws,
        ["День", "Слот", "Компании", "Длительность"],
        [
            ["1", MEETING_DATES.split("–")[0].strip() + " сентября 2026", "; ".join(COMPANIES_D1[:-1]), "1,5 часа"],
            ["1", "тот же день", COMPANIES_D1[-1], "0,5 часа"],
            ["2", "25 сентября 2026", COMPANIES_D2[0], "0,5 часа"],
            ["2", "25 сентября 2026", "; ".join(COMPANIES_D2[1:]), "0,5 часа"],
        ],
        [10, 22, 80, 16],
    )


def setup_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(1.6)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(12)


def add_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = None
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT


def add_meta(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    for idx, (key, value) in enumerate(rows):
        table.cell(idx, 0).text = key
        table.cell(idx, 1).text = value


def write_agenda(path: Path) -> None:
    doc = Document()
    setup_doc(doc)
    add_heading(doc, f"Повестка заседания Совета директоров по Группе компаний {MEETING}")
    doc.add_paragraph(f"Основание: {SOURCE}. Заседание — два дня, оба не позднее 25-го числа месяца, следующего за отчётным.")
    add_meta(
        doc,
        [
            ("Номер", MEETING),
            ("Отчётный период", PERIOD),
            ("Даты", MEETING_DATES),
            ("Председатель", "Председатель совета директоров"),
            ("Ответственный за подготовку", "Помощник ПСД"),
            ("Участники", "директора компаний, главные бухгалтеры, финансовый директор, руководитель инспекционной группы; иные — по повестке"),
            ("Крайний срок рассылки", DEADLINE_PACK),
        ],
    )
    doc.add_paragraph()
    doc.add_paragraph("День 1. Финансовая и управленческая консолидация (доклад 10 мин + 10 мин вопросы).")
    for idx, item in enumerate(
        [
            "Проверка комплектности обязательного пакета п. 6.4 и приложений А–Б.",
            "Краткое резюме по ГК.",
            "Финансы план/факт: ОПУ, ДДС, ключевые отклонения; приложения А и Б.",
            "Задолженность и претензии свыше 1 000 000 ₽ от одного контрагента — доклады директоров и главных бухгалтеров.",
            "Доклад руководителя инспекционной группы о проверке управленческой отчётности с применением ИИ.",
            f"Доклады компаний дня 1 ({'; '.join(COMPANIES_D1)}).",
            "Капитал и инвестиции; вопросы сверх DoA.",
            "Статус поручений предыдущего СД (Action Tracker).",
        ],
        start=1,
    ):
        doc.add_paragraph(f"{idx}. {item}")
    doc.add_paragraph()
    doc.add_paragraph("День 2. Оставшиеся юридические лица и решения.")
    for idx, item in enumerate(
        [
            f"Доклады компаний дня 2 ({'; '.join(COMPANIES_D2)}).",
            "Риски и комплаенс (ТОП-10).",
            "Решения на согласование класса А.",
            "Фиксация решений в протоколе и Decision Log; оформление поручений в Action Tracker.",
        ],
        start=9,
    ):
        doc.add_paragraph(f"{idx}. {item}")
    doc.save(path)


def write_resume(path: Path) -> None:
    doc = Document()
    setup_doc(doc)
    add_heading(doc, f"Краткое резюме к {MEETING} (1 страница)")
    doc.add_paragraph(f"Отчётный период: {PERIOD}. Норма: пакет 1 п. 6.4 {SOURCE}.")
    add_meta(
        doc,
        [
            ("Главный вывод", ""),
            ("Выручка ГК факт / план / откл.", ""),
            ("Прибыль ГК факт / план / откл.", ""),
            ("Денежный поток", ""),
            ("ДЗ / КЗ / претензии > 1 млн ₽", ""),
            ("Инвестиции и сверхлимит DoA", ""),
            ("Критические риски", ""),
            ("Поручения прошлого СД: закрыто / открыто / просрочено", ""),
            ("Решения, которые нужно принять на этом СД", ""),
        ],
    )
    doc.add_paragraph()
    doc.add_paragraph("Текст резюме (не более одной страницы):")
    doc.add_paragraph("")
    doc.save(path)


def write_protocol(path: Path) -> None:
    doc = Document()
    setup_doc(doc)
    add_heading(doc, f"Протокол заседания Совета директоров по ГК {MEETING}")
    doc.add_paragraph(f"Готовит помощник ПСД. Решения дублируются в Decision Log в течение 1 рабочего дня. {SOURCE}, п. 5.2 и 6.5.")
    add_meta(
        doc,
        [
            ("Номер", MEETING),
            ("Даты", MEETING_DATES),
            ("Место / формат", ""),
            ("Председатель", "Председатель совета директоров"),
            ("Секретарь / помощник ПСД", ""),
            ("Присутствовали", ""),
            ("Отсутствовали / замещение", ""),
            ("Кворум", ""),
        ],
    )
    doc.add_paragraph()
    doc.add_paragraph("Повестка (кратко): комплектность; финансы; претензии > 1 млн ₽; доклад инспекции; доклады компаний; инвестиции; риски; решения класса А.")
    doc.add_paragraph()
    doc.add_paragraph("Решения (DEC)")
    table = doc.add_table(rows=2, cols=5)
    table.style = "Table Grid"
    for idx, name in enumerate(["DEC", "Формулировка", "Основание / DoA", "Голосование", "Поручения AT"]):
        table.cell(0, idx).text = name
    doc.add_paragraph()
    doc.add_paragraph("Поручения переносятся в Action Tracker: владелец, срок, критерий приёмки, статус.")
    doc.add_paragraph()
    doc.add_paragraph("Председатель _________________     Помощник ПСД _________________")
    doc.save(path)


def write_act(path: Path) -> None:
    doc = Document()
    setup_doc(doc)
    add_heading(doc, f"Акт проверки комплекта {MEETING}")
    doc.add_paragraph(
        "Отчёт агента «Подготовка заседаний Совета директоров». "
        "Агент материалы не утверждает. Проверка — по ПЛ-34-242 версия 03, п. 5.2–5.3, 6.1, 6.3–6.5."
    )
    add_meta(
        doc,
        [
            ("Заседание", MEETING),
            ("Отчётный период", PERIOD),
            ("Даты СД", MEETING_DATES),
            ("Крайний срок рассылки", DEADLINE_PACK),
            ("Проверяющий агент", "Подготовка заседаний Совета директоров"),
            ("Получатель", "Помощник Председателя совета директоров"),
            ("Итог комплекта", "полный / неполный"),
        ],
    )
    doc.add_paragraph()
    doc.add_paragraph("1. Комплектность обязательного пакета п. 6.4")
    doc.add_paragraph("Пункты 1–10, приложения А и Б, БДР/БДДС/ДЗ-КЗ, подписи трёх лиц, доклад инспекции — отметить по чек-листу. Отсутствующие файлы перечислить, не выдумывать.")
    doc.add_paragraph()
    doc.add_paragraph("2. Срок рассылки")
    doc.add_paragraph(f"Норма: не позднее чем за 2 рабочих дня. Для дат {MEETING_DATES} крайний срок {DEADLINE_PACK}. Факт рассылки — из реестра рассылки, не из предположения.")
    doc.add_paragraph()
    doc.add_paragraph("3. Задолженность и претензии свыше 1 000 000 ₽")
    doc.add_paragraph("Проверить, что директора и главные бухгалтеры каждой компании дня 1 и дня 2 представили сведения. Пустые строки — замечание в акте.")
    doc.add_paragraph()
    doc.add_paragraph("4. Превышение DoA")
    doc.add_paragraph("Инвестиции и решения класса А: если заявка выше лимита — вынести на СД, не утверждать.")
    doc.add_paragraph()
    doc.add_paragraph("5. Протокол, Decision Log, Action Tracker")
    doc.add_paragraph("До заседания — только черновики и статус старых поручений. Новые DEC/AT заполнять после расшифровки заседания. Без расшифровки решения не выдумывать.")
    doc.add_paragraph()
    doc.add_paragraph("Замечания агента")
    doc.add_paragraph("1.")
    doc.add_paragraph("2.")
    doc.add_paragraph()
    doc.add_paragraph("Передать помощнику ПСД: чек-лист, резюме, пакет 1–10, приложения А–Б, реестр претензий, реестр рассылки, акт.")
    doc.save(path)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    COPY.mkdir(parents=True, exist_ok=True)

    wb = new_wb()
    add_checklist(wb)
    add_pack(wb)
    add_registers(wb)
    xlsx = OUT / f"{MEETING}_комплект_и_формы.xlsx"
    wb.save(xlsx)

    write_agenda(OUT / f"{MEETING}_повестка.docx")
    write_resume(OUT / f"{MEETING}_краткое_резюме.docx")
    write_protocol(OUT / f"{MEETING}_протокол.docx")
    write_act(OUT / f"{MEETING}_акт_проверки_агента.docx")

    for item in OUT.iterdir():
        target = COPY / item.name
        target.write_bytes(item.read_bytes())
    print(OUT)
    for item in sorted(OUT.iterdir()):
        print(item.name, item.stat().st_size)


if __name__ == "__main__":
    main()
