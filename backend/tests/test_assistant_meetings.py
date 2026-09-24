from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from kpi.sources.assistant_meetings import (
    SHEET_HEADERS,
    compute_meetings_schedule_kpi,
    score_meetings_schedule_kpi,
    violation_score,
    write_meetings_report_xlsx,
)

AS_OF = date(2026, 9, 22)
START = date(2026, 9, 1)
END = date(2026, 9, 30)


def _memo(**extra):
    row = {
        "Number": "СЗ-1",
        "DeletionMark": False,
        "Статус": "Согласована",
        "ТемаСлужебнойЗаписки_Name": "Организация совещаний (регл.)",
        "ТемаСовещания": "Планёрка директора",
        "ДатаПроведенияСовещания": "2026-09-10T00:00:00",
    }
    row.update(extra)
    return row


def _event(subject: str, start: str) -> dict:
    return {"role": "fact", "subject": subject, "start": start}


def test_violation_bands():
    assert violation_score(0) == 100
    assert violation_score(4) == 100
    assert violation_score(5) == 50
    assert violation_score(8) == 50
    assert violation_score(9) == 0


def test_approved_memos_must_be_in_meetings_calendar():
    rows = [
        {"role": "plan", **_memo(Number="СЗ-1", ТемаСовещания="Планёрка директора")},
        {
            "role": "plan",
            **_memo(
                Number="СЗ-2",
                ТемаСовещания="Закупки сентября",
                ДатаПроведенияСовещания="2026-09-12T09:00:00",
            ),
        },
        {
            "role": "plan",
            **_memo(Number="СЗ-3", Статус="На согласовании", ТемаСовещания="Черновик"),
        },
        {
            "role": "plan",
            **_memo(
                Number="СЗ-4",
                ТемаСлужебнойЗаписки_Name="Прочее",
                ТемаСовещания="Не эта тема",
            ),
        },
        _event("Планёрка директора", "2026-09-10T10:00:00"),
    ]
    report = score_meetings_schedule_kpi(rows, as_of=AS_OF, date_from=START, date_to=END)
    assert report["plan_total"] == 2
    assert report["in_calendar"] == 1
    assert report["violations"] == 1
    assert report["fact_pct"] == 50.0
    assert report["score_pct"] == 100.0
    assert report["contrib_pct"] == 40.0
    missing = [row for row in report["rows"] if not row["in_calendar"]]
    assert missing[0]["number"] == "СЗ-2"


def test_five_missing_memos_score_half():
    rows = [
        {"role": "plan", **_memo(Number=f"СЗ-{index}", ТемаСовещания=f"Тема {index}")}
        for index in range(5)
    ]
    report = score_meetings_schedule_kpi(rows, as_of=AS_OF, date_from=START, date_to=END)
    assert report["violations"] == 5
    assert report["score_pct"] == 50.0
    assert report["contrib_pct"] == 20.0


def test_nine_missing_memos_score_zero():
    rows = [
        {"role": "plan", **_memo(Number=f"СЗ-{index}", ТемаСовещания=f"Тема {index}")}
        for index in range(9)
    ]
    report = score_meetings_schedule_kpi(rows, as_of=AS_OF, date_from=START, date_to=END)
    assert report["score_pct"] == 0.0
    assert report["contrib_pct"] == 0.0


def test_same_title_on_another_day_still_counts():
    rows = [
        {"role": "plan", **_memo()},
        _event("Планёрка директора (переговорная)", "2026-09-11T10:00:00"),
    ]
    report = score_meetings_schedule_kpi(rows, as_of=AS_OF, date_from=START, date_to=END)
    assert report["in_calendar"] == 1
    assert report["rows"][0]["event_date"] == "2026-09-11"


def test_catalog_ref_counts_as_meeting_theme():
    rows = [
        {
            "role": "plan",
            **_memo(
                Number="СЗ-9",
                ТемаСлужебнойЗаписки="cad8df76-73cc-11ea-8341-ac1f6b05524d",
                ТемаСлужебнойЗаписки_Name="",
            ),
        },
        _event("Планёрка директора", "2026-09-10T10:00:00"),
    ]
    report = score_meetings_schedule_kpi(rows, as_of=AS_OF, date_from=START, date_to=END)
    assert report["plan_total"] == 1
    assert report["in_calendar"] == 1


def test_memo_outside_month_is_not_plan():
    rows = [
        {
            "role": "plan",
            **_memo(ДатаПроведенияСовещания="2026-08-31T00:00:00"),
        }
    ]
    report = score_meetings_schedule_kpi(rows, as_of=AS_OF, date_from=START, date_to=END)
    assert report["plan_total"] == 0
    assert report["fact_pct"] is None
    assert report["score_pct"] is None


def test_compute_reads_dontsova_plan_fact():
    class _Ctx:
        date_from = START
        date_to = END

        def load_for(self, extra):
            entity = extra.get("entity")
            if entity == "Catalog_Пользователи":
                return [{"Ref_Key": "leader", "Description": "Донцова Анна Егоровна"}]
            if entity == "Catalog_ТД_ТемыСовещаний":
                return [{"Ref_Key": "t1", "Description": "Еженедельное совещание"}]
            assert entity == "Document_ТД_Протокол"
            return [
                {
                    "ТемаСовещания_Key": "t1",
                    "Date": "2026-09-04T10:00:00",
                    "DeletionMark": False,
                    "ВидСовещания": "Отчетное",
                },
                {
                    "ТемаСовещания_Key": "t1",
                    "Date": "2026-09-11T10:00:00",
                    "DeletionMark": False,
                    "ВидСовещания": "Отчетное",
                },
                {
                    "ТемаСовещания_Key": "t1",
                    "Date": "2026-09-18T10:00:00",
                    "DeletionMark": False,
                    "ВидСовещания": "Внеплановое",
                },
            ]

    report = compute_meetings_schedule_kpi(_Ctx(), as_of=AS_OF)
    assert report["plan_total"] == 2
    assert report["fact_total"] == 2
    assert report["violations"] == 0
    assert report["fact_pct"] == 100.0
    assert report["score_pct"] == 100.0
    assert report["rows"][0]["name"] == "Еженедельное совещание"
    assert report["contrib_pct"] == 40.0


def test_excel_sheet_uses_unplanned_columns(tmp_path: Path) -> None:
    rows = [
        {"role": "plan", **_memo()},
        {
            "role": "plan",
            **_memo(
                Number="СЗ-2",
                ТемаСовещания="Закупки сентября",
                ДатаПроведенияСовещания="2026-09-12T09:00:00",
            ),
        },
        _event("Планёрка директора", "2026-09-11T10:00:00"),
    ]
    report = score_meetings_schedule_kpi(rows, as_of=AS_OF, date_from=START, date_to=END)
    assert report["sheet_headers"] == list(SHEET_HEADERS)
    assert report["sheet_rows"] == [
        ["Планёрка директора", 1, "2026-09-10", "2026-09-11"],
        ["Закупки сентября", 1, "2026-09-12", ""],
    ]
    path = write_meetings_report_xlsx(report, tmp_path / "sept.xlsx")
    sheet = load_workbook(path)["сентябрь 2026"]
    assert [sheet.cell(1, col).value for col in range(1, 5)] == list(SHEET_HEADERS)
    assert sheet["A2"].value == "Планёрка директора"
    assert sheet["B2"].value == 1
    assert sheet["C2"].value.date() == date(2026, 9, 10)
    assert sheet["D2"].value.date() == date(2026, 9, 11)
    assert sheet["D3"].value is None
