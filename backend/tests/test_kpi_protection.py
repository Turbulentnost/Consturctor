from __future__ import annotations

from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.services.position_kpi import protection
from app.services.position_kpi.bonus_form import BonusMoney, build_bonus_form_xlsx


@pytest.fixture(autouse=True)
def _store(tmp_path, monkeypatch):
    monkeypatch.setenv("KPI_PROTECTION_PATH", str(tmp_path / "kpi.json"))
    protection._attempts.clear()


def test_enable_requires_password():
    with pytest.raises(protection.KpiProtectionError):
        protection.update_protection("u1", enabled=True, password=None, current_password=None)
    state = protection.update_protection("u1", enabled=True, password="secret1", current_password=None)
    assert state.enabled and state.has_password


def test_password_not_stored_in_plain(tmp_path):
    protection.update_protection("u1", enabled=True, password="secret1", current_password=None)
    assert "secret1" not in (tmp_path / "kpi.json").read_text(encoding="utf-8")


def test_changes_need_current_password():
    protection.update_protection("u1", enabled=True, password="secret1", current_password=None)
    with pytest.raises(protection.KpiProtectionError) as exc:
        protection.update_protection("u1", enabled=False, password=None, current_password="wrong")
    assert exc.value.status_code == 403
    state = protection.update_protection("u1", enabled=False, password=None, current_password="secret1")
    assert not state.enabled and state.has_password


def test_unlock_token_bound_to_user_and_password():
    protection.update_protection("u1", enabled=True, password="secret1", current_password=None)
    token, _ = protection.unlock("u1", "secret1")
    assert protection.token_unlocks("u1", token)
    assert not protection.token_unlocks("u2", token)
    assert not protection.token_unlocks("u1", token[:-2] + "xx")
    protection.update_protection("u1", enabled=True, password="secret2", current_password="secret1")
    assert not protection.token_unlocks("u1", token)


def test_unlock_rejected_when_disabled_or_wrong():
    with pytest.raises(protection.KpiProtectionError):
        protection.unlock("u1", "any")
    protection.update_protection("u1", enabled=True, password="secret1", current_password=None)
    with pytest.raises(protection.KpiProtectionError):
        protection.unlock("u1", "nope")


def test_rate_limit_after_failures():
    protection.update_protection("u1", enabled=True, password="secret1", current_password=None)
    for _ in range(protection._MAX_ATTEMPTS):
        with pytest.raises(protection.KpiProtectionError):
            protection.unlock("u1", "nope")
    with pytest.raises(protection.KpiProtectionError) as exc:
        protection.unlock("u1", "secret1")
    assert exc.value.status_code == 429


def _rows():
    return [
        {"code": "a", "name": "Цель А", "weight": 60, "evidence": "", "earned": 60.0, "detail": {}},
        {"code": "b", "name": "Цель Б", "weight": 40, "evidence": "", "earned": 20.0, "detail": {}},
    ]


def test_bonus_form_money_columns():
    content = build_bonus_form_xlsx(
        fio="Иванов Иван Иванович",
        position="Инженер",
        period_from=date(2026, 5, 1),
        period_to=date(2026, 5, 31),
        rows=_rows(),
        money=BonusMoney(base=10000.0, oklad=50000.0, basis="Базовая премия из 1С:ЗУП"),
    )
    sheet = load_workbook(BytesIO(content))["ИЦПП"]
    assert sheet["H12"].value == "Премия, ₽"
    assert sheet["H13"].value == 6000.0
    assert sheet["H14"].value == 2000.0
    assert sheet["H15"].value == 8000.0
    assert sheet["H4"].value == 50000.0


def test_bonus_form_without_money_has_no_rubles():
    content = build_bonus_form_xlsx(
        fio="Иванов Иван Иванович",
        position="Инженер",
        period_from=date(2026, 5, 1),
        period_to=date(2026, 5, 31),
        rows=_rows(),
    )
    sheet = load_workbook(BytesIO(content))["ИЦПП"]
    assert sheet["H12"].value is None
    assert sheet["G15"].value == "80%"
