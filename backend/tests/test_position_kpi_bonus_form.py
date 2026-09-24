from __future__ import annotations

from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker

from app.clients.erp_sql import ErpSqlError
from app.db.base import Base
from app.services.position_kpi.bonus_form import (
    ReportSubjectError,
    build_bonus_form_xlsx,
    render_bonus_form,
    resolve_report_subject,
)
from kpi.seed_pl_npo_010 import upsert_catalog

PSD = "Помощник Председателя совета директоров"


def _session() -> Session:
    engine = create_engine("sqlite:///:memory:", future=True)
    Base.metadata.create_all(bind=engine)
    db = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)()
    upsert_catalog(db)
    db.commit()
    return db


def _sheet_text(content: bytes) -> str:
    workbook = load_workbook(BytesIO(content))
    sheet = workbook.active
    return "\n".join(str(cell.value or "") for row in sheet.iter_rows() for cell in row)


def test_ilchenko_resolves_without_erp(monkeypatch) -> None:
    monkeypatch.setattr("app.services.app_users.find_app_user_by_fio", lambda _fio: None)

    def boom(_fio: str):
        raise AssertionError("1С не нужна, должность Ильченко уже известна")

    monkeypatch.setattr("app.clients.erp_sql.get_user_profile_by_fio", boom)
    subject = resolve_report_subject(fio="Ильченко Екатерина Александровна")
    assert subject["fio"] == "Ильченко Екатерина Александровна"
    assert subject["position"] == PSD


def test_explicit_position_for_any_employee(monkeypatch) -> None:
    monkeypatch.setattr("app.services.app_users.find_app_user_by_fio", lambda _fio: None)

    def boom(_fio: str):
        raise AssertionError("должность передана явно")

    monkeypatch.setattr("app.clients.erp_sql.get_user_profile_by_fio", boom)
    subject = resolve_report_subject(fio="Смирнов Пётр Иванович", position="Офис-менеджер")
    assert subject["position"] == "Офис-менеджер"


def test_unknown_employee_without_position(monkeypatch) -> None:
    monkeypatch.setattr("app.services.app_users.find_app_user_by_fio", lambda _fio: None)

    def down(_fio: str):
        raise ErpSqlError("sql down")

    monkeypatch.setattr("app.clients.erp_sql.get_user_profile_by_fio", down)
    with pytest.raises(ReportSubjectError) as exc:
        resolve_report_subject(fio="Смирнов Пётр Иванович")
    assert exc.value.status_code == 404


def test_form_uses_selected_employee_not_azanova() -> None:
    content = build_bonus_form_xlsx(
        fio="Ильченко Екатерина Александровна",
        position=PSD,
        period_from=date(2026, 8, 1),
        period_to=date(2026, 8, 31),
        rows=[
            {
                "name": "Своевременность пакета к заседаниям (СД + РК)",
                "weight": 25,
                "evidence": "пакет вовремя",
                "earned": 25,
            }
        ],
    )
    text = _sheet_text(content)
    assert "Ильченко Екатерина Александровна" in text
    assert PSD in text
    assert "август 2026 г." in text
    assert "пакет вовремя" in text
    assert "25%" in text
    assert "Азанова" not in text


def test_report_column_links_to_detail_sheet() -> None:
    content = build_bonus_form_xlsx(
        fio="Ильченко Екатерина Александровна",
        position=PSD,
        period_from=date(2026, 9, 1),
        period_to=date(2026, 9, 30),
        rows=[
            {
                "code": "package_on_time",
                "name": "Своевременность пакета к заседаниям (СД + РК)",
                "weight": 25,
                "evidence": "Zвовремя / Zвсего = 4/4",
                "earned": 25,
                "detail": {
                    "z_total": 4,
                    "z_on_time": 4,
                    "fact_pct": 100,
                    "score_pct": 100,
                    "rows": [
                        {
                            "kind_label": "СД",
                            "subject": "Совет директоров 12.09",
                            "plan_date": "2026-09-12",
                            "deadline": "2026-09-10",
                            "due": True,
                            "protocol_number": "ПСД-14",
                            "on_time": True,
                        }
                    ],
                },
            }
        ],
    )
    workbook = load_workbook(BytesIO(content))
    cover = workbook["ИЦПП"]
    link = cover["D13"].hyperlink
    assert link is not None
    assert link.location
    detail_name = link.location.split("!")[0].strip("'")
    assert cover["D13"].value == detail_name
    detail = workbook[detail_name]
    blob = "\n".join(str(cell.value or "") for row in detail.iter_rows() for cell in row)
    assert "Совет директоров 12.09" in blob
    assert "ПСД-14" in blob
    assert "Zвсего" in blob
    assert "4" in blob
    assert detail["A2"].value == "К форме"
    assert detail["A2"].hyperlink is not None
    assert detail["A2"].hyperlink.location.startswith("'ИЦПП'")


def test_dpi_detail_sheet_shows_plan_and_fact() -> None:
    content = build_bonus_form_xlsx(
        fio="Комарькова",
        position="Помощник руководителя",
        period_from=date(2026, 9, 1),
        period_to=date(2026, 9, 30),
        rows=[
            {
                "code": "dpi_appointment",
                "name": "Назначение ДПИ за месяц",
                "weight": 10,
                "evidence": "В календаре 16/16, факт 100%",
                "earned": 10,
                "detail": {
                    "plan_total": 16,
                    "in_calendar": 16,
                    "fact_pct": 100,
                    "score_pct": 100,
                    "rows": [
                        {
                            "number": 1,
                            "name": "ДПИ СУП",
                            "rule": "5 числа каждого месяца",
                            "plan_date": "2026-09-07",
                            "fact_date": "2026-09-08",
                            "appointed": True,
                        }
                    ],
                },
            }
        ],
    )
    workbook = load_workbook(BytesIO(content))
    detail = workbook[workbook.sheetnames[1]]
    blob = "\n".join(str(cell.value or "") for row in detail.iter_rows() for cell in row)
    assert "План, шт" in blob
    assert "В календаре, шт" in blob
    assert "ДПИ СУП" in blob
    assert "5 числа каждого месяца" in blob
    assert "2026-09-07" in blob
    assert "2026-09-08" in blob
    assert "да" in blob


def test_instructions_fact_cell_stays_empty() -> None:
    content = build_bonus_form_xlsx(
        fio="Ильченко Екатерина Александровна",
        position=PSD,
        period_from=date(2026, 9, 1),
        period_to=date(2026, 9, 30),
        rows=[
            {
                "code": "instructions",
                "name": "Реестр и контроль исполнения поручений (СД + РК)",
                "weight": 25,
                "evidence": (
                    "2026-09-01 — 2026-09-30: 1С 31 позиций, в Excel 31. "
                    "Позиции 1С считаем внесёнными в файл. Файл: C:\\ActionTracker.xlsx"
                ),
                "earned": 25,
                "detail": {
                    "r_total": 31,
                    "r24": 31,
                    "fact_pct": 100,
                    "score_pct": 100,
                    "rows": [],
                },
            }
        ],
    )
    workbook = load_workbook(BytesIO(content))
    cover = workbook["ИЦПП"]
    assert cover["F13"].value in (None, "")
    detail = workbook[workbook.sheetnames[1]]
    blob = "\n".join(str(cell.value or "") for row in detail.iter_rows() for cell in row)
    assert "ActionTracker" not in blob
    assert "считаем внесёнными" not in blob
    assert "Rвсего" in blob


def test_orders_detail_sheet_is_not_a_meeting() -> None:
    content = build_bonus_form_xlsx(
        fio="Акинина",
        position="Помощник руководителя",
        period_from=date(2026, 9, 1),
        period_to=date(2026, 9, 30),
        rows=[
            {
                "code": "orders_registration",
                "name": "Регистрация приказов и распоряжений",
                "weight": 10,
                "evidence": "Приказов 1, распоряжений 0, всего 1.",
                "earned": 10,
                "detail": {
                    "fact_pct": 100,
                    "score_pct": 100,
                    "rows": [
                        {
                            "kind": "order",
                            "doc_kind": "Приказ",
                            "number": "НП00-000322",
                            "date": "2026-09-01",
                            "posted": True,
                            "content": "О назначении ответственного",
                            "ref": "30b21e5c",
                        }
                    ],
                },
            }
        ],
    )
    workbook = load_workbook(BytesIO(content))
    detail = workbook[workbook.sheetnames[1]]
    headers = [cell.value for cell in next(detail.iter_rows(min_row=7, max_row=7))]
    assert headers[:5] == ["Номер", "Вид", "Содержание", "Дата", "Проведён"]
    assert "Совещание" not in headers
    assert "30b21e5c" not in [cell.value for row in detail.iter_rows() for cell in row]


def test_render_lists_position_goals_for_ilchenko(monkeypatch) -> None:
    db = _session()
    monkeypatch.setattr("app.services.app_users.find_app_user_by_fio", lambda _fio: None)
    monkeypatch.setattr(
        "app.services.position_kpi.bonus_form.get_or_compute_position_kpi",
        lambda *_args, **_kwargs: {
            "tiles": [
                {
                    "code": "package_on_time",
                    "score": 100,
                    "contrib": 25,
                    "evidence": "пакет вовремя",
                }
            ]
        },
    )
    content, filename = render_bonus_form(
        db,
        fio="Ильченко Екатерина Александровна",
        position="",
        date_from=date(2026, 8, 1),
        date_to=date(2026, 8, 31),
    )
    assert filename == "ИЦПП_Ильченко_2026-08.xlsx"
    text = _sheet_text(content)
    assert "Своевременность пакета к заседаниям (СД + РК)" in text
    assert "Своевременность протоколов (СД + РК)" in text
    assert "Реестр и контроль исполнения поручений (СД + РК)" in text
    assert "Качество протокола и материалов" in text
    assert "пакет вовремя" in text
    assert "По части целей факт за период ещё не посчитан" in text
    assert "100%" in text
