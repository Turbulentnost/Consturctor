from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from app.services.position_kpi.report_fact import (
    ReportFactMissing,
    breakdown_for_metric,
    find_report,
    included_flags,
    load_metric_fact,
)


def _book() -> Workbook:
    book = Workbook()
    cover = book.active
    cover.title = "ИЦПП"
    cover["C5"] = "Ильченко Екатерина Александровна"
    sheet = book.create_sheet("3. Реестр")
    sheet["A1"] = "Реестр и контроль исполнения поручений (СД + РК)"
    sheet["B2"] = "01.10.2026 — 31.10.2026"
    for column, label in enumerate(("Вес", "План", "Факт", "Оценка", "Вклад"), start=1):
        sheet.cell(4, column, label)
    for column, value in enumerate(("25%", 95, 19.35, "20.37%", "5.09%"), start=1):
        sheet.cell(5, column, value)
    sheet["A7"] = "KPI3.1 10/40 = 25.0%. KPI3.2 6/31 = 19.35%. min 19.35% → оценка 20.37%."
    for column, header in enumerate(("Орган", "В Excel", "Вовремя"), start=1):
        sheet.cell(9, column, header)
    sheet.cell(10, 1, "АСТ")
    sheet.cell(10, 2, "нет")
    sheet.cell(10, 3, "нет")
    sheet.cell(11, 1, "ПСД")
    sheet.cell(11, 2, "да")
    sheet.cell(11, 3, "да")
    return book


def test_breakdown_reads_fact_and_rows() -> None:
    found = breakdown_for_metric(_book(), "Реестр и контроль исполнения поручений (СД + РК)")
    assert found is not None
    assert found["fact_text"].startswith("KPI3.1")
    assert found["summary"] == [
        {"label": "Вес", "value": "25%"},
        {"label": "План", "value": "95%"},
        {"label": "Факт", "value": "19.35%"},
        {"label": "Оценка", "value": "20.37%"},
        {"label": "Вклад", "value": "5.09%"},
    ]
    assert found["columns"] == ["Орган", "В Excel", "Вовремя"]
    assert found["rows"] == [["АСТ", "нет", "нет"], ["ПСД", "да", "да"]]
    assert found["period"] == "01.10.2026 — 31.10.2026"


def test_included_flags_mark_rows_outside_the_fact() -> None:
    assert included_flags(
        ["Срок наступил", "Вовремя"],
        [["да", "да"], ["да", "нет"], ["нет", "нет"]],
    ) == [True, False, True]
    assert included_flags(["Возврат"], [["нет"], ["да"]]) == [True, False]
    assert included_flags(
        ["В периоде", "Вовремя", "Активное", "Контроль"],
        [
            ["да", "да", "нет", "нет"],
            ["да", "нет", "да", "нет"],
            ["нет", "нет", "да", "да"],
        ],
    ) == [True, False, True]
    assert included_flags(
        ["В периоде", "Вовремя", "Активное", "Контроль"],
        [
            ["да", "да", "нет", "нет"],
            ["да", "нет", "нет", "нет"],
            ["нет", "нет", "да", "да"],
        ],
    ) == [True, False, True]


def test_find_report_matches_cover_fio(tmp_path: Path) -> None:
    book = _book()
    target = tmp_path / "ИЦПП_Ильченко_2026-10.xlsx"
    book.save(target)
    other = Workbook()
    other.active["C5"] = "Акинина Татьяна Владимировна"
    other.save(tmp_path / "ИЦПП_Акинина_2026-10.xlsx")

    found = find_report("Ильченко Екатерина Александровна", "2026-10", tmp_path)
    assert found == target
    assert find_report("Акинина Татьяна Владимировна", "2026-10", tmp_path) is not None
    assert find_report("Ильченко Екатерина Александровна", "2026-09", tmp_path) is None


def test_load_metric_fact_missing_sheet(tmp_path: Path, monkeypatch) -> None:
    book = _book()
    book.save(tmp_path / "ИЦПП_Ильченко_2026-10.xlsx")
    monkeypatch.setattr("app.services.position_kpi.report_fact.report_dir", lambda: tmp_path)
    try:
        load_metric_fact("Ильченко Екатерина Александровна", "2026-10", "Качество протокола")
    except ReportFactMissing as exc:
        assert "нет листа" in exc.message
    else:
        raise AssertionError("expected missing sheet")
