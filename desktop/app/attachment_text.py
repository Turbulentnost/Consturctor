"""Извлечение текста из вложений: офисные файлы, PDF и OCR для сканов/фото."""

from __future__ import annotations

from pathlib import Path

_TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".xml", ".html", ".htm", ".log"}
_IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff"}
_MAX_CHARS = 12_000
UNREADABLE = "текст не извлечён"
_UNREADABLE = UNREADABLE


def extract_attachment_text(
    path: str,
    *,
    max_chars: int = _MAX_CHARS,
    max_pages: int = 0,
    start_page: int = 1,
    ocr: bool = True,
) -> str:
    file_path = Path(path)
    name = file_path.name
    suffix = file_path.suffix.lower()
    if not file_path.is_file():
        return f"файл {name} прикреплён, {_UNREADABLE} (файл не найден)"
    try:
        if suffix in _TEXT_SUFFIXES:
            text = file_path.read_text(encoding="utf-8", errors="replace")
        elif suffix == ".pdf":
            text = _read_pdf(
                file_path, max_pages=max_pages, start_page=start_page, ocr=ocr
            )
        elif suffix == ".docx":
            text = _read_docx(file_path)
        elif suffix in {".xlsx", ".xlsm"}:
            text = _read_xlsx(file_path)
        elif suffix in _IMAGE_SUFFIXES:
            text = _ocr(file_path) if ocr else ""
        elif suffix == ".doc":
            text = _read_doc(file_path)
        else:
            text = ""
    except Exception:  # noqa: BLE001
        text = ""
    text = (text or "").strip()
    if not text:
        return f"файл {name} прикреплён, {_UNREADABLE}"
    limit = max(200, int(max_chars or _MAX_CHARS))
    if len(text) > limit:
        text = text[:limit].rstrip() + "\n…"
    return text


def write_extracted_sidecar(path: str | Path) -> str:
    """Write path + '.txt' when OCR/native extract produced real text.

    Returns the sidecar path as a posix string, or empty if nothing useful.
    """
    source = Path(path)
    if not source.is_file():
        return ""
    text = extract_attachment_text(str(source))
    if not text or _UNREADABLE in text:
        return ""
    target = source.with_name(source.name + ".txt")
    target.write_text(text, encoding="utf-8")
    return target.as_posix()


def format_attachments_block(paths: list[str]) -> str:
    if not paths:
        return ""
    parts: list[str] = ["", "Прикреплено:"]
    for path in paths:
        name = Path(path).name
        extracted = extract_attachment_text(path)
        parts.append(f"— {name}")
        parts.append(extracted)
    return "\n".join(parts).rstrip()


def _ocr(path: Path) -> str:
    try:
        from app.ocr_client import ocr_file
    except ImportError:
        return ""
    return ocr_file(path)


def _read_pdf(
    path: Path, *, max_pages: int = 0, start_page: int = 1, ocr: bool = True
) -> str:
    native = ""
    try:
        import fitz  # pymupdf
    except ImportError:
        fitz = None
    if fitz is not None:
        doc = fitz.open(path)
        try:
            pages = list(doc)
            start = max(1, int(start_page or 1))
            if start > len(pages):
                start = 1
            pages = pages[start - 1 :]
            if max_pages and max_pages > 0:
                pages = pages[:max_pages]
            native = "\n\n".join((page.get_text() or "") for page in pages)
        finally:
            doc.close()
    native = (native or "").strip()
    if native:
        return native
    return _ocr(path) if ocr else ""


def _read_docx(path: Path) -> str:
    try:
        import docx
    except ImportError:
        return ""
    document = docx.Document(str(path))
    return "\n".join(para.text for para in document.paragraphs)


def _read_xlsx(path: Path) -> str:
    try:
        import openpyxl
    except ImportError:
        return ""
    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        parts: list[str] = []
        for sheet in workbook.worksheets:
            parts.append(f"===== SHEET: {sheet.title} =====")
            for row in sheet.iter_rows(values_only=True):
                values = [str(cell).strip() if cell is not None else "" for cell in row]
                while values and not values[-1]:
                    values.pop()
                if values:
                    parts.append("\t".join(values))
        return "\n".join(parts)
    finally:
        workbook.close()


def _read_doc(path: Path) -> str:
    # Старый .doc без внешних утилит обычно не разбирается.
    return ""
