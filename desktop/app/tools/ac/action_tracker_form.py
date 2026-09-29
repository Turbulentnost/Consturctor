"""Собрать журнал поручений по форме со скриншота (OCR), 1С и Action Tracker.

Колонки и заголовок берутся из OCR скриншота или из уже лежащего журнала.
Строки не выдумываются: только поручения 1С и строки трекера.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff"}
_AST_RE = re.compile(r"АСТ\d{0,2}-?\d+", re.IGNORECASE)
_ACT_ID_RE = re.compile(r"ACT-(\d+)", re.IGNORECASE)
_PLACEHOLDER_OCR = re.compile(r"^\[изображение:", re.IGNORECASE)

_THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
_HEADER_FILL = PatternFill("solid", fgColor="D6DCE4")
_CRIT_FILL = PatternFill("solid", fgColor="FF6B6B")
_HIGH_FILL = PatternFill("solid", fgColor="F4B183")
_CLOSED_FILL = PatternFill("solid", fgColor="92D050")
_PROGRESS_FILL = PatternFill("solid", fgColor="9BC2E6")
_WRAP = Alignment(wrap_text=True, vertical="center", horizontal="left")
_CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")


@dataclass
class FormSpec:
    title: str = ""
    subtitle: str = ""
    headers: list[str] = field(default_factory=list)
    source: str = ""


def is_placeholder_extract(text: str) -> bool:
    raw = (text or "").strip()
    return not raw or bool(_PLACEHOLDER_OCR.match(raw))


def parse_form_from_ocr(text: str) -> FormSpec:
    """Вытащить шапку и колонки из текста OCR, без эталонного списка полей."""
    lines = [line.strip() for line in (text or "").splitlines() if line.strip()]
    title = ""
    subtitle = ""
    headers: list[str] = []
    for line in lines:
        folded = line.casefold().replace("ё", "е")
        compact = re.sub(r"[\s_\-]+", "", folded)
        if "приложен" in folded and not title:
            title = re.sub(r"\s+", " ", line).strip()
            continue
        if "actiontracker" in compact:
            if not subtitle:
                subtitle = re.sub(r"\s+", " ", line).strip()
            parts = _split_header_line(line)
            if _looks_like_headers(parts) and not headers:
                headers = parts
            continue
        parts = _split_header_line(line)
        if _looks_like_headers(parts) and not headers:
            headers = parts
    return FormSpec(title=title, subtitle=subtitle, headers=headers, source="ocr")


def headers_from_tracker_rows(rows: list[list]) -> list[str]:
    if not rows or not isinstance(rows[0], (list, tuple)):
        return []
    headers = [str(cell or "").strip() for cell in rows[0]]
    if _looks_like_headers(headers):
        return [item for item in headers if item]
    return []


def resolve_form_spec(*, ocr_text: str = "", tracker_rows: list[list] | None = None) -> FormSpec:
    spec = parse_form_from_ocr(ocr_text)
    tracker_headers = headers_from_tracker_rows(tracker_rows or [])
    if not spec.headers and tracker_headers:
        spec.headers = tracker_headers
        spec.source = "tracker" if not spec.source else f"{spec.source}+tracker"
    return spec


def find_form_screenshot(root: Path) -> Path | None:
    if not root.is_dir():
        return None
    candidates: list[Path] = []
    for path in root.rglob("*"):
        if not path.is_file() or path.suffix.lower() not in _IMAGE_SUFFIXES:
            continue
        relative = path.relative_to(root).as_posix()
        if relative.startswith("code/") or relative.startswith("tool_results/"):
            continue
        candidates.append(path)
    if not candidates:
        return None

    def score(path: Path) -> tuple[int, int, float]:
        text = _sidecar_text(path)
        folded = text.casefold().replace("ё", "е")
        formish = int(
            "action tracker" in folded
            or "приложен" in folded
            or "поручен" in folded
            or "actiontracker" in re.sub(r"[\s_\-]+", "", folded)
        )
        in_materials = int("materials" in path.relative_to(root).as_posix().lower())
        return (formish, in_materials, path.stat().st_mtime)

    return max(candidates, key=score)


def ocr_form_text(path: Path) -> str:
    sidecar = Path(str(path) + ".txt")
    if sidecar.is_file():
        text = sidecar.read_text(encoding="utf-8")
        if not is_placeholder_extract(text):
            return text
    try:
        from app.attachment_text import extract_attachment_text

        text = extract_attachment_text(str(path))
    except Exception:  # noqa: BLE001
        return ""
    if is_placeholder_extract(text):
        return ""
    sidecar.write_text(text, encoding="utf-8")
    return text


def _is_tracker_name(name: str) -> bool:
    folded = re.sub(r"[\s_\-]+", "", (name or "").casefold())
    return "actiontracker" in folded


def find_tracker_path(root: Path, requested: str = "") -> Path:
    name = (requested or "").strip()
    if name:
        candidate = (root / name).resolve()
        if root.resolve() in candidate.parents or candidate == root.resolve():
            return candidate
    for path in root.iterdir() if root.is_dir() else []:
        if path.is_file() and _is_tracker_name(path.name):
            return path
    return root / "ActionTracker.xlsx"


def merge_form_rows(
    headers: list[str],
    *,
    tracker_tasks: list[dict],
    assignments: list[dict],
) -> list[list]:
    """Свести строки трекера и карточки 1С в колонки формы. Статус трекера важнее."""
    roles = [_header_role(header) for header in headers]
    by_key: dict[str, dict[str, str]] = {}
    order: list[str] = []

    for task in tracker_tasks:
        values = _row_from_mapping(headers, roles, task)
        key = _row_key(values, roles) or f"tracker:{len(order)}"
        if key not in by_key:
            order.append(key)
            by_key[key] = values
        else:
            by_key[key] = _prefer_filled(by_key[key], values)

    next_id = _next_act_id(by_key.values(), roles)
    for assignment in assignments:
        incoming = _row_from_assignment(headers, roles, assignment)
        key = _row_key(incoming, roles) or _assignment_number(assignment)
        if key and key in by_key:
            by_key[key] = _overlay_assignment(by_key[key], incoming, roles)
            continue
        id_idx = _role_index(roles, "id")
        if id_idx is not None and not incoming.get(id_idx):
            incoming[id_idx] = f"ACT-{next_id:04d}"
            next_id += 1
        new_key = key or (incoming.get(id_idx) if id_idx is not None else "") or f"new:{len(order)}"
        order.append(str(new_key))
        by_key[str(new_key)] = incoming

    rows: list[list] = []
    for key in order:
        values = by_key[key]
        rows.append([values.get(idx, "") for idx in range(len(headers))])
    return rows


def write_form(path: Path, spec: FormSpec, rows: list[list]) -> Path:
    headers = spec.headers
    if not headers:
        raise ValueError("Нет колонок формы: нужен OCR скриншота или шапка Action Tracker.")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Action Tracker"[:31]
    cursor = 1
    last_col = get_column_letter(len(headers))
    if spec.title:
        sheet.merge_cells(f"A{cursor}:{last_col}{cursor}")
        cell = sheet.cell(cursor, 1, spec.title)
        cell.font = Font(name="Calibri", size=16, bold=True, color="1F4E79")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[cursor].height = 24
        cursor += 1
    if spec.subtitle:
        sheet.merge_cells(f"A{cursor}:{last_col}{cursor}")
        cell = sheet.cell(cursor, 1, spec.subtitle)
        cell.font = Font(name="Calibri", size=12, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[cursor].height = 20
        cursor += 1
    header_row = cursor
    roles = [_header_role(header) for header in headers]
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(header_row, col, header)
        cell.fill = _HEADER_FILL
        cell.font = Font(name="Calibri", size=9, bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = _THIN
        sheet.column_dimensions[get_column_letter(col)].width = min(46, max(12, len(header) + 4))
    sheet.row_dimensions[header_row].height = 32
    sheet.auto_filter.ref = f"A{header_row}:{last_col}{header_row + max(len(rows), 1)}"
    sheet.freeze_panes = f"A{header_row + 1}"

    for offset, row in enumerate(rows):
        r = header_row + 1 + offset
        for col, value in enumerate(row, start=1):
            cell = sheet.cell(r, col, value)
            role = roles[col - 1] if col - 1 < len(roles) else ""
            cell.font = Font(name="Calibri", size=10)
            cell.alignment = _CENTER if role in {"id", "date", "customer", "priority", "status"} else _WRAP
            cell.border = _THIN
            _paint_value(cell, role, str(value or ""))
        sheet.row_dimensions[r].height = 48

    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_view.showGridLines = False
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    return path


def _split_header_line(line: str) -> list[str]:
    if "\t" in line:
        return [part.strip() for part in line.split("\t") if part.strip()]
    parts = [part.strip() for part in re.split(r"\s{2,}", line) if part.strip()]
    if len(parts) >= 5:
        return parts
    return []


def _looks_like_headers(parts: list[str]) -> bool:
    if len(parts) < 4:
        return False
    blob = " ".join(parts).casefold().replace("ё", "е")
    hits = sum(
        1
        for hint in ("id", "источник", "поручен", "статус", "заказчик", "владелец", "срок")
        if hint in blob
    )
    return hits >= 3


def _header_role(header: str) -> str:
    folded = (header or "").casefold().replace("ё", "е")
    rules = (
        ("id", ("id",)),
        ("source", ("источник", "source")),
        ("date", ("дата решен", "дата")),
        ("result", ("ссылка", "док-т", "докумен")),
        ("task", ("поручен", "артефакт")),
        ("customer", ("заказчик", "customer")),
        ("owner", ("владелец", "owner")),
        ("due", ("срок",)),
        ("priority", ("приоритет",)),
        ("status", ("статус",)),
        ("risk", ("риск", "эскалац")),
        ("comment", ("комментар",)),
    )
    for role, aliases in rules:
        if any(alias in folded for alias in aliases):
            return role
    return ""


def _role_index(roles: list[str], role: str) -> int | None:
    try:
        return roles.index(role)
    except ValueError:
        return None


def _row_from_mapping(headers: list[str], roles: list[str], item: dict) -> dict[int, str]:
    values: dict[int, str] = {}
    lower = {str(key).casefold().replace("ё", "е"): key for key in item}
    for idx, header in enumerate(headers):
        folded = header.casefold().replace("ё", "е")
        raw = item.get(header)
        if raw in (None, "") and folded in lower:
            raw = item.get(lower[folded])
        if raw in (None, ""):
            role = roles[idx]
            raw = _value_for_role(role, item)
        values[idx] = _as_text(raw)
    return values


def _value_for_role(role: str, item: dict) -> object:
    aliases = {
        "id": ("ID", "id"),
        "source": ("Источник", "number", "source"),
        "date": ("Дата решения", "date", "Date"),
        "task": ("Поручение (результат/артефакт)", "topic", "text", "Поручение"),
        "customer": ("Заказчик", "customer"),
        "owner": ("Владелец", "owner", "executor"),
        "due": ("Срок", "due", "deadline"),
        "priority": ("Приоритет", "priority"),
        "status": ("Статус", "status"),
        "risk": ("Риск/эскалация", "risk"),
        "result": ("Ссылка на результат/документы", "files"),
        "comment": ("Комментарий", "comment"),
    }
    for key in aliases.get(role, ()):
        if item.get(key) not in (None, ""):
            return item.get(key)
    return ""


def _row_from_assignment(headers: list[str], roles: list[str], item: dict) -> dict[int, str]:
    mapped = {
        "source": _assignment_number(item),
        "date": _format_date(item.get("date")),
        "task": _assignment_topic(item),
        "customer": item.get("customer") or "",
        "owner": _assignment_owners(item),
        "due": _format_date(item.get("due")) or _line_due(item),
        "priority": _assignment_priority(item),
        "status": _map_status(item.get("status")),
        "risk": _assignment_risk(item),
        "result": _assignment_files(item),
        "comment": _assignment_comment(item),
    }
    values: dict[int, str] = {}
    for idx, role in enumerate(roles):
        if role == "id":
            values[idx] = ""
            continue
        values[idx] = _as_text(mapped.get(role, ""))
        if not values[idx]:
            values[idx] = _as_text(_value_for_role(role, item))
    return values


def _assignment_number(item: dict) -> str:
    return str(item.get("number") or item.get("Number") or "").strip()


def _assignment_topic(item: dict) -> str:
    topic = str(item.get("topic") or item.get("text") or "").strip()
    bits = [topic] if topic else []
    for line in item.get("lines") or []:
        if not isinstance(line, dict):
            continue
        text = str(line.get("text") or "").strip()
        if not text:
            continue
        number = line.get("line")
        bits.append(f"п.{number} {text}" if number else text)
    return "; ".join(bits)


def _assignment_owners(item: dict) -> str:
    names: list[str] = []
    for line in item.get("lines") or []:
        if not isinstance(line, dict):
            continue
        name = str(line.get("executor") or line.get("owner") or "").strip()
        if name and name not in names:
            names.append(name)
    extra = str(item.get("owner") or item.get("reporter") or "").strip()
    if extra and extra not in names:
        names.append(extra)
    return " ".join(names)


def _assignment_priority(item: dict) -> str:
    for line in item.get("lines") or []:
        if isinstance(line, dict) and str(line.get("priority") or "").strip():
            return str(line.get("priority") or "").strip()
    return str(item.get("priority") or "").strip()


def _assignment_risk(item: dict) -> str:
    notes: list[str] = []
    if item.get("overdue"):
        notes.append("просрочено")
    for line in item.get("lines") or []:
        if isinstance(line, dict) and line.get("overdue"):
            number = line.get("line")
            due = _format_date(line.get("due"))
            notes.append(f"п.{number} просрочен {due}".strip() if number else f"просрочен {due}".strip())
    return "; ".join(part for part in notes if part)


def _assignment_files(item: dict) -> str:
    names: list[str] = []
    for file in item.get("files") or []:
        if isinstance(file, dict):
            name = str(file.get("name") or "").strip()
            if name:
                names.append(name)
        elif str(file).strip():
            names.append(str(file).strip())
    return "; ".join(names)


def _assignment_comment(item: dict) -> str:
    status = str(item.get("status") or "").strip()
    extra = str(item.get("comment") or "").strip()
    parts = [f"Статус 1С: {status}" if status else "", extra]
    return ". ".join(part for part in parts if part)


def _line_due(item: dict) -> str:
    dates: list[str] = []
    for line in item.get("lines") or []:
        if isinstance(line, dict):
            formatted = _format_date(line.get("due"))
            if formatted:
                dates.append(formatted)
    return dates[0] if dates else ""


def _map_status(status: object) -> str:
    folded = "".join(str(status or "").casefold().replace("ё", "е").split())
    if any(token in folded for token in ("принят", "исполнен", "закрыт", "closed")):
        return "CLOSED принято и закрыто"
    if "проверк" in folded:
        return "IN PROGRESS в работе"
    if folded:
        return "IN PROGRESS в работе"
    return ""


def _format_date(value: object) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    for fmt, size in (("%Y-%m-%dT%H:%M:%S", 19), ("%Y-%m-%d", 10), ("%d.%m.%Y", 10)):
        try:
            return datetime.strptime(text[:size], fmt).strftime("%d.%m.%Y")
        except ValueError:
            continue
    return text


def _row_key(values: dict[int, str], roles: list[str]) -> str:
    source_idx = _role_index(roles, "source")
    raw = values.get(source_idx, "") if source_idx is not None else ""
    match = _AST_RE.search(raw)
    if match:
        return match.group(0).upper().replace("ACT", "АСТ")
    number = raw.strip()
    return number


def _next_act_id(rows, roles: list[str]) -> int:
    idx = _role_index(roles, "id")
    highest = 0
    if idx is None:
        return 1
    for row in rows:
        match = _ACT_ID_RE.search(str(row.get(idx) or ""))
        if match:
            highest = max(highest, int(match.group(1)))
    return highest + 1


def _prefer_filled(base: dict[int, str], extra: dict[int, str]) -> dict[int, str]:
    merged = dict(base)
    for key, value in extra.items():
        if not merged.get(key) and value:
            merged[key] = value
    return merged


def _overlay_assignment(tracker: dict[int, str], incoming: dict[int, str], roles: list[str]) -> dict[int, str]:
    merged = dict(tracker)
    status_idx = _role_index(roles, "status")
    for idx, value in incoming.items():
        if not value:
            continue
        if idx == status_idx and merged.get(idx):
            continue
        if not merged.get(idx):
            merged[idx] = value
    return merged


def _paint_value(cell, role: str, value: str) -> None:
    folded = value.casefold().replace("ё", "е")
    if role == "priority":
        if "критич" in folded:
            cell.fill = _CRIT_FILL
        elif "высок" in folded:
            cell.fill = _HIGH_FILL
    if role == "status":
        if folded.startswith("closed") or "принят" in folded or "закрыт" in folded:
            cell.fill = _CLOSED_FILL
        elif "progress" in folded or "работ" in folded or "проверк" in folded:
            cell.fill = _PROGRESS_FILL


def _as_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return "; ".join(_as_text(item) for item in value if _as_text(item))
    return str(value).strip()


def _sidecar_text(path: Path) -> str:
    sidecar = Path(str(path) + ".txt")
    if sidecar.is_file():
        return sidecar.read_text(encoding="utf-8")
    return ""
