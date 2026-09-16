"""Оформление Excel и Word: создать красивый файл или переоформить существующий."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from app.tools.ac.agent_workspace import AgentWorkspaceResolver, WorkspaceError
from app.tools.ac.base import BaseTool
from app.tools.ac.office_style import (
    as_kpis,
    as_sections,
    as_word_table,
    normalize_table,
    restyle_docx,
    restyle_excel_workbook,
    theme_names,
    write_docx,
    write_excel_sheet,
)
from app.tools.ac.registry import ToolRegistry
from app.tools.ac.tooling import (
    ToolCallResult,
    ToolDefinition,
    ToolExecutionMode,
    ToolSideEffectLevel,
)


def _docx_available() -> bool:
    import importlib.util

    return importlib.util.find_spec("docx") is not None


class OfficeFormatDocumentTool(BaseTool):
    """ChatGPT-подобный инструмент: красивый Excel/Word из данных или restyle файла."""

    def __init__(self, resolver: AgentWorkspaceResolver) -> None:
        super().__init__(
            ToolDefinition(
                name="office.format_document",
                title="Оформление Excel и Word",
                description=(
                    "Создаёт или переоформляет Excel (.xlsx) / Word (.docx) "
                    "корпоративным шаблоном: тема, титул, KPI, шапка таблицы, "
                    "зебра, статусы, колонтитулы. "
                    "Если файла ещё нет — передай headers/rows (Excel) или "
                    "sections/table (Word). Если файл уже есть и данных нет — "
                    "только применит оформление. "
                    "excel.create_workbook и report.export_document оформляют сами; "
                    "этот инструмент — когда нужно «сделать красиво» или сменить тему."
                ),
                side_effect_level=ToolSideEffectLevel.CREATE_DRAFT,
                execution_mode=ToolExecutionMode.LOCAL,
                requires_human_approval=False,
                input_schema={
                    "type": "object",
                    "properties": {
                        "filename": {
                            "type": "string",
                            "description": "Имя файла в папке агента, с расширением .xlsx или .docx",
                        },
                        "format": {
                            "type": "string",
                            "enum": ["xlsx", "docx"],
                            "description": "Формат, если его нет в имени файла",
                        },
                        "theme": {
                            "type": "string",
                            "enum": ["navy", "forest", "graphite", "wine", "sand"],
                            "description": "Цветовая тема. По умолчанию navy.",
                        },
                        "title": {"type": "string", "description": "Заголовок на обложке/баннере"},
                        "subtitle": {"type": "string", "description": "Подзаголовок или период"},
                        "sheet": {"type": "string", "description": "Имя листа Excel"},
                        "headers": {
                            "type": "array",
                            "items": {"type": "string"},
                            "description": "Заголовки таблицы Excel",
                        },
                        "rows": {"type": "array", "description": "Строки Excel: списки или объекты"},
                        "summary": {"type": "string", "description": "Резюме в начале Word"},
                        "sections": {
                            "type": "array",
                            "description": "Разделы Word: [{heading, body}]",
                        },
                        "table": {
                            "type": "object",
                            "description": "Таблица Word: {headers, rows}",
                        },
                        "kpis": {
                            "type": "array",
                            "description": "Плашки [{label, value, hint}]",
                        },
                    },
                    "required": ["filename"],
                },
                output_schema={"type": "object"},
            )
        )
        self._resolver = resolver

    def execute(self, input_data: dict) -> ToolCallResult:
        raw_name = str(input_data.get("filename") or "").strip()
        if not raw_name:
            return self._fail("INVALID_FILENAME", "Укажи filename.")
        kind = self._kind(raw_name, str(input_data.get("format") or ""))
        theme = str(input_data.get("theme") or "navy").strip()
        if theme not in theme_names():
            theme = "navy"
        title = str(input_data.get("title") or Path(raw_name).stem).strip()
        try:
            workspace = self._resolver.for_agent(
                self._resolver.agent_id_from_input(input_data)
            )
            path = workspace.resolve(self._with_suffix(raw_name, kind))
        except WorkspaceError as exc:
            return self._fail("WORKSPACE_ERROR", str(exc))

        try:
            if kind == "xlsx":
                return self._excel(path, input_data, title, theme)
            return self._word(path, input_data, title, theme)
        except WorkspaceError as exc:
            return self._fail("WORKSPACE_ERROR", str(exc))
        except Exception as exc:  # noqa: BLE001
            return self._fail("OFFICE_WRITE_ERROR", str(exc))

    def _excel(self, path: Path, input_data: dict, title: str, theme: str) -> ToolCallResult:
        from openpyxl import Workbook, load_workbook

        headers, rows = normalize_table(input_data.get("headers"), input_data.get("rows"))
        kpis = as_kpis(input_data.get("kpis"))
        subtitle = str(input_data.get("subtitle") or "").strip()
        sheet = str(input_data.get("sheet") or "Лист1")
        has_data = bool(headers or rows or kpis or title)
        if path.is_file() and not headers and input_data.get("rows") is None:
            workbook = load_workbook(path)
            try:
                restyle_excel_workbook(workbook, theme=theme, title=title)
                workbook.save(path)
            finally:
                workbook.close()
            action = "restyle"
        else:
            if not has_data and not path.is_file():
                return self._fail(
                    "OFFICE_EMPTY",
                    "Для нового Excel нужны headers/rows или существующий файл.",
                )
            workbook = Workbook()
            try:
                write_excel_sheet(
                    workbook,
                    sheet=sheet,
                    headers=headers,
                    rows=rows,
                    title=title,
                    subtitle=subtitle,
                    theme=theme,
                    kpis=kpis,
                )
                workbook.save(path)
            finally:
                workbook.close()
            action = "create"
        return self._ok(path, "xlsx", theme, action)

    def _word(self, path: Path, input_data: dict, title: str, theme: str) -> ToolCallResult:
        if not _docx_available():
            return self._fail(
                "DOCX_UNAVAILABLE",
                "python-docx не установлен — Word оформить нельзя.",
            )
        sections = as_sections(input_data.get("sections"))
        headers, rows = as_word_table(input_data.get("table"))
        summary = str(input_data.get("summary") or "").strip()
        kpis = as_kpis(input_data.get("kpis"))
        if path.is_file() and not sections and not summary and not rows and not kpis:
            restyle_docx(path, theme=theme, title=title)
            action = "restyle"
        else:
            if not sections and not summary and not rows and not kpis:
                return self._fail(
                    "OFFICE_EMPTY",
                    "Для нового Word нужны summary/sections/table или существующий файл.",
                )
            write_docx(
                path,
                title=title,
                summary=summary,
                sections=sections,
                headers=headers,
                rows=rows,
                theme=theme,
                kpis=kpis,
            )
            action = "create"
        return self._ok(path, "docx", theme, action)

    def _ok(self, path: Path, fmt: str, theme: str, action: str) -> ToolCallResult:
        return ToolCallResult(
            ok=True,
            tool_name=self.definition.name,
            output_data={
                "file": str(path),
                "filename": path.name,
                "path": str(path),
                "format": fmt,
                "theme": theme,
                "action": action,
                "generated_at": datetime.now().isoformat(timespec="seconds"),
            },
        )

    def _fail(self, error_type: str, message: str) -> ToolCallResult:
        return ToolCallResult(
            ok=False,
            tool_name=self.definition.name,
            error_type=error_type,
            error_message=message,
        )

    @staticmethod
    def _kind(filename: str, explicit: str) -> str:
        wanted = explicit.strip().casefold()
        if wanted in {"xlsx", "docx"}:
            return wanted
        suffix = Path(filename).suffix.casefold()
        if suffix in {".xlsx", ".xlsm"}:
            return "xlsx"
        if suffix in {".docx", ".doc"}:
            return "docx"
        return "xlsx"

    @staticmethod
    def _with_suffix(filename: str, kind: str) -> str:
        path = Path(filename)
        suffix = ".xlsx" if kind == "xlsx" else ".docx"
        if path.suffix.casefold() in {".xlsx", ".xlsm", ".docx", ".doc"}:
            return filename
        return f"{path.name}{suffix}"


def register_office_tools(
    registry: ToolRegistry,
    resolver: AgentWorkspaceResolver,
    *,
    skip_existing: bool = False,
) -> None:
    tool = OfficeFormatDocumentTool(resolver)
    if skip_existing and registry.has_tool(tool.definition.name):
        return
    registry.register(tool)
