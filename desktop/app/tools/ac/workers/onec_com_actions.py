"""Защищенный слой для 1C COMConnector."""

from __future__ import annotations

import base64
import io
import os
import re
import sys
from datetime import datetime, timedelta
from typing import Any

from app.tools.ac.workers.models import WorkerResult, WorkerTask
from app.tools.ac.workers.onec_actions import (
    ensure_onec_readonly_input,
    ensure_onec_readonly_tool,
)
from app.tools.ac.workers.onec_errors import OneCConnectionError
from app.tools.ac.workers.onec_meeting_notes import (
    MEETING_FIELDS,
    assert_select_only,
    build_meeting_notes_query,
    default_addressee,
    meeting_params_from_row,
    parse_note_period,
    pick_document_name,
)

FORBIDDEN_COM_METHOD_PARTS = (
    "write",
    "записать",
    "post",
    "провести",
    "delete",
    "удалить",
    "set_status",
)

DEFAULT_COM_PROGID = "V83.COMConnector"
ENV_COM_PROGID = "ONEC_COM_PROGID"
ENV_CONNECTION_STRING = "ONEC_COM_CONNECTION_STRING"
ENV_SERVER = "ONEC_COM_SERVER"
ENV_REF = "ONEC_COM_REF"
ENV_LOGIN = "ERP_LOGIN"
ENV_PASSWORD = "ERP_PASSWORD"
ENV_SEARCH_METHOD = "ONEC_COM_SEARCH_METHOD"
ENV_GET_DOCUMENT_CARD_METHOD = "ONEC_COM_GET_DOCUMENT_CARD_METHOD"
ENV_SEARCH_TASKS_METHOD = "ONEC_COM_SEARCH_TASKS_METHOD"
ENV_GET_TASK_CARD_METHOD = "ONEC_COM_GET_TASK_CARD_METHOD"


def execute_onec_com_readonly(task: WorkerTask) -> WorkerResult:
    """Выполнить read-only задачу 1С через COMConnector."""
    try:
        ensure_onec_readonly_tool(task.tool_name)
        ensure_onec_readonly_input(task.input_data)
    except Exception as exc:
        return WorkerResult(
            task_id=task.task_id,
            ok=False,
            error_type="ONEC_READONLY_POLICY_ERROR",
            error_message=str(exc),
        )

    if sys.platform != "win32":
        return _com_not_available(task, "1C COMConnector доступен только на Windows")

    pythoncom = None
    com_initialized = False
    session = None
    output_data: dict[str, Any] = {}
    used_com32 = False
    try:
        from app.tools.ac.workers.com_availability import prefers_com32

        if prefers_com32() or not _native_connector_creatable():
            print("[COM_DIAG] step=com32-helper", file=sys.stderr, flush=True)
            used_com32 = True
            output_data = _dispatch_via_com32(task)
        else:
            pythoncom, _ = _load_pywin32_modules()
            print("[COM_DIAG] step=pythoncom-loaded", file=sys.stderr, flush=True)
            pythoncom.CoInitialize()
            com_initialized = True
            print("[COM_DIAG] step=coinitialize", file=sys.stderr, flush=True)
            session = _connect_session()
            print("[COM_DIAG] step=connected", file=sys.stderr, flush=True)
            output_data = _dispatch_tool(session, task)
            print("[COM_DIAG] step=dispatched", file=sys.stderr, flush=True)
    except ImportError as exc:
        try:
            used_com32 = True
            output_data = _dispatch_via_com32(task)
        except Exception:
            return _com_not_available(task, f"pywin32 недоступен: {exc}")
    except OneCConnectionError as exc:
        if used_com32:
            return WorkerResult(
                task_id=task.task_id,
                ok=False,
                error_type="ONEC_CONNECTION_ERROR",
                error_message=str(exc),
            )
        try:
            used_com32 = True
            output_data = _dispatch_via_com32(task)
        except Exception as helper_exc:
            return WorkerResult(
                task_id=task.task_id,
                ok=False,
                error_type="ONEC_CONNECTION_ERROR",
                error_message=f"{exc}; 32-bit: {helper_exc}",
            )
    except Exception as exc:
        return WorkerResult(
            task_id=task.task_id,
            ok=False,
            error_type="ONEC_WORKER_ERROR",
            error_message=str(exc),
        )
    finally:
        session = None
        if pythoncom is not None and com_initialized:
            # In a short-lived worker process, let Python/Windows tear down COM
            # during process exit. Manual CoUninitialize after COM proxy release
            # has been unstable here and could crash or hang the helper process.
            print("[COM_DIAG] step=skip-co-uninitialize", file=sys.stderr, flush=True)

    return WorkerResult(task_id=task.task_id, ok=True, output_data=output_data)


def ensure_no_write_like_method(method_name: str) -> None:
    """Запретить write-like методы COMConnector."""
    normalized = method_name.strip().casefold()
    if any(part in normalized for part in FORBIDDEN_COM_METHOD_PARTS):
        raise ValueError(f"Метод {method_name!r} запрещён read-only политикой 1С")


def _load_pywin32_modules() -> tuple[Any, Any]:
    """Импортировать pythoncom / win32com.client только внутри worker-а."""
    import pythoncom  # type: ignore
    import win32com.client  # type: ignore

    return pythoncom, win32com.client


def _native_connector_creatable() -> bool:
    """64-bit in-process V83.COMConnector. На этой машине его обычно нет."""
    try:
        _, win32com_client = _load_pywin32_modules()
        progid = os.environ.get(ENV_COM_PROGID, DEFAULT_COM_PROGID).strip() or DEFAULT_COM_PROGID
        win32com_client.Dispatch(progid)
    except Exception:
        return False
    return True


def _dispatch_via_com32(task: WorkerTask) -> dict[str, Any]:
    """Чтение через 32-bit COMConnector (cscript). Без записи в 1С."""
    if task.tool_name == "onec.meeting_service_notes":
        return _list_meeting_service_notes_com32(task.input_data)
    if task.tool_name == "onec.search_documents":
        return _search_documents_com32(task.input_data)
    if task.tool_name == "onec.get_document_card":
        return _get_document_card_com32(task.input_data)
    raise OneCConnectionError(
        f"Инструмент {task.tool_name} ещё не переведён на 32-bit COMConnector (cscript). "
        "32-bit Python (py -3.12-32) для этого не нужен."
    )


def _com32_select(
    attempts: list[tuple[str, list[str]]],
    *,
    timeout: int | None = None,
) -> tuple[list[dict[str, str]], int]:
    from app.tools.ac.workers.onec_com32_helper import Com32TimeoutError, run_select_first

    try:
        return run_select_first(attempts, timeout=timeout)
    except Com32TimeoutError as exc:
        raise OneCConnectionError(str(exc)) from exc
    except Exception as exc:
        _reraise_transient_com32(exc)
        raise


def _reraise_transient_com32(exc: BaseException) -> None:
    message = str(exc).strip()
    low = message.casefold()
    if message.startswith(("CREATE", "CONNECT", "TIMEOUT")) or "timed out" in low:
        raise OneCConnectionError(_human_com32_error(message)) from exc


def _human_com32_error(message: str) -> str:
    text = str(message or "").strip()
    low = text.casefold()
    if "timed out" in low or low.startswith("timeout"):
        return (
            "1С не ответила через COM за отведённое время. "
            "Повтори тот же вызов — это таймаут сеанса, не пустой список документов."
        )
    if "cscript.exe" in low or "run.vbs" in low:
        return (
            "1С не ответила через COM. "
            "Повтори тот же вызов — не меняй параметры поиска."
        )
    if text.startswith("CONNECT"):
        detail = text[len("CONNECT") :].strip()
        return f"Не удалось открыть сеанс 1С: {detail}" if detail else "Не удалось открыть сеанс 1С."
    if text.startswith("CREATE"):
        detail = text[len("CREATE") :].strip()
        return f"Не удалось создать COMConnector: {detail}" if detail else "Не удалось создать COMConnector."
    return text


def _list_meeting_service_notes_com32(
    input_data: dict[str, Any],
    *,
    timeout: int | None = None,
) -> dict[str, Any]:
    from app.tools.ac.workers.onec_meeting_notes import (
        build_meeting_notes_query_latin,
        note_from_com32_row,
    )

    args = input_data if isinstance(input_data, dict) else {}
    date_from, date_to = parse_note_period(args)
    addressee = str(args.get("fio") or "").strip() or default_addressee()
    limit = int(args.get("max_results") or args.get("limit") or 50)
    specs: list[tuple[str, list[str], str, bool, bool]] = []
    variants = (
        {
            "presentation": True,
            "deref": False,
            "include_addressee": True,
            "include_meeting_fields": True,
            "include_schedule": False,
            "person_field": "МенеджерКому",
        },
        {
            "presentation": True,
            "deref": False,
            "include_addressee": False,
            "include_meeting_fields": True,
            "include_schedule": False,
            "person_field": "МенеджерКому",
        },
        {
            "presentation": False,
            "deref": True,
            "include_addressee": False,
            "include_meeting_fields": False,
            "include_schedule": False,
            "person_field": "МенеджерКому",
        },
    )
    for document_name in ("ТД_СлужебнаяЗаписка", "СлужебнаяЗаписка"):
        for variant in variants:
            query, columns = build_meeting_notes_query_latin(
                document_name=document_name,
                date_from=date_from,
                date_to=date_to,
                fio=addressee,
                limit=limit,
                include_addressee=bool(variant["include_addressee"]),
                deref=bool(variant["deref"]),
                presentation=bool(variant["presentation"]),
                person_field=str(variant["person_field"]),
                include_meeting_fields=bool(variant["include_meeting_fields"]),
                include_schedule=bool(variant["include_schedule"]),
            )
            specs.append(
                (
                    query,
                    columns,
                    document_name,
                    bool(variant["include_addressee"]),
                    bool(variant["deref"] or variant["presentation"]),
                )
            )
    try:
        rows, chosen = _com32_select(
            [(query, columns) for query, columns, *_ in specs],
            timeout=timeout,
        )
    except OneCConnectionError:
        raise
    except Exception as exc:  # noqa: BLE001
        message = str(exc).strip()
        return {
            "notes": [],
            "count": 0,
            "source": "onec_com32",
            "readonly": True,
            "date_from": date_from.isoformat(),
            "date_to": date_to.isoformat(),
            "addressee": addressee,
            "theme": "организация совещаний",
            "error": message or "32-bit COM не вернул служебные записки",
        }
    document_name, include_addressee, deref = specs[chosen][2:]
    notes = [
        note_from_com32_row(row, document_name=document_name, addressee=addressee)
        for row in rows
    ]
    return {
        "notes": notes,
        "count": len(notes),
        "source": "onec_com32",
        "readonly": True,
        "date_from": date_from.isoformat(),
        "date_to": date_to.isoformat(),
        "addressee": addressee,
        "theme": "организация совещаний",
        "document_type": document_name,
        "addressed_filter": include_addressee,
        "deref": deref,
        "method": "select_meeting_service_notes_com32",
    }


def _search_documents_com32(
    input_data: dict[str, Any],
    *,
    timeout: int | None = None,
) -> dict[str, Any]:
    """Поиск документов 1С через 32-bit SELECT. Без записи."""
    from app.tools.ac.workers.onec_meeting_notes import (
        build_document_search_query_latin,
        build_incoming_search_query_latin,
        document_from_com32_row,
    )

    args = input_data if isinstance(input_data, dict) else {}
    query = str(
        args.get("number")
        or args.get("query")
        or args.get("document_ref")
        or ""
    ).strip()
    limit = int(args.get("max_results") or args.get("limit") or 10)
    exact_number = bool(str(args.get("number") or "").strip()) or _looks_like_document_number(query)
    specs: list[tuple[str, list[str], str, str]] = []
    for document_name in ("ТД_СлужебнаяЗаписка", "СлужебнаяЗаписка"):
        if exact_number and query:
            text, columns = build_document_search_query_latin(
                document_name=document_name,
                query=query,
                limit=limit,
                exact_number=True,
                include_meeting_fields=True,
            )
            specs.append((text, columns, document_name, "Служебная записка"))
        text, columns = build_document_search_query_latin(
            document_name=document_name,
            query=query,
            limit=limit,
            exact_number=False,
            include_meeting_fields=True,
        )
        specs.append((text, columns, document_name, "Служебная записка"))
        text, columns = build_document_search_query_latin(
            document_name=document_name,
            query=query,
            limit=limit,
            exact_number=False,
            include_meeting_fields=False,
        )
        specs.append((text, columns, document_name, "Служебная записка"))
    if query:
        incoming, incoming_columns = build_incoming_search_query_latin(
            query=query,
            limit=limit,
            exact_number=exact_number,
        )
        specs.append((incoming, incoming_columns, "ТД_ВходящаяКорреспонденция", "Входящая корреспонденция"))
        if exact_number:
            incoming_like, incoming_like_columns = build_incoming_search_query_latin(
                query=query,
                limit=limit,
                exact_number=False,
            )
            specs.append(
                (
                    incoming_like,
                    incoming_like_columns,
                    "ТД_ВходящаяКорреспонденция",
                    "Входящая корреспонденция",
                )
            )
    try:
        rows, chosen = _com32_select(
            [(text, columns) for text, columns, *_ in specs],
            timeout=timeout,
        )
    except OneCConnectionError:
        raise
    except Exception as exc:  # noqa: BLE001
        message = str(exc).strip()
        return {
            "documents": [],
            "count": 0,
            "source": "onec_com32",
            "method": "select_documents_com32",
            "query": query,
            "readonly": True,
            "error": message or "32-bit COM не вернул документы",
        }
    document_name, document_type = specs[chosen][2:]
    documents = [
        document_from_com32_row(row, document_name=document_name, document_type=document_type)
        for row in rows
    ]
    return {
        "documents": documents,
        "count": len(documents),
        "source": "onec_com32",
        "method": "select_documents_com32",
        "query": query,
        "readonly": True,
        "document_type": document_type,
        "metadata_name": document_name,
    }


def _get_document_card_com32(
    input_data: dict[str, Any],
    *,
    timeout: int | None = None,
) -> dict[str, Any]:
    """Карточка документа 1С через 32-bit SELECT. Без записи."""
    args = input_data if isinstance(input_data, dict) else {}
    query = str(
        args.get("number")
        or args.get("document_ref")
        or args.get("query")
        or ""
    ).strip()
    if not query:
        raise OneCConnectionError("Для get_document_card нужен number, document_ref или query")
    raw = _search_documents_com32(
        {
            "number": query if _looks_like_document_number(query) else "",
            "query": query,
            "max_results": 1,
        },
        timeout=timeout,
    )
    documents = raw.get("documents") if isinstance(raw.get("documents"), list) else []
    document = documents[0] if documents and isinstance(documents[0], dict) else {}
    result: dict[str, Any] = {
        "document": document,
        "source": "onec_com32",
        "method": "select_document_card_com32",
        "readonly": True,
        "query": query,
    }
    if raw.get("error"):
        result["error"] = raw["error"]
    return result


def _connect_session() -> Any:
    """Подключиться к 1С через COMConnector."""
    _, win32com_client = _load_pywin32_modules()
    progid = os.environ.get(ENV_COM_PROGID, DEFAULT_COM_PROGID).strip() or DEFAULT_COM_PROGID
    connection_string = _connection_string()
    if not connection_string:
        raise OneCConnectionError(
            "Не заданы параметры подключения к 1С COMConnector. "
            f"Укажите {ENV_CONNECTION_STRING} или {ENV_SERVER}/{ENV_REF}."
        )
    try:
        print(f"[COM_DIAG] step=dispatch progid={progid}", file=sys.stderr, flush=True)
        connector = win32com_client.Dispatch(progid)
    except Exception as exc:  # noqa: BLE001
        raise OneCConnectionError(
            f"Не удалось создать COMConnector {progid!r}: {exc}"
        ) from exc
    try:
        print("[COM_DIAG] step=connect", file=sys.stderr, flush=True)
        session = connector.Connect(connection_string)
        del connector
        return session
    except Exception as exc:  # noqa: BLE001
        raise OneCConnectionError(
            f"Не удалось подключиться к 1С через COMConnector: {exc}"
        ) from exc


def _connection_string() -> str:
    """Собрать connection string из env."""
    explicit = os.environ.get(ENV_CONNECTION_STRING, "").strip()
    if explicit:
        return explicit
    server = os.environ.get(ENV_SERVER, "").strip()
    ref = os.environ.get(ENV_REF, "").strip()
    if not server or not ref:
        return ""
    parts = [f"Srvr={_quote_conn_value(server)}", f"Ref={_quote_conn_value(ref)}"]
    login = os.environ.get(ENV_LOGIN, "").strip()
    password = os.environ.get(ENV_PASSWORD, "").strip()
    if login and password:
        parts.append(f"Usr={_quote_conn_value(login)}")
        parts.append(f"Pwd={_quote_conn_value(password)}")
    return ";".join(parts) + ";"


def _quote_conn_value(value: str) -> str:
    """Кавычить значение connection string, если в нём есть пробелы или спецсимволы."""
    if not value:
        return value
    if any(ch in value for ch in (' ', ';', '"')):
        return '"' + value.replace('"', '""') + '"'
    return value


def _dispatch_tool(session: Any, task: WorkerTask) -> dict[str, Any]:
    """Маршрутизировать read-only tool_name к COM-методу."""
    if task.tool_name == "onec.search_documents":
        query = str(task.input_data.get("number") or task.input_data.get("query") or "").strip()
        document_type = str(
            task.input_data.get("document_type")
            or task.input_data.get("type")
            or task.input_data.get("section")
            or ""
        ).strip() or None
        limit = int(task.input_data.get("max_results") or task.input_data.get("limit") or 10)
        if not query:
            return {
                "documents": [],
                "count": 0,
                "source": "onec_com",
                "method": "get_incoming_correspondence",
                "note": "Укажите номер документа или текст для поиска",
            }
        raw = _search_documents_across_specs(
            session,
            query=query,
            limit=limit,
            document_type=document_type,
        )
        documents = raw.get("documents") if isinstance(raw.get("documents"), list) else []
        return {
            "documents": documents,
            "count": len(documents),
            "document_type": raw.get("document_type"),
            "browse_only": raw.get("browse_only", False),
            "source": raw.get("source", "onec_com"),
            "method": raw.get("method", "query_documents"),
            "query": query,
            **({"error": raw.get("error")} if raw.get("error") else {}),
        }
    if task.tool_name == "onec.get_document_card":
        query = str(
            task.input_data.get("query")
            or task.input_data.get("document_type")
            or task.input_data.get("type")
            or task.input_data.get("section")
            or task.input_data.get("number")
            or ""
        ).strip()
        if not query:
            raise OneCConnectionError("Для get_document_card нужен query или document_type")
        document_type = str(
            task.input_data.get("document_type")
            or task.input_data.get("type")
            or task.input_data.get("section")
            or ""
        ).strip() or None
        raw = _search_documents_across_specs(session, query=query, limit=1, document_type=document_type)
        documents = raw.get("documents") if isinstance(raw.get("documents"), list) else []
        document = documents[0] if documents and isinstance(documents[0], dict) else {}
        if document:
            document = _attach_com_tabular_parts(session, document)
            document = _populate_record_attachments(session, document)
        return {
            "document": document,
            "source": "onec_com",
            "method": "get_document_card_from_search",
        }
    if task.tool_name == "onec.list_attachments":
        return _list_attachments_tool(session, task.input_data)
    if task.tool_name == "onec.read_attachment":
        args = task.input_data if isinstance(task.input_data, dict) else {}
        return _read_attachment_for_owner(
            session,
            metadata_name=str(args.get("metadata_name") or args.get("document_type") or ""),
            kind=str(args.get("kind") or "document"),
            attachment_ref=str(args.get("attachment_ref") or args.get("ref") or ""),
            filename=str(args.get("filename") or args.get("name") or ""),
            owner_ref=str(args.get("owner_ref") or args.get("document_ref") or ""),
            number=str(args.get("number") or args.get("query") or ""),
        )
    if task.tool_name == "onec.search_tasks":
        mine_only = bool(task.input_data.get("mine_only", True))
        limit = int(task.input_data.get("max_results") or task.input_data.get("limit") or 10)
        tasks, source = query_performer_tasks(session, mine_only=mine_only, limit=limit)
        return {
            "tasks": tasks,
            "count": len(tasks),
            "source": "onec_com",
            "task_source": source,
            "method": "query_performer_tasks",
        }
    if task.tool_name == "onec.get_task_card":
        number = str(task.input_data.get("number") or task.input_data.get("task_number") or "").strip()
        if not number:
            raise OneCConnectionError("Для get_task_card нужен номер задачи")
        raw = get_task_details(session, number=number)
        task_record = raw.get("fields", {}) if raw.get("found") else {}
        return {
            "task": task_record,
            "source": "onec_com",
            "method": "get_task_details",
        }
    if task.tool_name == "onec.meeting_service_notes":
        return _list_meeting_service_notes(session, task.input_data)
    raise OneCConnectionError(f"Неизвестный COM tool_name: {task.tool_name}")


def _safe_str(value: Any, limit: int = 500) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.startswith("0001-01-01"):
        return ""
    if len(text) > limit:
        return text[: limit - 3] + "..."
    return text


def _looks_like_document_number(value: str) -> bool:
    if not value:
        return False
    if " " in value:
        return False
    return any(ch.isdigit() for ch in value) and bool(re.fullmatch(r"[\w./\\\-()]+", value))


def _normalize_search_text(value: str) -> str:
    return " ".join(str(value or "").split()).casefold()


METADATA_COLLECTION_NAMES = ("Документы", "Documents", "Справочники", "Catalogs")
SELECTABLE_FIELD_HINTS = (
    "номер",
    "наимен",
    "код",
    "коммент",
    "содерж",
    "тема",
    "описан",
    "статус",
    "ответствен",
    "организац",
    "контрагент",
    "клиент",
    "поставщик",
    "сумм",
    "дата",
    "текст",
    "соглаш",
    "договор",
    "вид",
    "тип",
    "участник",
    "инициатор",
    "автор",
    "организатор",
    "длительн",
    "формат",
    "время",
    "совеща",
    "место",
    "планир",
)
TEXT_SEARCH_FIELD_HINTS = (
    "номер",
    "наимен",
    "код",
    "коммент",
    "содерж",
    "тема",
    "описан",
    "статус",
    "текст",
    "соглаш",
    "договор",
    "участник",
    "инициатор",
    "совеща",
)


def _tokenize_search_text(value: str) -> list[str]:
    return [token for token in re.findall(r"[0-9A-Za-zА-Яа-яЁё]+", _normalize_search_text(value)) if token]


def _stem_token(token: str) -> str:
    token = token.casefold()
    for suffix in ("иями", "ями", "ами", "ями", "ого", "ему", "ому", "ыми", "ими", "ее", "ие", "ые", "ой", "ый", "ий", "ая", "яя", "ое", "ее", "ов", "ев", "ам", "ям", "ах", "ях", "ом", "ем", "ую", "юю", "а", "я", "ы", "и", "о", "е", "у", "ю", "й"):
        if token.endswith(suffix) and len(token) - len(suffix) >= 4:
            return token[: -len(suffix)]
    return token


def _stem_set(value: str) -> set[str]:
    return {_stem_token(token) for token in _tokenize_search_text(value)}


def _metadata_collection_kind(collection_name: str) -> str:
    return "catalog" if "справоч" in collection_name.casefold() or "catalog" in collection_name.casefold() else "document"


def _get_metadata_root(session: Any) -> Any | None:
    return getattr(session, "Метаданные", None) or getattr(session, "Metadata", None)


def _iter_metadata_collection_items(collection: Any) -> list[Any]:
    try:
        count = int(collection.Count())
    except Exception:
        return []
    return [collection.Get(i) for i in range(count)]


def _metadata_name(item: Any) -> str:
    for attr in ("Name", "Имя"):
        value = _safe_str(getattr(item, attr, ""))
        if value:
            return value
    return ""


def _metadata_synonym(item: Any) -> str:
    for attr in ("Synonym", "Синоним"):
        value = _safe_str(getattr(item, attr, ""))
        if value:
            return value
    return ""


def _metadata_requisite_names(item: Any) -> list[str]:
    for attr in ("Реквизиты", "Requisites", "Attributes"):
        requisites = getattr(item, attr, None)
        if requisites is None:
            continue
        names: list[str] = []
        for requisite in _iter_metadata_collection_items(requisites):
            name = _metadata_name(requisite)
            if name:
                names.append(name)
        if names:
            return names
    return []


def _metadata_tabular_sections(item: Any) -> list[dict[str, Any]]:
    for attr in ("ТабличныеЧасти", "TabularSections"):
        collection = getattr(item, attr, None)
        if collection is None:
            continue
        sections: list[dict[str, Any]] = []
        for section in _iter_metadata_collection_items(collection):
            name = _metadata_name(section)
            if not name:
                continue
            sections.append(
                {
                    "name": name,
                    "synonym": _metadata_synonym(section),
                    "fields": _metadata_requisite_names(section),
                }
            )
        if sections:
            return sections
    return []


def _is_browse_query(query: str, candidate_name: str, candidate_synonym: str) -> bool:
    query_stems = _stem_set(query)
    if not query_stems:
        return False
    label_stems = _stem_set(f"{candidate_name} {candidate_synonym}")
    if not label_stems:
        return False
    overlap = len(query_stems & label_stems)
    return overlap >= max(1, round(len(query_stems) * 0.6))


def _discover_metadata_candidates(session: Any, query: str, document_type: str | None = None) -> list[dict[str, Any]]:
    metadata_root = _get_metadata_root(session)
    if metadata_root is None:
        return []

    normalized_query = _normalize_search_text(document_type or query)
    query_stems = _stem_set(normalized_query)
    candidates: list[dict[str, Any]] = []
    for collection_name in METADATA_COLLECTION_NAMES:
        try:
            collection = getattr(metadata_root, collection_name)
        except Exception:
            continue
        kind = _metadata_collection_kind(collection_name)
        for item in _iter_metadata_collection_items(collection):
            name = _metadata_name(item)
            synonym = _metadata_synonym(item)
            label_text = _normalize_search_text(f"{name} {synonym}")
            if not label_text:
                continue
            label_stems = _stem_set(label_text)
            overlap = len(query_stems & label_stems)
            if normalized_query and overlap == 0 and normalized_query not in label_text:
                continue
            score = overlap * 10
            if normalized_query and normalized_query == _normalize_search_text(name):
                score += 50
            if normalized_query and normalized_query == _normalize_search_text(synonym):
                score += 40
            if normalized_query and normalized_query in label_text:
                score += 20
            if not score:
                continue
            candidates.append(
                {
                    "collection_name": collection_name,
                    "kind": kind,
                    "name": name,
                    "synonym": synonym,
                    "label": synonym or name or label_text,
                    "item": item,
                    "score": score,
                    "browse_only": _is_browse_query(normalized_query, name, synonym),
                    "requisites": _metadata_requisite_names(item),
                    "tabular_sections": _metadata_tabular_sections(item),
                }
            )
    candidates.sort(key=lambda candidate: candidate["score"], reverse=True)
    return candidates


def _pick_select_fields(candidate: dict[str, Any]) -> list[tuple[str, str]]:
    requisites = candidate.get("requisites") or []
    fields: list[tuple[str, str]] = [("ref", "Ссылка")]

    if requisites:
        for field_name in (name for name in requisites if _is_selectable_field_name(name)):
            if len(fields) >= 24:
                break
            if all(existing_name != field_name for _, existing_name in fields):
                fields.append((_field_alias(field_name), field_name))
    return fields


def _field_alias(field_name: str) -> str:
    alias = re.sub(r"[^0-9A-Za-zА-Яа-яЁё]+", "_", field_name).strip("_")
    return alias[:60] or "field"


def _is_selectable_field_name(field_name: str) -> bool:
    normalized = field_name.casefold()
    return any(hint in normalized for hint in SELECTABLE_FIELD_HINTS)


def _is_text_search_field_name(field_name: str) -> bool:
    normalized = field_name.casefold()
    return any(hint in normalized for hint in TEXT_SEARCH_FIELD_HINTS)


def _build_metadata_query(
    candidate: dict[str, Any],
    query: str,
    *,
    limit: int,
    browse_only: bool,
) -> tuple[str, list[tuple[str, str]]]:
    select_fields = _pick_select_fields(candidate)
    select_clause = ",\n        ".join(f"Д.{field_name} КАК {alias}" for alias, field_name in select_fields)
    query_lines = [f"ВЫБРАТЬ ПЕРВЫЕ {limit}", f"        {select_clause}", f"        ИЗ {candidate['collection_name'][:-1] if False else ''}"]
    object_expr = f"{candidate['collection_name']}."  # placeholder to keep structure valid? will be replaced below
    object_expr = candidate["name"]
    # Build the fully qualified metadata reference from the collection kind.
    if candidate["kind"] == "catalog":
        object_expr = f"Справочник.{candidate['name']}"
    else:
        object_expr = f"Документ.{candidate['name']}"
    query_lines[2] = f"        ИЗ {object_expr} КАК Д"

    conditions: list[str] = []
    if candidate["kind"] in {"document", "catalog"}:
        conditions.append("НЕ Д.ПометкаУдаления")
    if not browse_only:
        safe_query = query.replace('"', '""')
        search_fields = [
            field_name
            for _, field_name in select_fields
            if _is_text_search_field_name(field_name)
        ]
        if search_fields:
            conditions.append(
                "("
                + " ИЛИ ".join(f"Д.{field_name} ПОДОБНО \"%{safe_query}%\"" for field_name in search_fields)
                + ")"
            )
    if conditions:
        query_lines.append("        ГДЕ " + " И ".join(conditions))
    order_field = "Наименование" if candidate["kind"] == "catalog" else "Дата"
    if order_field not in {field_name for _, field_name in select_fields}:
        order_field = select_fields[0][1]
    query_lines.append(f"        УПОРЯДОЧИТЬ ПО Д.{order_field} УБЫВ")
    return "\n".join(query_lines), select_fields


def _collect_metadata_rows(
    session: Any,
    table: Any,
    candidate: dict[str, Any],
    select_fields: list[tuple[str, str]],
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for i in range(table.Count()):
        row = table.Get(i)
        fields: dict[str, str] = {}
        for alias, attr in select_fields:
            val = getattr(row, attr, None)
            fields[alias] = _safe_str(getattr(val, "Наименование", None) or val, 5000)
        item = {
            "found": True,
            "document_type": candidate["synonym"] or candidate["name"],
            "metadata_name": candidate["name"],
            "kind": candidate["kind"],
            "ref": _safe_str(getattr(row, "Ref", "") or getattr(row, "Ссылка", "")),
            "number": _safe_str(getattr(row, "Number", "") or getattr(row, "Номер", "")),
            "fields": fields,
            "tabular_sections": candidate.get("tabular_sections") or [],
            "attachments": [],
        }
        rows.append(_populate_record_attachments(session, item))
    return rows


def _object_expr(candidate_kind: str, metadata_name: str) -> str:
    if candidate_kind == "catalog":
        return f"Справочник.{metadata_name}"
    return f"Документ.{metadata_name}"


def _attach_com_tabular_parts(session: Any, document: dict[str, Any]) -> dict[str, Any]:
    """Дочитать табличные части карточки (участники совещания и т.п.)."""
    if not isinstance(document, dict):
        return document
    number = str(document.get("number") or "").strip()
    metadata_name = str(document.get("metadata_name") or "").strip()
    kind = str(document.get("kind") or "document")
    sections = document.get("tabular_sections") or []
    if not number or not metadata_name or not isinstance(sections, list):
        return document
    safe_number = number.replace('"', '""')
    object_expr = _object_expr(kind, metadata_name)
    parts: dict[str, list[dict[str, str]]] = {}
    for section in sections[:8]:
        if not isinstance(section, dict):
            continue
        section_name = str(section.get("name") or "").strip()
        if not section_name:
            continue
        field_names = [str(name) for name in (section.get("fields") or []) if str(name).strip()]
        select_fields = field_names[:12] or ["НомерСтроки"]
        select_clause = ",\n        ".join(
            f"Т.{field_name} КАК {_field_alias(field_name)}" for field_name in select_fields
        )
        query_text = (
            f"ВЫБРАТЬ ПЕРВЫЕ 50\n        {select_clause}\n"
            f"        ИЗ {object_expr}.{section_name} КАК Т\n"
            f"        ГДЕ Т.Ссылка.Номер = \"{safe_number}\""
        )
        try:
            table = session.NewObject("Query", query_text).Execute().Unload()
        except Exception:
            continue
        rows_out: list[dict[str, str]] = []
        for index in range(table.Count()):
            row = table.Get(index)
            item: dict[str, str] = {}
            for field_name in select_fields:
                alias = _field_alias(field_name)
                val = getattr(row, alias, None)
                if val is None:
                    val = getattr(row, field_name, None)
                item[field_name] = _safe_str(getattr(val, "Наименование", None) or val, 2000)
            if any(item.values()):
                rows_out.append(item)
        if rows_out:
            label = str(section.get("synonym") or section_name)
            parts[label] = rows_out
    document = dict(document)
    if parts:
        document["tabular_parts"] = parts
    document.pop("tabular_sections", None)
    return document


_ATTACHMENT_SUFFIX = "ПрисоединенныеФайлы"
_OWNER_FIELDS = ("ВладелецФайла", "Owner", "Владелец")
_ATTACHMENT_NAME_FIELDS = ("Наименование", "Description", "ИмяФайла", "FileName")
_ATTACHMENT_EXT_FIELDS = ("Расширение", "Extension")
_ATTACHMENT_SIZE_FIELDS = ("Размер", "Size")


def _attachment_catalog_name(metadata_name: str) -> str:
    name = (metadata_name or "").strip()
    if not name:
        return ""
    if name.casefold().endswith(_ATTACHMENT_SUFFIX.casefold()):
        return name
    return f"{name}{_ATTACHMENT_SUFFIX}"


def _discover_attachment_catalog(session: Any, metadata_name: str) -> dict[str, Any] | None:
    catalog_name = _attachment_catalog_name(metadata_name)
    if not catalog_name:
        return None
    candidates = _discover_metadata_candidates(session, catalog_name, catalog_name)
    for candidate in candidates:
        if _ATTACHMENT_SUFFIX.casefold() in str(candidate.get("name") or "").casefold():
            return candidate
    return None


def _resolve_owner_number(session: Any, *, metadata_name: str, kind: str, number: str) -> str:
    safe_number = number.replace('"', '""')
    object_expr = _object_expr(kind or "document", metadata_name)
    query_text = (
        f"ВЫБРАТЬ ПЕРВЫЕ 1\n"
        f"        Д.Ссылка КАК Ref\n"
        f"        ИЗ {object_expr} КАК Д\n"
        f"        ГДЕ Д.Номер = \"{safe_number}\""
    )
    try:
        table = session.NewObject("Query", query_text).Execute().Unload()
        if table.Count():
            return _safe_str(getattr(table.Get(0), "Ref", ""))
    except Exception:
        return ""
    return ""


def _list_attachments_for_owner(
    session: Any,
    *,
    metadata_name: str,
    kind: str = "document",
    owner_ref: str = "",
    number: str = "",
    limit: int = 30,
) -> list[dict[str, Any]]:
    catalog = _discover_attachment_catalog(session, metadata_name)
    if catalog is None:
        return []
    owner = (owner_ref or "").strip()
    if not owner and number.strip():
        owner = _resolve_owner_number(session, metadata_name=metadata_name, kind=kind, number=number)
    if not owner:
        return []

    catalog_name = str(catalog.get("name") or "")
    object_expr = f"Справочник.{catalog_name}"
    limit = max(1, min(100, int(limit)))
    attachments: list[dict[str, Any]] = []
    for owner_field in _OWNER_FIELDS:
        query_text = (
            f"ВЫБРАТЬ ПЕРВЫЕ {limit}\n"
            f"        Ф.Ссылка КАК Ref,\n"
            f"        Ф.Наименование КАК Name,\n"
            f"        Ф.Расширение КАК Ext,\n"
            f"        Ф.Размер КАК Size\n"
            f"        ИЗ {object_expr} КАК Ф\n"
            f"        ГДЕ Ф.{owner_field}.Ссылка = &Owner"
        )
        try:
            query = session.NewObject("Query", query_text)
            query.SetParameter("Owner", owner)
            table = query.Execute().Unload()
        except Exception:
            continue
        for index in range(table.Count()):
            row = table.Get(index)
            name = _safe_str(getattr(row, "Name", "") or getattr(row, "Наименование", ""), 500)
            ext = _safe_str(getattr(row, "Ext", "") or getattr(row, "Extension", ""), 20)
            size_raw = getattr(row, "Size", None)
            try:
                size = int(size_raw)
            except (TypeError, ValueError):
                size = 0
            filename = name
            if ext and not filename.casefold().endswith(f".{ext.casefold()}"):
                filename = f"{filename}.{ext}" if filename else f"file.{ext}"
            attachments.append(
                {
                    "ref": _safe_str(getattr(row, "Ref", "")),
                    "name": filename or name or "file",
                    "extension": ext,
                    "size": size,
                    "metadata_name": catalog_name,
                    "owner_ref": owner,
                }
            )
        if attachments:
            break
    return attachments


def _populate_record_attachments(session: Any, record: dict[str, Any]) -> dict[str, Any]:
    if not isinstance(record, dict):
        return record
    metadata_name = str(record.get("metadata_name") or "").strip()
    if not metadata_name:
        return record
    kind = str(record.get("kind") or "document")
    attachments = _list_attachments_for_owner(
        session,
        metadata_name=metadata_name,
        kind=kind,
        owner_ref=str(record.get("ref") or ""),
        number=str(record.get("number") or ""),
    )
    updated = dict(record)
    updated["attachments"] = attachments
    return updated


def _extract_text_from_attachment_bytes(filename: str, raw: bytes) -> str:
    suffix = os.path.splitext(filename)[1].casefold()
    if suffix == ".pdf":
        try:
            import fitz

            doc = fitz.open(stream=raw, filetype="pdf")
            parts = [page.get_text() or "" for page in doc]
            doc.close()
            text = "\n\n".join(parts).strip()
            if text:
                return text
        except Exception:
            pass
    if suffix == ".docx":
        try:
            import docx

            document = docx.Document(io.BytesIO(raw))
            parts = [para.text for para in document.paragraphs if para.text.strip()]
            return "\n".join(parts).strip()
        except Exception:
            pass
    if suffix in {".xlsx", ".xlsm"}:
        try:
            import openpyxl

            workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            parts: list[str] = []
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    values = [str(cell).strip() if cell is not None else "" for cell in row]
                    if any(values):
                        parts.append("\t".join(values))
            workbook.close()
            return "\n".join(parts).strip()
        except Exception:
            pass
    for encoding in ("utf-8", "utf-8-sig", "cp1251", "latin-1"):
        try:
            text = raw.decode(encoding).strip()
            if text:
                return text
        except UnicodeDecodeError:
            continue
    return ""


def _read_attachment_binary(session: Any, *, catalog_name: str, attachment_ref: str) -> bytes | None:
    ref = (attachment_ref or "").strip()
    if not ref:
        return None
    for accessor in (
        lambda: session.РаботаСФайлами.ДвоичныеДанныеФайла(ref),
        lambda: session.РаботаСФайлами.ДвоичныеДанныеФайла(session.NewObject("СправочникСсылка." + catalog_name, ref)),
    ):
        try:
            data = accessor()
            if data is None:
                continue
            if isinstance(data, (bytes, bytearray)):
                return bytes(data)
            if hasattr(data, "GetData"):
                return bytes(data.GetData())
        except Exception:
            continue
    return None


def _read_attachment_for_owner(
    session: Any,
    *,
    metadata_name: str,
    kind: str,
    attachment_ref: str = "",
    filename: str = "",
    owner_ref: str = "",
    number: str = "",
    max_bytes: int = 8 * 1024 * 1024,
) -> dict[str, Any]:
    catalog = _discover_attachment_catalog(session, metadata_name)
    if catalog is None:
        return {"ok": False, "error": "Каталог вложений не найден в метаданных 1С"}
    listed = _list_attachments_for_owner(
        session,
        metadata_name=metadata_name,
        kind=kind,
        owner_ref=owner_ref,
        number=number,
        limit=100,
    )
    target = None
    ref = (attachment_ref or "").strip()
    name_hint = (filename or "").strip().casefold()
    for item in listed:
        if ref and str(item.get("ref") or "") == ref:
            target = item
            break
        if name_hint and name_hint in str(item.get("name") or "").casefold():
            target = item
            break
    if target is None and listed and not ref and not name_hint:
        target = listed[0]
    if target is None:
        return {"ok": False, "error": "Вложение не найдено", "attachments": listed}

    catalog_name = str(catalog.get("name") or "")
    raw = _read_attachment_binary(session, catalog_name=catalog_name, attachment_ref=str(target.get("ref") or ""))
    if not raw:
        return {
            "ok": False,
            "error": "Не удалось прочитать двоичные данные вложения через COM",
            "attachment": target,
            "attachments": listed,
        }
    if len(raw) > max_bytes:
        return {
            "ok": False,
            "error": f"Вложение слишком большое ({len(raw)} байт), лимит {max_bytes}",
            "attachment": target,
        }
    file_name = str(target.get("name") or filename or "attachment")
    text = _extract_text_from_attachment_bytes(file_name, raw)
    return {
        "ok": True,
        "attachment": target,
        "filename": file_name,
        "size": len(raw),
        "text": text,
        "text_chars": len(text),
        "data_b64": base64.b64encode(raw).decode("ascii") if len(raw) <= 512 * 1024 else "",
        "note": "" if text else "Текст не извлечён — возможно скан; попробуй OCR на стороне сервера.",
        "source": "onec_com",
        "readonly": True,
    }


def _list_attachments_tool(session: Any, input_data: dict[str, Any]) -> dict[str, Any]:
    args = input_data if isinstance(input_data, dict) else {}
    metadata_name = str(args.get("metadata_name") or args.get("document_type") or "").strip()
    number = str(args.get("number") or args.get("document_ref") or args.get("query") or "").strip()
    owner_ref = str(args.get("owner_ref") or args.get("ref") or "").strip()
    kind = str(args.get("kind") or "document").strip() or "document"
    if not metadata_name and number:
        card = _get_document_card_any(session, number=number)
        doc = (card or {}).get("document") if isinstance(card, dict) else {}
        if isinstance(doc, dict):
            metadata_name = str(doc.get("metadata_name") or "")
            kind = str(doc.get("kind") or kind)
            owner_ref = owner_ref or str(doc.get("ref") or "")
            number = number or str(doc.get("number") or "")
    if not metadata_name:
        return {
            "attachments": [],
            "count": 0,
            "error": "Укажите metadata_name или number документа 1С",
            "source": "onec_com",
            "readonly": True,
        }
    attachments = _list_attachments_for_owner(
        session,
        metadata_name=metadata_name,
        kind=kind,
        owner_ref=owner_ref,
        number=number,
        limit=int(args.get("max_results") or args.get("limit") or 30),
    )
    return {
        "attachments": attachments,
        "count": len(attachments),
        "metadata_name": metadata_name,
        "source": "onec_com",
        "readonly": True,
        "method": "list_attachments",
    }


def _session_user_fio(session: Any) -> str:
    for name in ("ПолноеИмяПользователя", "ИмяПользователя"):
        fn = getattr(session, name, None)
        if callable(fn):
            try:
                value = _safe_str(fn())
            except Exception:
                value = ""
            if value:
                return value
    return ""


def _list_meeting_service_notes(session: Any, input_data: dict[str, Any]) -> dict[str, Any]:
    """Прочитать служебные записки. Только ВЫБРАТЬ, без записи в 1С."""
    args = input_data if isinstance(input_data, dict) else {}
    date_from, date_to = parse_note_period(args)
    addressee = str(args.get("fio") or "").strip() or _session_user_fio(session) or default_addressee()
    limit = int(args.get("max_results") or args.get("limit") or 50)
    candidates = _discover_metadata_candidates(session, "служебная записка", "служебная записка")
    candidate = pick_document_name(candidates)
    if candidate is None:
        return {
            "notes": [],
            "count": 0,
            "source": "onec_com",
            "readonly": True,
            "date_from": date_from.isoformat(),
            "date_to": date_to.isoformat(),
            "addressee": addressee,
            "theme": "организация совещаний",
            "error": "В метаданных 1С не найден документ «служебная записка»",
        }
    query_text, theme_fields, person_fields = build_meeting_notes_query(
        document_name=str(candidate["name"]),
        requisites=list(candidate.get("requisites") or []),
        date_from=date_from,
        date_to=date_to,
        fio=addressee,
        limit=limit,
    )
    assert_select_only(query_text)
    table = session.NewObject("Query", query_text).Execute().Unload()
    notes: list[dict[str, Any]] = []
    for index in range(table.Count()):
        row = table.Get(index)
        item: dict[str, Any] = {
            "document_type": candidate.get("synonym") or candidate.get("name"),
            "metadata_name": candidate.get("name"),
            "ref": _safe_str(getattr(row, "Ссылка", None) or getattr(row, "Ref", None)),
            "number": _safe_str(getattr(row, "Номер", None) or getattr(row, "Number", None)),
            "date": _safe_str(getattr(row, "Дата", None) or getattr(row, "Date", None)),
            "theme": "",
            "addressee": "",
            "fields": {},
        }
        fields: dict[str, str] = {}
        for name in (*(theme_fields or []), *(person_fields or []), "Комментарий"):
            alias = _field_alias(name)
            val = getattr(row, alias, None)
            if val is None:
                val = getattr(row, name, None)
            text = _safe_str(getattr(val, "Наименование", None) or val, 2000)
            if text:
                fields[name] = text
        item["fields"] = fields
        item["theme"] = next((fields[name] for name in theme_fields if fields.get(name)), "")
        item["addressee"] = next((fields[name] for name in person_fields if fields.get(name)), addressee)
        meeting_row: dict[str, str] = {}
        for field_name, alias, _kind in MEETING_FIELDS:
            val = getattr(row, alias, None)
            if val is None:
                val = getattr(row, field_name, None)
            meeting_row[alias] = _safe_str(getattr(val, "Наименование", None) or val, 2000)
        meeting = meeting_params_from_row(meeting_row)
        item["meeting_topic"] = meeting["topic"]
        item["place"] = meeting["place"]
        item["meeting"] = meeting
        notes.append(item)
    return {
        "notes": notes,
        "count": len(notes),
        "source": "onec_com",
        "readonly": True,
        "date_from": date_from.isoformat(),
        "date_to": date_to.isoformat(),
        "addressee": addressee,
        "theme": "организация совещаний",
        "document_type": candidate.get("synonym") or candidate.get("name"),
        "addressed_filter": bool(person_fields),
        "method": "select_meeting_service_notes",
    }


def _search_documents_across_specs(
    session: Any,
    *,
    query: str,
    limit: int,
    document_type: str | None = None,
) -> dict[str, Any]:
    candidates = _discover_metadata_candidates(session, query, document_type)
    if not candidates:
        return {
            "found": False,
            "query": query,
            "document_type": document_type or "",
            "browse_only": False,
            "documents": [],
            "count": 0,
            "source": "onec_com",
            "method": "query_documents",
            "error": "Не удалось найти подходящие типы документов в метаданных 1С",
        }

    last_error: str | None = None
    for candidate in candidates[:10]:
        try:
            query_text, select_fields = _build_metadata_query(
                candidate,
                query,
                limit=limit,
                browse_only=candidate["browse_only"],
            )
            table = session.NewObject("Query", query_text).Execute().Unload()
            documents = _collect_metadata_rows(session, table, candidate, select_fields)
            if documents:
                if limit == 1:
                    enriched = _attach_com_tabular_parts(session, documents[0])
                    documents = [_populate_record_attachments(session, enriched)]
                return {
                    "found": True,
                    "query": query,
                    "document_type": candidate["synonym"] or candidate["name"],
                    "browse_only": candidate["browse_only"],
                    "documents": documents,
                    "count": len(documents),
                    "source": "onec_com",
                    "method": "query_documents",
                }
        except Exception as exc:  # noqa: BLE001
            last_error = str(exc)
            continue

    return {
        "found": False,
        "query": query,
        "document_type": candidates[0]["synonym"] or candidates[0]["name"],
        "browse_only": candidates[0]["browse_only"],
        "documents": [],
        "count": 0,
        "source": "onec_com",
        "method": "query_documents",
        "error": last_error,
    }


def _build_card_query(candidate: dict[str, Any], number: str) -> tuple[str, list[tuple[str, str]]]:
    select_fields = _pick_select_fields(candidate)
    select_clause = ",\n        ".join(f"Д.{field_name} КАК {alias}" for alias, field_name in select_fields)
    object_expr = f"Справочник.{candidate['name']}" if candidate["kind"] == "catalog" else f"Документ.{candidate['name']}"
    safe_number = number.replace('"', '""')
    query_lines = [
        f"ВЫБРАТЬ ПЕРВЫЕ 1",
        f"        {select_clause}",
        f"        ИЗ {object_expr} КАК Д",
        "        ГДЕ НЕ Д.ПометкаУдаления",
        f"            И Д.Номер = \"{safe_number}\"",
    ]
    return "\n".join(query_lines), select_fields


def _get_document_card_any(
    session: Any,
    *,
    number: str,
    document_type: str | None = None,
) -> dict[str, Any] | None:
    candidates = _discover_metadata_candidates(session, document_type or number, document_type)
    for candidate in candidates[:10]:
        try:
            query_text, select_fields = _build_card_query(candidate, number)
            table = session.NewObject("Query", query_text).Execute().Unload()
            if not table.Count():
                continue
            row = table.Get(0)
            fields: dict[str, str] = {}
            for alias, attr in select_fields:
                val = getattr(row, attr, None)
                fields[alias] = _safe_str(getattr(val, "Наименование", None) or val, 5000)
            doc = _populate_record_attachments(
                session,
                {
                    "found": True,
                    "document_type": candidate["synonym"] or candidate["name"],
                    "metadata_name": candidate["name"],
                    "kind": candidate["kind"],
                    "number": _safe_str(getattr(row, "Number", "") or getattr(row, "Номер", "")),
                    "ref": _safe_str(getattr(row, "Ref", "") or getattr(row, "Ссылка", "")),
                    "fields": fields,
                    "attachments": [],
                },
            )
            return {
                "document": doc,
                "source": "onec_com",
                "method": "query_document_card",
            }
        except Exception:
            continue
    return None


def get_current_user_name(app: Any) -> str:
    try:
        user = app.ПользователиИнформационнойБазы.ТекущийПользователь()
        for attr in ("ПолноеИмя", "Name", "Имя"):
            value = _safe_str(getattr(user, attr, ""))
            if value:
                return value
        return _safe_str(user)
    except Exception:
        return ""


def _rows_from_table(table: Any, *, source: str) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for i in range(table.Count()):
        row = table.Get(i)
        rows.append(
            {
                "number": _safe_str(getattr(row, "Number", "") or getattr(row, "Номер", "")),
                "description": _safe_str(getattr(row, "Description", "") or getattr(row, "Наименование", "")),
                "date": _safe_str(getattr(row, "Date", "") or getattr(row, "Дата", "")),
                "due_date": _safe_str(getattr(row, "DueDate", "") or getattr(row, "СрокИсполнения", "")),
                "executor": _safe_str(getattr(row, "Executor", "") or getattr(row, "Исполнитель", "")),
                "source": source,
            }
        )
    return rows


def query_performer_tasks(app: Any, *, mine_only: bool = True, limit: int = 30) -> tuple[list[dict[str, Any]], str]:
    limit = max(1, min(100, int(limit)))
    user_name = get_current_user_name(app) if mine_only else ""
    if mine_only and user_name:
        query_text = f"""ВЫБРАТЬ ПЕРВЫЕ {limit}
            Т.Номер КАК Number,
            Т.Наименование КАК Description,
            Т.Дата КАК Date,
            Т.СрокИсполнения КАК DueDate,
            Т.Исполнитель.Наименование КАК Executor
            ИЗ Задача.ЗадачаИсполнителя КАК Т
            ГДЕ НЕ Т.Выполнена
                И Т.Исполнитель.Наименование = "{user_name.replace('"', '""')}"
            УПОРЯДОЧИТЬ ПО Т.Дата УБЫВ"""
        table = app.NewObject("Query", query_text).Execute().Unload()
        rows = _rows_from_table(table, source="erp_задача_исполнителя")
        return rows, "erp_задача_исполнителя"

    query_text = f"""ВЫБРАТЬ ПЕРВЫЕ {limit}
        Т.Номер КАК Number,
        Т.Наименование КАК Description,
        Т.Дата КАК Date,
        Т.СрокИсполнения КАК DueDate,
        Т.Исполнитель.Наименование КАК Executor
        ИЗ Задача.ЗадачаИсполнителя КАК Т
        ГДЕ НЕ Т.Выполнена
        УПОРЯДОЧИТЬ ПО Т.Дата УБЫВ"""
    table = app.NewObject("Query", query_text).Execute().Unload()
    return _rows_from_table(table, source="erp_задача_исполнителя_all"), "erp_задача_исполнителя_all"


def get_task_details(app: Any, *, number: str) -> dict[str, Any]:
    safe_number = number.replace('"', '""')
    query_text = f"""ВЫБРАТЬ ПЕРВЫЕ 1
        Т.Ссылка КАК Ref,
        Т.Номер КАК Number,
        Т.Наименование КАК Description,
        Т.Дата КАК Date,
        Т.СрокИсполнения КАК DueDate,
        Т.Исполнитель.Наименование КАК Executor,
        Т.Автор.Наименование КАК Author,
        Т.Выполнена КАК Done,
        Т.РезультатВыполнения КАК Result,
        Т.Предмет КАК Subject,
        Т.Описание КАК Details
        ИЗ Задача.ЗадачаИсполнителя КАК Т
        ГДЕ Т.Номер = "{safe_number}" """
    table = app.NewObject("Query", query_text).Execute().Unload()
    if not table.Count():
        return {"found": False, "number": number}

    row = table.Get(0)
    fields: dict[str, str] = {}
    for alias, attr in (
        ("number", "Number"),
        ("description", "Description"),
        ("date", "Date"),
        ("due_date", "DueDate"),
        ("executor", "Executor"),
        ("author", "Author"),
        ("done", "Done"),
        ("result", "Result"),
        ("details", "Details"),
    ):
        val = getattr(row, attr, None)
        fields[alias] = _safe_str(getattr(val, "Наименование", None) or val, 2000)

    record = {
        "found": True,
        "number": number,
        "metadata_name": "ЗадачаИсполнителя",
        "kind": "task",
        "fields": fields,
        "attachments": [],
    }
    return record


def get_incoming_correspondence(app: Any, *, number: str) -> dict[str, Any]:
    safe_number = number.replace('"', '""')
    query_text = f"""ВЫБРАТЬ ПЕРВЫЕ 1
        Д.Ссылка КАК Ref,
        Д.Номер КАК Number,
        Д.Дата КАК DocDate,
        Д.Комментарий КАК Comment,
        Д.Организация.Наименование КАК Org,
        Д.Контрагент.Наименование КАК Counterparty,
        Д.Содержание КАК Content,
        Д.ТемаСлужебнойЗаписки КАК MemoSubject,
        Д.EmailОтправителяПисьма КАК EmailFrom,
        Д.EmailПолучателяПисьма КАК EmailTo,
        Д.Кому КАК MailTo,
        Д.НомерИсходящий КАК OutNumber,
        Д.ДатаИсходящая КАК OutDate,
        Д.Ответственный.Наименование КАК Responsible,
        Д.ТекстHTML КАК HtmlText
        ИЗ Документ.ТД_ВходящаяКорреспонденция КАК Д
        ГДЕ Д.Номер = "{safe_number}"
            И НЕ Д.ПометкаУдаления"""
    table = app.NewObject("Query", query_text).Execute().Unload()
    if not table.Count():
        return {"found": False, "number": number}

    row = table.Get(0)
    fields: dict[str, str] = {}
    for alias, attr in (
        ("number", "Number"),
        ("date", "DocDate"),
        ("comment", "Comment"),
        ("org", "Org"),
        ("counterparty", "Counterparty"),
        ("content", "Content"),
        ("memo_subject", "MemoSubject"),
        ("email_from", "EmailFrom"),
        ("email_to", "EmailTo"),
        ("mail_to", "MailTo"),
        ("out_number", "OutNumber"),
        ("out_date", "OutDate"),
        ("responsible", "Responsible"),
        ("html_text", "HtmlText"),
    ):
        val = getattr(row, attr, None)
        fields[alias] = _safe_str(getattr(val, "Наименование", None) or val, 5000)

    return {
        "found": True,
        "number": number,
        "fields": fields,
        "attachments": [],
    }


def search_incoming_correspondence(app: Any, *, query: str, limit: int = 10) -> dict[str, Any]:
    safe_query = query.replace('"', '""')
    limit = max(1, min(50, int(limit)))
    query_text = f"""ВЫБРАТЬ ПЕРВЫЕ {limit}
        Д.Ссылка КАК Ref,
        Д.Номер КАК Number,
        Д.Дата КАК DocDate,
        Д.Комментарий КАК Comment,
        Д.Организация.Наименование КАК Org,
        Д.Контрагент.Наименование КАК Counterparty,
        Д.Содержание КАК Content,
        Д.ТемаСлужебнойЗаписки КАК MemoSubject,
        Д.EmailОтправителяПисьма КАК EmailFrom,
        Д.EmailПолучателяПисьма КАК EmailTo,
        Д.Кому КАК MailTo,
        Д.НомерИсходящий КАК OutNumber,
        Д.ДатаИсходящая КАК OutDate,
        Д.Ответственный.Наименование КАК Responsible,
        Д.ТекстHTML КАК HtmlText
        ИЗ Документ.ТД_ВходящаяКорреспонденция КАК Д
        ГДЕ НЕ Д.ПометкаУдаления
            И (
                Д.Номер ПОДОБНО "%{safe_query}%"
                ИЛИ Д.Комментарий ПОДОБНО "%{safe_query}%"
                ИЛИ Д.Содержание ПОДОБНО "%{safe_query}%"
                ИЛИ Д.ТемаСлужебнойЗаписки ПОДОБНО "%{safe_query}%"
                ИЛИ Д.Организация.Наименование ПОДОБНО "%{safe_query}%"
                ИЛИ Д.Контрагент.Наименование ПОДОБНО "%{safe_query}%"
                ИЛИ Д.Ответственный.Наименование ПОДОБНО "%{safe_query}%"
            )
        УПОРЯДОЧИТЬ ПО Д.Дата УБЫВ"""
    table = app.NewObject("Query", query_text).Execute().Unload()
    documents: list[dict[str, Any]] = []
    for i in range(table.Count()):
        row = table.Get(i)
        fields: dict[str, str] = {}
        for alias, attr in (
            ("number", "Number"),
            ("date", "DocDate"),
            ("comment", "Comment"),
            ("org", "Org"),
            ("counterparty", "Counterparty"),
            ("content", "Content"),
            ("memo_subject", "MemoSubject"),
            ("email_from", "EmailFrom"),
            ("email_to", "EmailTo"),
            ("mail_to", "MailTo"),
            ("out_number", "OutNumber"),
            ("out_date", "OutDate"),
            ("responsible", "Responsible"),
            ("html_text", "HtmlText"),
        ):
            val = getattr(row, attr, None)
            fields[alias] = _safe_str(getattr(val, "Наименование", None) or val, 5000)
        documents.append(
            {
                "found": True,
                "ref": _safe_str(getattr(row, "Ref", "")),
                "number": _safe_str(getattr(row, "Number", "")),
                "fields": fields,
                "attachments": [],
            }
        )
    return {
        "found": bool(documents),
        "query": query,
        "documents": documents,
        "count": len(documents),
    }


def _method_name(env_key: str, default: str) -> str:
    """Получить имя COM-метода из env с безопасным default."""
    method_name = os.environ.get(env_key, default).strip() or default
    ensure_no_write_like_method(method_name)
    return method_name


def _invoke_method(session: Any, method_name: str, input_data: dict[str, Any]) -> Any:
    """Позвать COM-метод несколькими безопасными способами."""
    method = getattr(session, method_name, None)
    if method is None:
        raise OneCConnectionError(f"В COM-сессии нет метода {method_name!r}")

    payload = dict(input_data or {})
    attempts: list[tuple[tuple[Any, ...], dict[str, Any]]] = []
    if payload:
        attempts.append(((payload,), {}))
        attempts.append(((), payload))
        ordered = tuple(
            payload[key]
            for key in ("query", "document_ref", "task_ref", "number")
            if key in payload
        )
        if ordered:
            attempts.append((ordered, {}))
    attempts.append(((), {}))

    last_error: Exception | None = None
    for args, kwargs in attempts:
        try:
            return method(*args, **kwargs)
        except TypeError as exc:
            last_error = exc
        except Exception as exc:  # noqa: BLE001
            last_error = exc
    raise OneCConnectionError(
        f"Не удалось вызвать COM-метод {method_name!r}: {last_error}"
    ) from last_error


def _normalize_collection(raw: Any, preferred_key: str) -> list[dict[str, Any]]:
    """Привести ответ поиска документов/задач к стабильному виду."""
    items = _extract_collection(raw, preferred_key)
    collection = [_normalize_record(item) for item in items]
    if not collection and raw is not None:
        collection.append(_normalize_record(raw))
    return collection


def _normalize_single_record(raw: Any, *, preferred_key: str) -> dict[str, Any]:
    """Достать одну карточку документа/задачи из COM-ответа."""
    if isinstance(raw, dict):
        nested = raw.get(preferred_key)
        if isinstance(nested, dict):
            return _normalize_record(nested)
        if nested is not None and not isinstance(nested, (list, tuple)):
            return {"value": _normalize_value(nested)}
        if raw:
            return _normalize_record(raw)
    if isinstance(raw, list) and raw:
        first = raw[0]
        if isinstance(first, dict):
            return _normalize_record(first)
        return {"value": _normalize_value(first)}
    return _normalize_record(raw)


def _extract_collection(raw: Any, preferred_key: str) -> list[Any]:
    """Достать список записей из COM-ответа."""
    if isinstance(raw, dict):
        for key in (preferred_key, "documents", "tasks", "items", "rows", "value", "result"):
            value = raw.get(key)
            if isinstance(value, list):
                return value
            if isinstance(value, tuple):
                return list(value)
        nested = raw.get(preferred_key)
        if nested is not None and not isinstance(nested, (dict, list, tuple)):
            return [nested]
    if isinstance(raw, list):
        return raw
    if isinstance(raw, tuple):
        return list(raw)
    return []


def _normalize_record(value: Any) -> dict[str, Any]:
    """Преобразовать COM-объект в обычный dict."""
    if isinstance(value, dict):
        return {str(key): _normalize_value(item) for key, item in value.items()}
    if isinstance(value, list):
        return {"items": [_normalize_value(item) for item in value]}
    if isinstance(value, tuple):
        return {"items": [_normalize_value(item) for item in value]}
    if value is None:
        return {}
    return {"value": _normalize_value(value)}


def _normalize_value(value: Any) -> Any:
    """Рекурсивно привести значение COM к JSON-friendly виду."""
    if isinstance(value, dict):
        return {str(key): _normalize_value(item) for key, item in value.items()}
    if isinstance(value, list):
        return [_normalize_value(item) for item in value]
    if isinstance(value, tuple):
        return [_normalize_value(item) for item in value]
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _com_not_available(task: WorkerTask, message: str) -> WorkerResult:
    """Вернуть COM_NOT_AVAILABLE."""
    return WorkerResult(
        task_id=task.task_id,
        ok=False,
        error_type="COM_NOT_AVAILABLE",
        error_message=message,
    )

