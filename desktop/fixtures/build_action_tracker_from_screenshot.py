"""Recreate the Action Tracker annex from the provided screenshot."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

OUT = Path(__file__).resolve().parent / "Action Tracker — единый журнал Поручений.xlsx"

THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
HEADER_FILL = PatternFill("solid", fgColor="D6DCE4")
TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="1F4E79")
SUB_BLACK = InlineFont(rFont="Calibri", sz=12, b=True, color="000000")
SUB_RED = InlineFont(rFont="Calibri", sz=12, b=True, color="C00000")
HEADER_FONT = Font(name="Calibri", size=9, bold=True, color="1F1F1F")
CELL_FONT = Font(name="Calibri", size=10, color="000000")
LINK_FONT = Font(name="Calibri", size=10, color="0563C1", underline="single")
BLUE_FONT = Font(name="Calibri", size=10, color="0563C1")
CRIT_FILL = PatternFill("solid", fgColor="FF6B6B")
HIGH_FILL = PatternFill("solid", fgColor="F4B183")
CLOSED_FILL = PatternFill("solid", fgColor="92D050")
PROGRESS_FILL = PatternFill("solid", fgColor="9BC2E6")
WRAP = Alignment(wrap_text=True, vertical="center", horizontal="left")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")

HEADERS = [
    "ID",
    "Источник (CEO/ОКМ/комитет)",
    "Дата решения",
    "Поручение (результат/артефакт)",
    "Заказчик",
    "Владелец",
    "Срок",
    "Приоритет",
    "Статус",
    "Риск/эскалация",
    "Ссылка на результат/док-ты",
    "Комментарий",
]

WIDTHS = [12, 28, 14, 46, 12, 24, 42, 14, 26, 36, 46, 38]


def _file_link(path: str) -> str:
    return "file:///" + path.replace("\\", "/")


def build() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Action Tracker"

    ws.merge_cells("A1:L1")
    title = ws["A1"]
    title.value = "ПРИЛОЖЕНИЕ №1"
    title.font = TITLE_FONT
    title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    ws.merge_cells("A2:L2")
    subtitle = ws["A2"]
    subtitle.value = CellRichText(
        TextBlock(SUB_RED, "Action Tracker"),
        TextBlock(SUB_BLACK, " — единый журнал Поручений (шаблон)"),
    )
    subtitle.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(3, col, header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = THIN
        ws.column_dimensions[get_column_letter(col)].width = WIDTHS[col - 1]
    ws.row_dimensions[3].height = 32
    ws.auto_filter.ref = "A3:L8"
    ws.freeze_panes = "A4"
    ws.auto_filter.add_filter_column(0, [])

    rows = [
        {
            "id": "ACT-0001",
            "source": "Поручение",
            "date": "16.01.2026",
            "task": "О подготовке локальных нормативных документов к утверждению на заседании Совета директоров 27.01.2026",
            "customer": "CEO",
            "owner": "Донцова А.Е. Тищенко М.Н.",
            "deadline": "23.01.2026",
            "priority": "критический",
            "status": "CLOSED принято и закрыто",
            "risk": "",
            "result": r"..\..\..\..\..\1.Руководство\Амураль Игорь Борисович\2026\Реорганизация\совет директоров 27.01.26",
            "comment": "Утверждено на совете директоров 27.01.2026 г.",
            "result_is_link": True,
            "comment_blue": False,
        },
        {
            "id": "ACT-0002",
            "source": "вх.НП00-000528",
            "date": "29.01.2026",
            "task": "О проведении опытно-промышленных испытаний номер",
            "customer": "CEO",
            "owner": "Мегрелишвили МЭ Соломичева С.В.",
            "deadline_rich": True,
            "priority": "высокий",
            "status": "IN PROGRESS в работе",
            "risk": "НП00-001585.msg",
            "result": "",
            "comment": "НП00-000528.msg",
            "risk_is_link": True,
            "comment_is_link": True,
        },
        {
            "id": "ACT-0003",
            "source": "Распоряжение ГК00-000002 от 02.02.26",
            "date": "02.02.2026",
            "task": "О подключении электронного документооборота в ООО Авион и ООО ИТЦ. На отдельной базе установить электронный документооборот в ООО Авион и ООО ИТЦ, а также",
            "customer": "CEO",
            "owner": "Мегрелишвили МЭ Ростовцева А.В.",
            "deadline": "06.02.2026 перенос срока на 09.02.26 до 13.02.26 на 17.02.26 на 18.02.26 на 24.02.26",
            "priority": "критический",
            "status": "CLOSED принято и закрыто",
            "risk": "",
            "result": "",
            "comment": "закрыто 27.02.26",
            "comment_blue": True,
        },
        {
            "id": "ACT-0004",
            "source": "Распоряжение ГК00-000002 от 02.02.26",
            "date": "02.02.2026",
            "task": "Организовать обучение руководителей",
            "customer": "CEO",
            "owner": "Мегрелишвили МЭ Ростовцева А.В.",
            "deadline": "13.02.2026 перенос срока на 17.02.26 на 18.02.26",
            "priority": "критический",
            "status": "CLOSED принято и закрыто",
            "risk": "",
            "result": "",
            "comment": "закрыто 18.02.26",
            "comment_blue": True,
        },
        {
            "id": "ACT-0005",
            "source": "номер вх. НП00-000684",
            "date": "05.02.2026",
            "task": 'Запрос ТКП от ПАО "Газпром"',
            "customer": "CEO",
            "owner": "Милютенко О.Н.",
            "deadline": "10.02.2026",
            "priority": "критический",
            "status": "CLOSED принято и закрыто",
            "risk": r"C:\Users\a.kudryavtseva\Desktop\Исх. № НП00-000825",
            "result": r"C:\Users\a.kudryavtseva\Desktop\2026.02.10 МЦР УИРГ ЗПР 132 этап 1,2,3,4.pdf",
            "comment": r"C:\Users\a.kudryavtseva\Desktop\письмо с резолюцией ПСД.pdf",
            "risk_is_link": True,
            "result_is_link": True,
            "comment_is_link": True,
        },
    ]

    for offset, item in enumerate(rows):
        r = 4 + offset
        values = [
            item["id"],
            item["source"],
            item["date"],
            item["task"],
            item["customer"],
            item["owner"],
            item.get("deadline", ""),
            item["priority"],
            item["status"],
            item["risk"],
            item["result"],
            item["comment"],
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(r, col, value)
            cell.font = CELL_FONT
            cell.alignment = CENTER if col in {1, 3, 5, 8, 9} else WRAP
            cell.border = THIN

        if item.get("deadline_rich"):
            deadline = ws.cell(r, 7)
            deadline.value = CellRichText(
                TextBlock(
                    InlineFont(rFont="Calibri", sz=10, color="000000"),
                    "в период с 25 по 30 число каждого месяца, след 29.05.26 ",
                ),
                TextBlock(
                    InlineFont(rFont="Calibri", sz=10, color="FF0000"),
                    "на 30.06.26 на 28.07.26 на 31.08.26",
                ),
            )
            deadline.alignment = WRAP

        priority = ws.cell(r, 8)
        priority.fill = CRIT_FILL if item["priority"] == "критический" else HIGH_FILL
        priority.alignment = CENTER

        status = ws.cell(r, 9)
        status.fill = CLOSED_FILL if item["status"].startswith("CLOSED") else PROGRESS_FILL
        status.alignment = CENTER

        for col, key, flag in (
            (10, "risk", "risk_is_link"),
            (11, "result", "result_is_link"),
            (12, "comment", "comment_is_link"),
        ):
            if not item.get(flag) or not item[key]:
                continue
            cell = ws.cell(r, col)
            target = item[key]
            if target.startswith("C:\\"):
                cell.hyperlink = Hyperlink(ref=cell.coordinate, target=_file_link(target))
            cell.font = LINK_FONT

        if item.get("comment_blue") and not item.get("comment_is_link"):
            ws.cell(r, 12).font = BLUE_FONT

        ws.row_dimensions[r].height = 48

    ws.row_dimensions[5].height = 56
    ws.row_dimensions[6].height = 64

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:3"
    ws.sheet_view.showGridLines = False

    wb.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
