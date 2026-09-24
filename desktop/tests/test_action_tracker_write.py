"""excel.write_action_tracker writes every 1C row without the model listing them."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from app.tools.ac.action_tracker import ActionTrackerError, write_action_tracker
from app.tools.ac.dispatch import invoke_ac_tool


def _dump(folder: Path, name: str, payload: dict) -> None:
    folder.mkdir(parents=True, exist_ok=True)
    (folder / name).write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")


def test_write_action_tracker_keeps_every_card(tmp_path: Path) -> None:
    results = tmp_path / "tool_results"
    _dump(
        results,
        "onec.erp_assignments_aaa.json",
        {
            "count": 2,
            "truncated": False,
            "assignments": [
                {
                    "number": "АСТ00-00002",
                    "date": "2026-09-01T10:00:00",
                    "due": "2026-09-24T00:00:00",
                    "status": "Создано",
                    "customer": "Амураль Игорь Борисович",
                    "open": True,
                    "overdue": True,
                    "lines": [{"text": "Первое поручение", "executor": "Иванов Иван"}],
                },
                {
                    "number": "АСТ00-00001",
                    "date": "2026-08-01T10:00:00",
                    "due": "2026-08-10T00:00:00",
                    "status": "Закрыт",
                    "customer": "Амураль Игорь Борисович",
                    "open": False,
                    "overdue": False,
                    "topic": "Закрытое",
                },
            ],
        },
    )
    _dump(
        results,
        "onec.meeting_protocols_bbb.json",
        {
            "count": 2,
            "truncated": False,
            "protocols": [
                {
                    "number": "ПСД_001_О_001",
                    "date": "2026-09-02T00:00:00",
                    "status": "НаИсполнении",
                    "meeting_topic": "Совет",
                },
                {
                    "number": "ПСД_001_О_002",
                    "date": "2026-01-02T00:00:00",
                    "status": "Закрыт",
                    "meeting_topic": "Старый",
                },
            ],
        },
    )
    written = write_action_tracker(tmp_path, "ActionTracker.xlsx")
    assert written["assignments"] == 2
    assert written["protocols"] == 2
    assert written["rows"] == 4
    book = load_workbook(tmp_path / "ActionTracker.xlsx", data_only=True)
    sheet = book.active
    sources = []
    for row in sheet.iter_rows(min_row=1, values_only=True):
        if row and str(row[1] or "").startswith(("АСТ", "ПСД")):
            sources.append(row[1])
    book.close()
    assert sources == ["АСТ00-00002", "АСТ00-00001", "ПСД_001_О_001", "ПСД_001_О_002"]


def test_write_action_tracker_refuses_truncated_extract(tmp_path: Path) -> None:
    results = tmp_path / "tool_results"
    _dump(results, "onec.erp_assignments_a.json", {"assignments": [], "truncated": False})
    _dump(
        results,
        "onec.meeting_protocols_b.json",
        {"protocols": [{"number": "ПСД_1"}], "truncated": True},
    )
    try:
        write_action_tracker(tmp_path)
    except ActionTrackerError as exc:
        assert "обрезана" in str(exc)
    else:
        raise AssertionError("truncated extract must not be written")


def test_tool_writes_into_agent_workspace(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path))
    from app.tools.ac import dispatch

    dispatch._REGISTRY = None
    dispatch._workspaces_root = lambda: tmp_path / "Constructor" / "agent_workspaces"  # type: ignore[method-assign]
    workspace = tmp_path / "Constructor" / "agent_workspaces" / "agent-1" / "tool_results"
    _dump(
        workspace,
        "onec.erp_assignments_a.json",
        {
            "truncated": False,
            "assignments": [
                {
                    "number": "АСТ00-00010",
                    "status": "Создано",
                    "customer": "Амураль Игорь Борисович",
                    "open": True,
                    "topic": "Карточка",
                }
            ],
        },
    )
    _dump(
        workspace,
        "onec.meeting_protocols_b.json",
        {
            "truncated": False,
            "protocols": [
                {"number": "ПСД_001_О_010", "status": "Закрыт", "meeting_topic": "Протокол"}
            ],
        },
    )
    result = invoke_ac_tool(
        "excel.write_action_tracker",
        {"filename": "ActionTracker.xlsx", "workflow_id": "agent-1"},
    )
    assert result["assignments"] == 1
    assert result["protocols"] == 1
    assert (tmp_path / "Constructor" / "agent_workspaces" / "agent-1" / "ActionTracker.xlsx").is_file()
