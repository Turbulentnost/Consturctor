from __future__ import annotations

from pathlib import Path

import pytest

from app.tools.ac.agent_workspace import AgentWorkspaceResolver
from app.tools.ac.office_read import OfficeReadFileTool
from app.tools.ac.readable_files import (
    artifact_cache_dir,
    kind_for_suffix,
    read_tool_for_suffix,
    resolve_readable_path,
)


def _tool(tmp_path: Path) -> OfficeReadFileTool:
    return OfficeReadFileTool(AgentWorkspaceResolver(tmp_path))


def _write_docx(path: Path, text: str) -> None:
    import docx

    document = docx.Document()
    document.add_paragraph(text)
    document.save(str(path))


def _write_pdf(path: Path, text: str) -> None:
    import fitz

    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    document.save(str(path))
    document.close()


def test_kind_and_reader_for_suffixes() -> None:
    assert kind_for_suffix(".docx") == "word"
    assert kind_for_suffix(".pdf") == "pdf"
    assert kind_for_suffix(".jpg") == "image"
    assert kind_for_suffix(".xlsx") == "excel"
    assert read_tool_for_suffix(".docx") == "office.read_file"
    assert read_tool_for_suffix(".pdf") == "office.read_file"
    assert read_tool_for_suffix(".png") == "office.read_file"
    assert read_tool_for_suffix(".xlsx") == "excel.read_workbook"
    assert read_tool_for_suffix(".txt") == ""


def test_read_docx_from_workspace(tmp_path: Path) -> None:
    resolver = AgentWorkspaceResolver(tmp_path)
    workspace = resolver.for_agent("wf-office")
    source = workspace.directory / "materials" / "attachments" / "note.docx"
    source.parent.mkdir(parents=True)
    _write_docx(source, "тема поручения АСТ00")

    result = _tool(tmp_path).execute(
        {"workflow_id": "wf-office", "filename": "note.docx"}
    )
    assert result.ok
    assert result.output_data["kind"] == "word"
    assert "тема поручения АСТ00" in result.output_data["text"]


def test_read_pdf_from_workspace(tmp_path: Path) -> None:
    resolver = AgentWorkspaceResolver(tmp_path)
    workspace = resolver.for_agent("wf-office")
    source = workspace.directory / "scan.pdf"
    _write_pdf(source, "protocol AST00")

    result = _tool(tmp_path).execute(
        {"workflow_id": "wf-office", "filename": "scan.pdf"}
    )
    assert result.ok
    assert result.output_data["kind"] == "pdf"
    assert "protocol AST00" in result.output_data["text"]


def _write_image(path: Path, label: str = "scan") -> None:
    import fitz

    document = fitz.open()
    page = document.new_page(width=240, height=120)
    page.insert_text((16, 64), label)
    pix = page.get_pixmap()
    pix.save(str(path))
    document.close()


def _write_scan_pdf(path: Path) -> None:
    import fitz

    document = fitz.open()
    page = document.new_page()
    page.draw_rect(page.rect, color=(0.1, 0.1, 0.1), width=3)
    document.save(str(path))
    document.close()


def test_read_image_uses_cursor_sdk_vision(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    resolver = AgentWorkspaceResolver(tmp_path)
    workspace = resolver.for_agent("wf-office")
    source = workspace.directory / "photo.png"
    _write_image(source, "photo")
    monkeypatch.setattr(
        "app.attachment_text._ocr",
        lambda path: (_ for _ in ()).throw(AssertionError("LM Studio OCR не должен вызываться")),
    )

    result = _tool(tmp_path).execute(
        {"workflow_id": "wf-office", "filename": "photo.png"}
    )
    assert result.ok
    assert result.output_data["kind"] == "image"
    assert result.output_data["ocr"] is False
    assert result.output_data["vision"] is True
    assert result.output_data["vision_source"] == "cursor_sdk"
    pages = result.output_data["vision_pages"]
    assert pages
    assert Path(workspace.directory / pages[0]["path"]).is_file()
    assert result.output_data["vision_source"] == "cursor_sdk"
    assert "СПИСОК_ГОТОВ" not in result.output_data["next_step"]
    assert "KPI" not in result.output_data["next_step"]
    assert "прочитан" in result.output_data["next_step"]


def test_read_file_from_artifact_cache(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("tempfile.gettempdir", lambda: str(tmp_path))
    cache = artifact_cache_dir()
    cache.mkdir(parents=True)
    source = cache / "RaspPriemDokum.docx"
    _write_docx(source, "лист записи")

    result = _tool(tmp_path).execute(
        {
            "workflow_id": "wf-office",
            "filename": str(source),
        }
    )
    assert result.ok
    assert "лист записи" in result.output_data["text"]
    assert Path(result.output_data["path"]) == source.resolve()


def test_read_file_finds_cache_by_basename(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("tempfile.gettempdir", lambda: str(tmp_path))
    cache = artifact_cache_dir()
    cache.mkdir(parents=True)
    source = cache / "road-map.pdf"
    _write_pdf(source, "road map")

    result = _tool(tmp_path).execute(
        {"workflow_id": "wf-office", "filename": "road-map.pdf"}
    )
    assert result.ok
    assert "road map" in result.output_data["text"]


def test_resolve_rejects_path_outside_workspace_and_cache(tmp_path: Path) -> None:
    workspace = AgentWorkspaceResolver(tmp_path).for_agent("wf-office")
    outsider = tmp_path / "secret.docx"
    _write_docx(outsider, "тайна")
    with pytest.raises(Exception, match="constructor-onec-artifacts"):
        resolve_readable_path(workspace, str(outsider))


def test_xlsx_is_not_office_read(tmp_path: Path) -> None:
    resolver = AgentWorkspaceResolver(tmp_path)
    workspace = resolver.for_agent("wf-office")
    source = workspace.directory / "plan.xlsx"
    source.write_bytes(b"PK")

    result = _tool(tmp_path).execute(
        {"workflow_id": "wf-office", "filename": "plan.xlsx"}
    )
    assert not result.ok
    assert result.error_type == "UNSUPPORTED_TYPE"
    assert "excel.read_workbook" in result.error_message


def test_old_doc_is_read(tmp_path: Path, monkeypatch) -> None:
    import app.attachment_text as attachment_text

    monkeypatch.setattr(attachment_text, "_read_doc_pieces", lambda path: "Извещение СТО-28-020\nУтвердил И.Б. Амураль")
    resolver = AgentWorkspaceResolver(tmp_path)
    workspace = resolver.for_agent("wf-office")
    (workspace.directory / "legacy.doc").write_bytes(b"\xd0\xcf\x11\xe0")

    result = _tool(tmp_path).execute(
        {"workflow_id": "wf-office", "filename": "legacy.doc"}
    )
    assert result.ok
    assert result.output_data["kind"] == "word"
    assert "Утвердил И.Б. Амураль" in result.output_data["text"]


def _write_scan_pdf_pages(path: Path, count: int) -> None:
    import fitz

    document = fitz.open()
    for _ in range(count):
        page = document.new_page()
        page.draw_rect(page.rect, color=(0.1, 0.1, 0.1), width=3)
    document.save(str(path))
    document.close()


def test_two_scans_in_one_run_both_return_pages(tmp_path: Path) -> None:
    resolver = AgentWorkspaceResolver(tmp_path)
    workspace = resolver.for_agent("wf-closure")
    _write_scan_pdf(workspace.directory / "card.pdf")
    _write_scan_pdf(workspace.directory / "act.pdf")
    tool = _tool(tmp_path)

    first = tool.execute({"workflow_id": "wf-closure", "filename": "card.pdf"})
    second = tool.execute({"workflow_id": "wf-closure", "filename": "act.pdf"})

    assert first.ok and second.ok
    assert first.output_data["vision_pages"]
    assert second.output_data["vision_pages"]
    assert second.output_data.get("cached") is not True
    assert "СПИСОК_ГОТОВ" not in str(second.output_data.get("next_step") or "")


def test_docx_table_is_text(tmp_path: Path) -> None:
    import docx

    resolver = AgentWorkspaceResolver(tmp_path)
    workspace = resolver.for_agent("wf-closure")
    source = workspace.directory / "act.docx"
    document = docx.Document()
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Утверждено"
    table.rows[0].cells[1].text = "СТО-18-231"
    document.save(str(source))

    result = _tool(tmp_path).execute({"workflow_id": "wf-closure", "filename": "act.docx"})
    assert result.ok
    assert "Утверждено" in result.output_data["text"]
    assert "СТО-18-231" in result.output_data["text"]
    assert result.output_data.get("vision") is not True


def test_scan_pdf_uses_cursor_sdk_vision(tmp_path: Path) -> None:
    resolver = AgentWorkspaceResolver(tmp_path)
    workspace = resolver.for_agent("wf-office")
    source = workspace.directory / "scan.pdf"
    _write_scan_pdf(source)

    result = _tool(tmp_path).execute({"workflow_id": "wf-office", "filename": "scan.pdf"})
    assert result.ok
    assert result.output_data["kind"] == "pdf"
    assert result.output_data["vision"] is True
    assert result.output_data["page_count"] == 1
    pages = result.output_data["vision_pages"]
    assert len(pages) == 1
    assert Path(workspace.directory / pages[0]["path"]).is_file()


def test_text_pdf_does_not_attach_vision(tmp_path: Path) -> None:
    resolver = AgentWorkspaceResolver(tmp_path)
    workspace = resolver.for_agent("wf-office")
    source = workspace.directory / "note.pdf"
    _write_pdf(
        source,
        "protocol AST00 native text layer for a real document not a scan page extra words here",
    )

    result = _tool(tmp_path).execute({"workflow_id": "wf-office", "filename": "note.pdf"})
    assert result.ok
    assert result.output_data.get("vision") is False
    assert "vision_pages" not in result.output_data
    assert "protocol AST00" in result.output_data["text"]


def test_unreadable_image_explains_sdk_render(tmp_path: Path) -> None:
    resolver = AgentWorkspaceResolver(tmp_path)
    workspace = resolver.for_agent("wf-office")
    source = workspace.directory / "photo.jpg"
    source.write_bytes(b"\xff\xd8\xff")

    result = _tool(tmp_path).execute({"workflow_id": "wf-office", "filename": "photo.jpg"})
    assert not result.ok
    assert result.error_type == "UNREADABLE"
    assert "картинка" in result.error_message
    assert "Cursor SDK" in result.error_message
    assert "LM Studio" not in result.error_message


def test_empty_filename_fails(tmp_path: Path) -> None:
    result = _tool(tmp_path).execute({"workflow_id": "wf-office"})
    assert not result.ok
    assert result.error_type == "INVALID_FILENAME"


def test_reader_does_not_speak_kpi(tmp_path: Path) -> None:
    resolver = AgentWorkspaceResolver(tmp_path)
    workspace = resolver.for_agent("kpi-build-demo")
    source = workspace.directory / "scan.pdf"
    _write_scan_pdf_pages(source, 3)

    result = _tool(tmp_path).execute(
        {"workflow_id": "kpi-build-demo", "filename": "scan.pdf", "max_pages": 1}
    )
    assert result.ok
    pages = result.output_data["vision_pages"]
    assert [item["page"] for item in pages] == [1]
    step = str(result.output_data.get("next_step") or "")
    assert "СПИСОК_ГОТОВ" not in step
    assert "KPI" not in step
    assert "start_page=2" in step


def test_second_file_is_not_blocked_by_the_first(tmp_path: Path) -> None:
    resolver = AgentWorkspaceResolver(tmp_path)
    workspace = resolver.for_agent("kpi-build-cache")
    _write_scan_pdf(workspace.directory / "first.pdf")
    _write_scan_pdf(workspace.directory / "second.pdf")
    tool = _tool(tmp_path)
    first = tool.execute({"workflow_id": "kpi-build-cache", "filename": "first.pdf"})
    second = tool.execute({"workflow_id": "kpi-build-cache", "filename": "second.pdf"})
    assert first.ok and second.ok
    assert first.output_data["vision_pages"]
    assert second.output_data["vision_pages"]
    assert "СПИСОК_ГОТОВ" not in str(first.output_data.get("next_step") or "")
    assert "СПИСОК_ГОТОВ" not in str(second.output_data.get("next_step") or "")


def test_scan_pdf_vision_reads_one_page_then_next(tmp_path: Path) -> None:
    resolver = AgentWorkspaceResolver(tmp_path)
    workspace = resolver.for_agent("wf-pages")
    source = workspace.directory / "scan.pdf"
    _write_scan_pdf_pages(source, 6)
    tool = _tool(tmp_path)
    first = tool.execute(
        {"workflow_id": "wf-pages", "filename": "scan.pdf", "max_pages": 1, "start_page": 1}
    )
    second = tool.execute(
        {
            "workflow_id": "wf-pages",
            "filename": "scan.pdf",
            "max_pages": 1,
            "start_page": first.output_data["next_start"],
        }
    )
    assert first.ok and second.ok
    assert [item["page"] for item in first.output_data["vision_pages"]] == [1]
    assert first.output_data["next_start"] == 2
    assert [item["page"] for item in second.output_data["vision_pages"]] == [2]
    assert second.output_data["next_start"] == 3


def test_office_read_folder_picks_pdf_inside(tmp_path: Path) -> None:
    resolver = AgentWorkspaceResolver(tmp_path)
    workspace = resolver.for_agent("wf-folder")
    folder = workspace.directory / "materials" / "attachments"
    folder.mkdir(parents=True)
    _write_scan_pdf_pages(folder / "001_method.pdf", 2)

    result = _tool(tmp_path).execute(
        {"workflow_id": "wf-folder", "filename": "materials/attachments", "max_pages": 2}
    )
    assert result.ok
    assert result.output_data["filename"] == "001_method.pdf"
    assert [item["page"] for item in result.output_data["vision_pages"]] == [1, 2]


def test_one_page_scan_reads_that_page(tmp_path: Path) -> None:
    resolver = AgentWorkspaceResolver(tmp_path)
    workspace = resolver.for_agent("kpi-build-one")
    source = workspace.directory / "scan.pdf"
    _write_scan_pdf(source)

    result = _tool(tmp_path).execute(
        {"workflow_id": "kpi-build-one", "filename": "scan.pdf"}
    )
    assert result.ok
    pages = result.output_data["vision_pages"]
    assert [item["page"] for item in pages] == [1]
