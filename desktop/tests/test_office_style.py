from pathlib import Path

from app.tools.ac.agent_workspace import AgentWorkspaceResolver
from app.tools.ac.document_tools import ReportExportDocumentTool
from app.tools.ac.excel_tools import (
    ExcelCreateWorkbookTool,
    ExcelReadWorkbookTool,
)
from app.tools.ac.office_style import THEMES
from app.tools.ac.office_tools import OfficeFormatDocumentTool


def test_create_workbook_applies_header_theme(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    from app.tools.ac.office_style import data_start_row

    resolver = AgentWorkspaceResolver(tmp_path)
    created = ExcelCreateWorkbookTool(resolver).execute(
        {
            "workflow_id": "wf-style",
            "filename": "plan.xlsx",
            "headers": ["Исполнитель", "Статус", "Сумма"],
            "rows": [["Иванов", "Выполнено", 1200], ["Петров", "Просрочено", 80]],
            "theme": "navy",
        }
    )
    assert created.ok
    path = Path(created.output_data["path"])
    workbook = load_workbook(path)
    try:
        sheet = workbook.active
        start = data_start_row(workbook, sheet)
        assert sheet.cell(1, 1).value == "plan"
        assert sheet.cell(start, 1).value == "Исполнитель"
        fill = str(sheet.cell(start, 1).fill.fgColor.rgb or "")
        assert THEMES["navy"].header in fill.upper()
        assert sheet.auto_filter.ref
        assert sheet.freeze_panes == f"A{start + 1}"
        status_fill = str(sheet.cell(start + 1, 2).fill.fgColor.rgb or "")
        assert "1E8449" in status_fill.upper()
    finally:
        workbook.close()


def test_create_workbook_always_adds_banner_even_without_title(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    resolver = AgentWorkspaceResolver(tmp_path)
    created = ExcelCreateWorkbookTool(resolver).execute(
        {
            "workflow_id": "wf-style",
            "filename": "daily_control.xlsx",
            "headers": ["Код"],
            "rows": [["1"]],
            "styled": False,
        }
    )
    assert created.ok
    workbook = load_workbook(Path(created.output_data["path"]))
    try:
        assert workbook.active["A1"].value == "daily control"
        fill = str(workbook.active["A1"].fill.fgColor.rgb or "")
        assert THEMES["navy"].primary in fill.upper()
    finally:
        workbook.close()
    read = ExcelReadWorkbookTool(resolver).execute(
        {"workflow_id": "wf-style", "filename": "daily_control.xlsx"}
    )
    assert read.ok
    assert read.output_data["rows"][0][0] == "Код"


def test_create_workbook_title_skipped_on_read(tmp_path: Path) -> None:
    resolver = AgentWorkspaceResolver(tmp_path)
    created = ExcelCreateWorkbookTool(resolver).execute(
        {
            "workflow_id": "wf-style",
            "filename": "banner.xlsx",
            "title": "Контроль поручений",
            "subtitle": "за неделю",
            "kpis": [{"label": "Всего", "value": 2}],
            "headers": ["Код", "Тема"],
            "rows": [["1", "Первое"], ["2", "Второе"]],
        }
    )
    assert created.ok
    read = ExcelReadWorkbookTool(resolver).execute(
        {"workflow_id": "wf-style", "filename": "banner.xlsx"}
    )
    assert read.ok
    assert read.output_data["rows"][0][0] == "Код"
    assert read.output_data["row_count"] == 3


def test_office_format_creates_excel_and_restyles(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    resolver = AgentWorkspaceResolver(tmp_path)
    tool = OfficeFormatDocumentTool(resolver)
    created = tool.execute(
        {
            "workflow_id": "wf-style",
            "filename": "pretty.xlsx",
            "title": "Сводка",
            "theme": "forest",
            "headers": ["A", "B"],
            "rows": [{"A": "x", "B": "y"}],
        }
    )
    assert created.ok
    assert created.output_data["action"] == "create"
    restyled = tool.execute(
        {
            "workflow_id": "wf-style",
            "filename": "pretty.xlsx",
            "theme": "wine",
            "title": "Сводка",
        }
    )
    assert restyled.ok
    assert restyled.output_data["action"] == "restyle"
    from app.tools.ac.office_style import data_start_row

    workbook = load_workbook(Path(restyled.output_data["path"]))
    try:
        start = data_start_row(workbook, workbook.active)
        fill = str(workbook.active.cell(start, 1).fill.fgColor.rgb or "")
        assert "7B2D4B" in fill.upper()
    finally:
        workbook.close()


def test_export_document_styles_word_table(tmp_path: Path) -> None:
    import importlib.util

    if importlib.util.find_spec("docx") is None:
        return
    from docx import Document
    from docx.oxml.ns import qn

    resolver = AgentWorkspaceResolver(tmp_path)
    result = ReportExportDocumentTool(resolver).execute(
        {
            "workflow_id": "wf-style",
            "filename": "report.docx",
            "title": "Протокол",
            "summary": "Кратко по итогам",
            "theme": "navy",
            "sections": [{"heading": "Решения", "body": "- Первое\n- Второе"}],
            "table": {"headers": ["Пункт", "Ответственный"], "rows": [["1", "Иванов"]]},
        }
    )
    assert result.ok
    assert result.output_data["format"] == "docx"
    document = Document(result.output_data["path"])
    assert any("Протокол" in (p.text or "") for p in document.paragraphs)
    assert document.tables
    header = document.tables[-1].rows[0].cells[0]
    fill = header._tc.get_or_add_tcPr().find(qn("w:shd")).get(qn("w:fill"))
    assert fill.upper() == "1B4F72"


def test_office_format_word_from_sections(tmp_path: Path) -> None:
    import importlib.util

    if importlib.util.find_spec("docx") is None:
        return
    resolver = AgentWorkspaceResolver(tmp_path)
    result = OfficeFormatDocumentTool(resolver).execute(
        {
            "workflow_id": "wf-style",
            "filename": "memo.docx",
            "title": "Служебная записка",
            "theme": "graphite",
            "sections": [{"heading": "Суть", "body": "Текст записки"}],
        }
    )
    assert result.ok
    assert result.output_data["format"] == "docx"
    assert Path(result.output_data["path"]).is_file()
